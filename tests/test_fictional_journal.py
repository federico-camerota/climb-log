# SPDX-License-Identifier: GPL-3.0-or-later
"""The fictional journal workbooks the importer and hand-entry tests use (D-034, D-060)."""

import datetime as dt
import tempfile
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from fictional_journal import LEGACY, LegacyJournal, legacy_journal, main, new_journal
from openpyxl.worksheet.worksheet import Worksheet
from parity import recalculate, soffice, work_dir

from climblog.modules import load_module
from climblog.storage import XlsxBackend, load_data

JOURNAL = Path(__file__).parent.parent / "modules" / "journal.yaml"


@pytest.fixture
def legacy(tmp_path: Path) -> LegacyJournal:
    return legacy_journal(tmp_path / "legacy.xlsx", seed=1)


def cells(path: Path) -> dict[str, list[tuple[Any, ...]]]:
    book = openpyxl.load_workbook(path)
    return {s.title: list(s.iter_rows(values_only=True)) for s in book.worksheets}


def test_the_legacy_layout(legacy: LegacyJournal) -> None:
    book = openpyxl.load_workbook(legacy.path)
    assert book.sheetnames == ["Start Here", *(tab.name for tab in LEGACY)]
    assert {n: d.attr_text for n, d in book.defined_names.items()} == {
        "LowMax": "'Start Here'!$F$46",
        "WatchMax": "'Start Here'!$F$47",
        "AttnMin": "'Start Here'!$F$48",
        "NoteMin": "'Start Here'!$F$49",
        "TrendDelta": "'Start Here'!$F$50",
    }
    session = book["Session"]
    assert isinstance(session, Worksheet)
    assert [c.value for c in session[4]][:4] == [
        "Date",
        "Venue",
        "Indoor / outdoor",
        "Session type",
    ]
    assert session["K5"].value.startswith('=IF(A5="","",CHOOSE(WEEKDAY(A5,2)')
    assert session["M6"].value == '=IF(A6="","",IF(COUNT(A$4:A5)=0,"",A6-MAX(A$4:A5)))'
    wellbeing = book["Wellbeing"]
    assert isinstance(wellbeing, Worksheet)
    assert wellbeing["A304"].value == '=IF(Session!$A304="","",Session!$A304)'  # linked by row
    assert wellbeing["A305"].value is None
    assert wellbeing.freeze_panes == "B5"


def test_the_same_seed_gives_the_same_workbook(tmp_path: Path, legacy: LegacyJournal) -> None:
    again = legacy_journal(tmp_path / "again.xlsx", seed=1)
    other = legacy_journal(tmp_path / "other.xlsx", seed=2)
    assert cells(again.path) == cells(legacy.path)
    assert again.sessions == legacy.sessions
    assert cells(other.path) != cells(legacy.path)


def test_the_legacy_data_has_the_hand_entry_cases(legacy: LegacyJournal) -> None:
    days = [s["Date"] for s in legacy.sessions]
    assert len(days) == 30 and len(legacy.rest_days) == 8
    assert len(set(days)) < len(days)  # two sessions on one day
    sleeps = [s["wb:Sleep\n(hours)"] for s in legacy.sessions]
    assert any(isinstance(v, str) and "," in v for v in sleeps)  # "7,5" typed as text
    notes = [v for s in legacy.sessions for k, v in s.items() if "note" in k.lower() and v]
    assert any(n.startswith("'") for n in notes) and any(n.startswith("=") for n in notes)
    book = openpyxl.load_workbook(legacy.path)
    typed = [
        c
        for row in book["Wellbeing"].iter_rows(min_row=5, max_row=34)
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=non")
    ]
    assert typed and all(c.data_type == "s" for c in typed)  # text, not a formula


def test_the_new_layout_loads_with_the_expected_issues(tmp_path: Path) -> None:
    fixture = new_journal(tmp_path / "new.xlsx")
    data = load_data(load_module(JOURNAL), XlsxBackend(fixture.path))
    row_issues = [(i.tab, i.row, i.field) for i in data.issues if i.row is not None]
    assert row_issues == fixture.issues
    sessions = data.entities["session"].records
    assert [(s["session_no"], s["session_type"], s["setting"]) for s in sessions] == [
        (1, "Climbing", "Indoor"),  # blank = 1; choices typed in another case
        (2, "Hangboard", "Home wall"),
        (1, "Climbing", "Outdoor"),
    ]
    assert [s.get("notes") for s in sessions] == [
        "'po' stanco",
        "=non una formula",
        "[ESEMPIO] così",
    ]
    assert sessions[0]["wb_sleep_hours"] == 7.5  # "7,5", read from a reordered tab


def test_make_fixtures_writes_both(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    main(["fictional_journal.py", str(tmp_path)])
    assert sorted(p.name for p in tmp_path.iterdir()) == ["legacy-journal.xlsx", "new-journal.xlsx"]
    assert "fictional" in capsys.readouterr().out


@pytest.mark.skipif(soffice() is None, reason="no working LibreOffice")
def test_the_legacy_formulas_compute_without_errors() -> None:
    command = soffice()
    assert command is not None
    with tempfile.TemporaryDirectory(dir=work_dir()) as folder:
        written = legacy_journal(Path(folder) / "legacy.xlsx", seed=3)
        sheets = recalculate(command, written.path)
    errors = [
        (tab, value)
        for tab, rows in sheets.items()
        for row in rows
        for value in row
        if isinstance(value, str) and value.startswith("#")
    ]
    assert errors == []
    session = sheets["Session"]
    assert session[4][9] in ("Complete",) or str(session[4][9]).startswith("Missing: ")
    assert isinstance(sheets["Wellbeing"][4][0], dt.datetime)  # the date linked from Session
