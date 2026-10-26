# SPDX-License-Identifier: GPL-3.0-or-later
# ruff: noqa: RUF001, RUF003  (the legacy labels use en dashes, as in "(1-10)")
"""Importing the legacy journal, and its parity report, on fictional workbooks (D-061)."""

import datetime as dt
import shutil
import tempfile
from pathlib import Path

import openpyxl
import pytest
from fictional_journal import LegacyJournal, legacy_journal
from parity import recalculate, soffice, work_dir
from typer.testing import CliRunner

from climblog import cli
from climblog.expr import Value
from climblog.importers.legacy_journal import (
    ImportRefused,
    LegacyLayoutError,
    import_legacy,
    parity,
    read_legacy,
)
from climblog.modules import Module, load_module
from climblog.modules.compute import compute_module
from climblog.storage import StorageError, XlsxBackend, load_data

JOURNAL = Path(__file__).parent.parent / "modules" / "journal.yaml"
runner = CliRunner()


@pytest.fixture(scope="module")
def module() -> Module:
    return load_module(JOURNAL)


@pytest.fixture
def legacy(tmp_path: Path) -> LegacyJournal:
    return legacy_journal(tmp_path / "Climbing_Journal.xlsx", seed=1)


def test_every_session_and_rest_day_is_read(module: Module, legacy: LegacyJournal) -> None:
    data = read_legacy(legacy.path, module)
    assert (data.count("session"), data.count("rest_day")) == (30, 8)
    assert data.skipped == [] and data.unread == []
    assert data.settings == {
        "low_max": 4,
        "watch_max": 6,
        "attn_min": 8,
        "note_min": 6,
        "trend_delta": 0.5,
    }
    first, written = data.records[0], legacy.sessions[0]
    assert first.values["date"] == written["Date"].date()
    assert first.values.get("venue") == written["Venue"]
    assert first.values.get("wb_sleep_quality") == written["wb:Sleep quality\n(1–10)"]
    sheet = openpyxl.load_workbook(legacy.path)["Physical"]
    assert first.values.get("ph_pre_fingers") == sheet["B5"].value  # pre, not post (H5)
    assert first.values.get("ph_post_fingers") == sheet["H5"].value
    assert first.rows["Wellbeing"] == 5  # linked to the Session row


def test_values_are_read_as_the_module_reads_them(module: Module, legacy: LegacyJournal) -> None:
    data = read_legacy(legacy.path, module)
    sessions = [r for r in data.records if r.entity == "session"]
    for record, written in zip(sessions, legacy.sessions, strict=True):
        typed = written["wb:Sleep\n(hours)"]
        if isinstance(typed, str):  # "7,5" typed as text
            assert record.values["wb_sleep_hours"] == float(typed.replace(",", "."))
        for key, field in (("Notes", "notes"), ("wb:Pre-session note", "wb_note_pre")):
            if written[key] is not None:
                assert record.values[field] == written[key]  # "'po'…", "=non…" as typed
    grades = {r.values.get("hardest_grade") for r in sessions} - {None}
    assert grades  # French, gym and slash grades parse


def test_a_second_session_on_a_day_is_number_2(module: Module, legacy: LegacyJournal) -> None:
    data = read_legacy(legacy.path, module)
    keys = [
        (r.values["date"], r.values["session_no"]) for r in data.records if r.entity == "session"
    ]
    assert len(set(keys)) == len(keys)
    assert any(n == 2 for _, n in keys)


def test_bad_values_and_dateless_rows_are_reported_by_place(
    module: Module, legacy: LegacyJournal
) -> None:
    book = openpyxl.load_workbook(legacy.path)
    book["Wellbeing"]["C5"] = "molto"  # not a 1–10 score
    book["Session"]["A40"] = None
    book["Mental"]["C40"] = 7  # a value in a row without a date
    book["Rest Days"]["E40"] = 5
    book.save(legacy.path)
    data = read_legacy(legacy.path, module)
    assert data.unread == ["Wellbeing!C5: not a valid value for this column, left out"]
    assert data.skipped == [
        "Session row 40: data in the session tabs but no date",
        "Rest Days row 40: data but no date",
    ]
    assert "wb_sleep_quality" not in data.records[0].values
    assert not any("molto" in line for line in [*data.unread, *data.skipped])


def test_a_different_layout_is_refused(
    module: Module, legacy: LegacyJournal, tmp_path: Path
) -> None:
    book = openpyxl.load_workbook(legacy.path)
    book["Mental"]["C4"] = "Fiducia"
    book.save(legacy.path)
    with pytest.raises(LegacyLayoutError, match=r"Mental!C4 is 'Fiducia', expected 'Confidence"):
        read_legacy(legacy.path, module)
    del book["Rest Days"]
    book.save(legacy.path)
    with pytest.raises(LegacyLayoutError, match="no tab 'Rest Days'"):
        read_legacy(legacy.path, module)
    (tmp_path / "not.xlsx").write_text("not a workbook")
    with pytest.raises(StorageError, match=r"cannot read not\.xlsx"):
        read_legacy(tmp_path / "not.xlsx", module)


def test_the_import_writes_a_new_workbook(
    module: Module, legacy: LegacyJournal, tmp_path: Path
) -> None:
    before = legacy.path.read_bytes()
    repo = XlsxBackend(tmp_path / "journal.xlsx")
    report = import_legacy(legacy.path, module, repo)
    assert report.counts == {"session": 30, "rest_day": 8}
    assert legacy.path.read_bytes() == before  # the legacy file is only read
    loaded = load_data(module, repo)
    assert loaded.issues == []
    assert len(loaded.entities["session"].records) == 30
    assert len(loaded.entities["rest_day"].records) == 8
    expected = [r.values for r in read_legacy(legacy.path, module).records if r.entity == "session"]
    got = loaded.entities["session"].records
    assert [g["date"] for g in got] == [e["date"] for e in expected]
    assert got[0]["wb_sleep_quality"] == expected[0].get("wb_sleep_quality")


def test_an_import_never_mixes_with_records(
    module: Module, legacy: LegacyJournal, tmp_path: Path
) -> None:
    repo = XlsxBackend(tmp_path / "journal.xlsx")
    import_legacy(legacy.path, module, repo)
    with pytest.raises(ImportRefused, match="already holds 38 records"):
        import_legacy(legacy.path, module, repo)


# ---- Parity ----


def with_computed(module: Module, legacy: LegacyJournal, wrong: str | None = None) -> None:
    """Write the values the app computes into the legacy computed cells, as a spreadsheet
    would have (optionally one wrong), so parity can run without LibreOffice."""
    from climblog.importers.legacy_journal import COMPUTED, REST_COMPUTED, REST_TAB

    data = read_legacy(legacy.path, module)
    tables: dict[str, list[dict[str, Value]]] = {}
    for record in data.records:
        tables.setdefault(record.entity, []).append(record.values)
    settings: dict[str, Value] = {k: s.default for k, s in module.settings.items()}
    derived = compute_module(module, tables, settings)
    book = openpyxl.load_workbook(legacy.path)
    sessions = [r for r in data.records if r.entity == "session"]
    for record, values in zip(sessions, derived["session"], strict=True):
        row = record.rows["Session"]
        for tab, columns in COMPUTED.items():
            for column, _, name in columns:
                book[tab][f"{column}{row}"] = values.get(name)
    rests = [r for r in data.records if r.entity == "rest_day"]
    for record, values in zip(rests, derived["rest_day"], strict=True):
        for column, _, name in REST_COMPUTED:
            book[REST_TAB][f"{column}{record.rows[REST_TAB]}"] = values.get(name)
    if wrong:
        book["Wellbeing"][wrong] = "LOW" if book["Wellbeing"][wrong].value != "LOW" else "GOOD"
    book.save(legacy.path)


def test_parity_when_the_sheets_agree(module: Module, legacy: LegacyJournal) -> None:
    with_computed(module, legacy)
    data = read_legacy(legacy.path, module)
    assert data.computed
    columns = parity(data, module)
    assert {c.field for c in columns} >= {"wb_score", "overall_status", "log_check", "weekday"}
    assert all(c.mismatches == [] for c in columns)
    assert sum(c.compared for c in columns) > 300


def test_parity_finds_a_difference(module: Module, legacy: LegacyJournal) -> None:
    with_computed(module, legacy, wrong="K6")  # the pre status of the 2nd session
    (column,) = [c for c in parity(read_legacy(legacy.path, module), module) if c.mismatches]
    assert column.field == "wb_pre_status" and len(column.mismatches) == 1
    assert column.matching == column.compared - 1


def test_a_workbook_never_computed_has_no_parity(module: Module, legacy: LegacyJournal) -> None:
    assert not read_legacy(legacy.path, module).computed


@pytest.mark.skipif(soffice() is None, reason="no working LibreOffice")
def test_parity_with_the_legacy_formulas(module: Module) -> None:
    """The legacy formulas, computed by LibreOffice, against the app (only known differences)."""
    command = soffice()
    assert command is not None
    with tempfile.TemporaryDirectory(dir=work_dir()) as folder:
        written = legacy_journal(Path(folder) / "legacy.xlsx", seed=4)
        recalculate(command, written.path)
        data = read_legacy(Path(folder) / "recalculated" / "legacy.xlsx", module)
    columns = parity(data, module)
    assert data.computed
    unexpected = [
        (c.entity, c.field, c.mismatches[:2]) for c in columns if c.mismatches and not c.known
    ]
    assert unexpected == []


# ---- The command ----


@pytest.fixture
def config(tmp_path: Path) -> Path:
    file = tmp_path / "config.toml"
    file.write_text(
        f'[modules.journal]\ndefinition = "{JOURNAL}"\nworkbook = "journal.xlsx"\n'
        '[backups]\nfolder = "backups"\n',
        encoding="utf-8",
    )
    return file


def test_import_legacy_command(module: Module, legacy: LegacyJournal, config: Path) -> None:
    args = ["import-legacy", "--journal", str(legacy.path), "--config", str(config)]
    result = runner.invoke(cli.app, args)
    assert result.exit_code == 0, result.output
    assert result.output.splitlines()[:2] == [
        f"imported into {config.parent / 'journal.xlsx'}",
        "  30 sessions, 8 rest days",
    ]
    assert "holds no computed values" in result.output
    again = runner.invoke(cli.app, args)
    assert again.exit_code == 1 and "already holds 38 records" in again.output


def test_dry_run_with_parity(module: Module, legacy: LegacyJournal, config: Path) -> None:
    with_computed(module, legacy, wrong="K6")
    args = ["import-legacy", "--journal", str(legacy.path), "--config", str(config), "--dry-run"]
    result = runner.invoke(cli.app, args)
    assert result.exit_code == 1  # a difference that isn't known
    assert "read 30 sessions, 8 rest days" in result.output
    assert "DIFFERS: session.wb_pre_status: " in result.output
    assert "  ok: session.wb_score: " in result.output
    assert not (config.parent / "journal.xlsx").exists()  # a dry run writes nothing


def test_known_differences_pass(module: Module, legacy: LegacyJournal, config: Path) -> None:
    with_computed(module, legacy)
    book = openpyxl.load_workbook(legacy.path)
    book["Rest Days"]["V5"] = 0  # the legacy sheet counts a session on the same day
    book.save(legacy.path)
    args = ["import-legacy", "--journal", str(legacy.path), "--config", str(config), "--dry-run"]
    result = runner.invoke(cli.app, args)
    assert result.exit_code == 0, result.output
    assert "known difference: the legacy sheet counts a session on the same day" in result.output


def test_command_errors(legacy: LegacyJournal, config: Path, tmp_path: Path) -> None:
    base = ["import-legacy", "--journal", str(legacy.path)]
    missing = runner.invoke(cli.app, [*base, "--config", str(tmp_path / "none.toml")])
    assert missing.exit_code == 1 and "no config file here" in missing.output
    other = runner.invoke(cli.app, [*base, "--config", str(config), "--module", "fingers"])
    assert other.exit_code == 1 and "not in the config: fingers" in other.output
    broken = tmp_path / "broken.xlsx"
    shutil.copy(legacy.path, broken)
    book = openpyxl.load_workbook(broken)
    book["Session"]["B4"] = "Luogo"
    book.save(broken)
    layout = runner.invoke(
        cli.app, ["import-legacy", "--journal", str(broken), "--config", str(config)]
    )
    assert layout.exit_code == 1 and "the layout differs" in layout.output


def test_report_lines_never_hold_note_contents(module: Module, legacy: LegacyJournal) -> None:
    book = openpyxl.load_workbook(legacy.path)
    book["Rest Days"]["E5"] = "[ESEMPIO] segreto"  # text in a score column
    book.save(legacy.path)
    data = read_legacy(legacy.path, module)
    assert data.unread == ["Rest Days!E5: not a valid value for this column, left out"]
    assert dt.date(2026, 6, 1) <= data.records[0].values["date"]  # type: ignore[operator]


def test_numbers_shown_in_date_formats_compare_as_numbers() -> None:
    from climblog.importers.legacy_journal import _as_number

    assert _as_number(dt.timedelta(days=2), 2) == 2
    assert _as_number(dt.datetime(1900, 1, 1), 1) == 1  # serial 1, shown as a date
    assert _as_number(dt.time(0, 0), 0) == 0
    assert _as_number("Mon", "Mon") == "Mon"
    assert _as_number(dt.timedelta(days=1), None) == dt.timedelta(days=1)


def test_a_missing_setting_keeps_the_default(module: Module, legacy: LegacyJournal) -> None:
    book = openpyxl.load_workbook(legacy.path)
    del book.defined_names["TrendDelta"]
    book.save(legacy.path)
    assert "trend_delta" not in read_legacy(legacy.path, module).settings
