# SPDX-License-Identifier: GPL-3.0-or-later
"""Formula parity (03 §2.2b): the sheet formulas, recalculated by LibreOffice, must give what the
app computes, on fictional workbooks generated from a seed (D-034).

LibreOffice is found through `CLIMBLOG_SOFFICE` (a command, e.g. `snap run libreoffice`) or
`soffice` on the PATH, and used only if it can open and recalculate an .xlsx; otherwise the
parity tests are skipped. A snap LibreOffice can't see /tmp: set `CLIMBLOG_SOFFICE_TMP` to a
folder in your home (the repo's git-ignored `build/` works).
"""

import datetime as dt
import functools
import os
import random
import shlex
import shutil
import subprocess
import tempfile
from collections.abc import Mapping
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from climblog.expr import Value
from climblog.modules import Module
from climblog.modules.compute import compute_module
from climblog.modules.formulas import XLSX, compile_formulas, formula_columns
from climblog.modules.schema import (
    BaseField,
    BoolField,
    ChoiceField,
    DateField,
    MultichoiceField,
    NumberField,
    NumericScale,
    OrdinalScale,
    ScaleField,
    TextField,
)
from climblog.storage import XlsxBackend, load_data, sync_module

WORDS = ["[ESEMPIO]", "bene", "stanco", "forte", "presa", "giù", "così così"]


@functools.cache
def soffice() -> list[str] | None:
    """The LibreOffice command, if one works here."""
    command = os.environ.get("CLIMBLOG_SOFFICE")
    words = shlex.split(command) if command else (["soffice"] if shutil.which("soffice") else [])
    if not words:
        return None
    with tempfile.TemporaryDirectory(dir=work_dir()) as folder:
        book = openpyxl.Workbook()
        sheet = book.active
        assert isinstance(sheet, Worksheet)
        sheet["A1"], sheet["A2"] = 2, "=A1*3"
        book.save(Path(folder) / "probe.xlsx")
        try:
            values = recalculate(words, Path(folder) / "probe.xlsx")
        except (OSError, subprocess.SubprocessError, KeyError, ValueError):
            return None
    return words if values["Sheet"][1][0] == 6 else None


def work_dir() -> str | None:
    return os.environ.get("CLIMBLOG_SOFFICE_TMP") or None


def recalculate(command: list[str], path: Path) -> dict[str, list[tuple[Any, ...]]]:
    """Open the workbook in LibreOffice, recalculate, and read back every cell's value."""
    out = path.parent / "recalculated"
    profile = path.parent / "profile"  # a fresh profile: no dialogs, no locks with a running one
    subprocess.run(
        [
            *command,
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--calc",
            "--convert-to",
            "xlsx:Calc MS Excel 2007 XML",
            "--outdir",
            str(out),
            str(path),
        ],
        check=True,
        capture_output=True,
        timeout=180,
    )
    book = openpyxl.load_workbook(out / path.name, data_only=True)
    return {s.title: list(s.iter_rows(values_only=True)) for s in book.worksheets}


# ---- Fictional data ----


@dataclass
class Generated:
    path: Path
    rows: dict[str, list[dict[str, Value]]]  # entity → stored values per record


def generate(module: Module, folder: Path, seed: int, records: int = 40) -> Generated:
    """A workbook made by schema sync, with formula columns and random fictional rows."""
    path = folder / "parity.xlsx"
    sync_module(module, XlsxBackend(path))
    rng = random.Random(seed)
    rows = _records(module, rng, records)
    book = openpyxl.load_workbook(path)
    columns = _add_formula_columns(module, book)
    for name, entity in module.entities.items():
        for tab, names in entity.tab_fields().items():
            sheet = book[tab]
            assert isinstance(sheet, Worksheet)
            keys = set(entity.key or [])
            row = 3
            for record in rows[name]:
                values = {k: record.get(k) for k in names}
                section = entity.is_split and tab != entity.main_tab
                empty = all(values[k] is None for k in names if k not in keys)
                if section and (rng.random() < 0.2 or empty):
                    continue  # a section not filled in
                for key, value in values.items():
                    if key in columns[tab]:
                        sheet.cell(row, columns[tab][key]).value = _cell(value)
                row += 1
    formulas = compile_formulas(module, columns, XLSX)
    for (tab, key), template in formulas.items():
        sheet = book[tab]
        assert isinstance(sheet, Worksheet)
        for row in range(3, records + 3):
            sheet.cell(row, columns[tab][key]).value = template.at(row)
    book.save(path)
    return Generated(path, rows)


def _add_formula_columns(module: Module, book: Any) -> dict[str, dict[str, int]]:
    """Append the `auto:` columns at the right of each tab (schema sync does it from #19)."""
    columns: dict[str, dict[str, int]] = {}
    for sheet in book.worksheets:
        columns[sheet.title] = {
            str(c.value): c.column for c in sheet[2] if isinstance(c.value, str) and c.value
        }
    placed, _ = formula_columns(module)
    for column in placed:
        sheet = book[column.tab]
        index = sheet.max_column + 1
        sheet.cell(1, index).value = column.label
        sheet.cell(2, index).value = column.key
        columns[column.tab][column.key] = index
    return columns


def _records(module: Module, rng: random.Random, count: int) -> dict[str, list[dict[str, Value]]]:
    rows: dict[str, list[dict[str, Value]]] = {}
    parents = [n for n, e in module.entities.items() if e.parent is None]
    children = [n for n, e in module.entities.items() if e.parent is not None]
    start = dt.date(2026, 9, 1)
    for name in parents:
        entity = module.entities[name]
        records: list[dict[str, Value]] = []
        day = start
        while len(records) < count:
            day += dt.timedelta(days=rng.choice([0, 1, 1, 2, 3]))
            record = {k: _value(module, f, rng) for k, f in _fields(module, name).items()}
            record[entity.date_field] = day
            for part in (entity.key or [])[1:]:
                if isinstance(entity.fields[part], NumberField):
                    record[part] = rng.choice([None, None, 1, 2])  # a second session now and then
            records.append(record)
        rows[name] = _unique(module, name, records)
    for name in children:
        entity = module.entities[name]
        assert entity.parent is not None
        parent_rows = [
            r
            for r in rows[entity.parent.entity]
            if all(r.get(k) == v for k, v in entity.parent.fixed.items())
        ]
        records = []
        parent_rows = _unique(module, entity.parent.entity, parent_rows)
        for parent in parent_rows:
            for n in range(rng.randint(0, 3)):
                record = {k: _value(module, f, rng) for k, f in _fields(module, name).items()}
                for k in entity.parent.key:
                    record[k] = parent[k]
                if "set_no" in record:
                    record["set_no"] = n + 1
                records.append(record)
        rows[name] = records[:count]
    return rows


def _unique(module: Module, name: str, records: list[dict[str, Value]]) -> list[dict[str, Value]]:
    """Records with distinct keys (the app would report a duplicate and skip it)."""
    entity = module.entities[name]
    key = entity.key or []
    seen, unique = set(), []
    for record in records:
        k = tuple(_read_as(entity.fields[p], record.get(p)) for p in key)
        if k not in seen:
            seen.add(k)
            unique.append(record)
    return unique


def _read_as(field: object, value: Value) -> Value:
    """A key part as the app reads it: blank = its default."""
    default = getattr(field, "default", None)
    return default if value is None else value


def _fields(module: Module, name: str) -> dict[str, BaseField]:
    return {k: f for k, f in module.entities[name].fields.items() if isinstance(f, BaseField)}


def _value(module: Module, field: BaseField, rng: random.Random) -> Value:
    if not field.required and rng.random() < 0.25:
        return None
    match field:
        case DateField(type="time"):
            return dt.time(rng.randint(6, 21), rng.choice([0, 15, 30, 45]))
        case DateField():
            return dt.date(2026, 9, 1) + dt.timedelta(days=rng.randint(0, 60))
        case ScaleField():
            scale = module.scales[field.scale]
            if isinstance(scale, OrdinalScale):
                return rng.choice(scale.choices)
            assert isinstance(scale, NumericScale)
            return rng.randint(int(scale.min), int(scale.max))
        case ChoiceField():
            choices = field.choices
            if choices is None:
                scale = module.scales[field.scale or ""]
                assert isinstance(scale, OrdinalScale)
                choices = scale.choices
            return rng.choice(choices)
        case BoolField():
            return rng.choice([True, False])
        case NumberField():
            low = field.min if field.min is not None else -10
            high = field.max if field.max is not None else 100
            if field.type == "integer":
                return rng.randint(int(low), int(high))
            return round(rng.uniform(low, high) * 2) / 2
        case MultichoiceField():
            return "; ".join(rng.sample(field.choices, k=1))
        case TextField():
            return " ".join(rng.sample(WORDS, k=2))
    if getattr(field, "type", None) == "grade":
        return rng.choice(["6a", "6b+", "7a", "5"])
    if getattr(field, "type", None) == "duration":
        return rng.choice([45, 60, 90, 120])
    return None


def _cell(value: Value) -> Any:
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, dt.date) and not isinstance(value, dt.datetime):
        return dt.datetime.combine(value, dt.time())
    return value


# ---- Comparing ----


def mismatches(module: Module, generated: Generated, command: list[str]) -> list[str]:
    """Every formula cell whose recalculated value differs from the app's."""
    repo = XlsxBackend(generated.path)
    data = load_data(module, repo)
    tables = {name: e.records for name, e in data.entities.items()}
    derived = compute_module(module, tables, data.settings)
    sheets = recalculate(command, generated.path)
    placed, _ = formula_columns(module)
    found = []
    for name, entity_data in data.entities.items():
        for column in (c for c in placed if c.key != "auto:_key" and not c.hand_written):
            key = column.key.removeprefix("auto:")
            if key not in module.entities[name].derived:
                continue
            rows = sheets[column.tab]
            index = rows[1].index(column.key)
            for record, place, values in zip(
                entity_data.records, entity_data.places, derived[name], strict=True
            ):
                if column.tab not in place:
                    continue
                got = rows[place[column.tab] - 1][index]
                want = values.get(key)
                if not same(want, got):
                    found.append(
                        f"{column.tab}!{column.key} row {place[column.tab]}: "
                        f"app {want!r}, sheet {got!r} ({_short(record)})"
                    )
    return found


def same(app: Value, sheet: Any) -> bool:
    if app is None or app == "":
        return sheet is None or sheet == ""
    if isinstance(app, bool) or isinstance(sheet, bool):
        return bool(app) == bool(sheet) and isinstance(sheet, bool)
    if isinstance(app, int | float) and isinstance(sheet, int | float):
        return abs(app - sheet) < 1e-9
    if isinstance(app, dt.date) and isinstance(sheet, dt.datetime):  # a day is its midnight
        return dt.datetime.combine(app, dt.time()) == sheet
    if isinstance(app, dt.time) and isinstance(sheet, dt.datetime | dt.timedelta):
        sheet = (
            (dt.datetime.min + sheet).time() if isinstance(sheet, dt.timedelta) else sheet.time()
        )
    return bool(app == sheet)


def _short(record: Mapping[str, Value]) -> str:
    return ", ".join(f"{k}={v!r}" for k, v in record.items() if v is not None)
