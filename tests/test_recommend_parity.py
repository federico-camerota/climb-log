# SPDX-License-Identifier: GPL-3.0-or-later
"""The recommendation's sheet formulas give what the app recommends (D-071, D-076),
recalculated by LibreOffice: on fictional journals (the status a formula column) and on a
workbook designed to reach every rule; skipped without a working LibreOffice (`parity.py`)."""

import datetime as dt
import os
import random
import tempfile
from collections.abc import Sequence
from pathlib import Path

import openpyxl
import pytest
from conftest import LoadYaml
from openpyxl.utils import get_column_letter
from parity import generate, recalculate, soffice, work_dir

from climblog.expr import Value
from climblog.modules import Module, load_module
from climblog.modules.checks import recommendation_plan
from climblog.modules.compute import compute_module
from climblog.modules.formulas import default_columns, recommendation_formulas
from climblog.modules.recommend import recommend
from climblog.storage import XlsxBackend, load_data, sync_module
from climblog.storage.drift import same

JOURNAL = load_module(Path(__file__).parent.parent / "modules" / "journal.yaml")
REQUIRED = bool(os.environ.get("CLIMBLOG_REQUIRE_SOFFICE"))
pytestmark = pytest.mark.skipif(soffice() is None and not REQUIRED, reason="no working LibreOffice")


def command() -> list[str]:
    found = soffice()
    assert found is not None, "CLIMBLOG_REQUIRE_SOFFICE is set but no LibreOffice works"
    return found


def mismatches(module: Module, path: Path, days: Sequence[dt.date]) -> tuple[list[str], set[str]]:
    """A `Check` tab with one row per day (the day, each value, headline, text) is added to
    the workbook and recalculated; every cell must equal the app's. Also the texts' headlines
    (or messages), to tell which rules the days reached."""
    plan = recommendation_plan(module)
    assert plan is not None
    data = load_data(module, XlsxBackend(path))
    tables = {name: e.records for name, e in data.entities.items()}
    derived = compute_module(module, tables, data.settings)
    full = {
        name: [{**row, **values} for row, values in zip(rows, derived[name], strict=True)]
        for name, rows in tables.items()
    }
    names = ["date", *(step.key for step in plan.values.steps), "headline", "text"]
    book = openpyxl.load_workbook(path)
    sheet = book.create_sheet("Check")
    sheet.append(names)
    for r, day in enumerate(days, start=2):
        refs = {n: f"{get_column_letter(i)}{r}" for i, n in enumerate(names, start=1)}
        formulas = recommendation_formulas(module, default_columns(module), refs)
        sheet.cell(r, 1).value = day
        for i, name in enumerate(names[1:], start=2):
            sheet.cell(r, i).value = formulas[name]
    book.save(path)

    rows = recalculate(command(), path)["Check"]
    wrong, reached = [], set()
    for day, row in zip(days, rows[1:], strict=True):
        advice = recommend(module, plan, full, data.settings, day)
        reached.add(advice.text.split(" — ")[0])
        want: dict[str, Value] = {**advice.values, "headline": advice.headline}
        want["text"] = advice.text
        for i, name in enumerate(names[1:], start=1):
            got = row[i]
            if name == "last_session" and isinstance(got, int | float) and got:
                got = (dt.datetime(1899, 12, 30) + dt.timedelta(days=got)).date()
            if not same(want[name], got):
                wrong.append(f"{day} {name}: app {want[name]!r}, sheet {got!r}")
    return wrong, reached


@pytest.mark.parametrize("seed", [1, 2, 3])
def test_the_journal_sheet_recommends_as_the_app(seed: int) -> None:
    with tempfile.TemporaryDirectory(dir=work_dir()) as folder:
        generated = generate(JOURNAL, Path(folder), seed)
        dates = [r.get("date") for r in generated.rows["session"]]
        days = {d for d in dates if isinstance(d, dt.date)}
        days |= {min(days) - dt.timedelta(days=1), max(days) + dt.timedelta(days=30)}
        for d in [d for d in dates if isinstance(d, dt.date)]:
            days |= {d + dt.timedelta(days=1), d + dt.timedelta(days=15)}
        wrong, _ = mismatches(JOURNAL, generated.path, sorted(days))
        assert wrong == []


# A status typed by hand, so the rows can reach every rule of the journal's recommendation.
DESIGNED = """\
module: test_recommend_sheet
name: Test recommend sheet
version: 1
kind: journal
entities:
  session:
    label: Session
    tab: Sessions
    key: [date, session_no]
    fields:
      date: {type: date, required: true}
      session_no: {type: integer, min: 1, default: 1}
      status: {type: choice, choices: [GOOD, WATCH, LOW]}
recommendation:
  values:
    last_session: "last(session.date, until=date)"
    latest_status: "last(session.status, until=date)"
    low_sessions: "count(window(session.status, days=14, until=date), equals='LOW')"
  rules:
    - {if: "not notnull(last_session)", message: Log a session to get guidance}
    - {if: "latest_status == 'LOW' or low_sessions >= 2", headline: RECOVER, message: Rest.}
    - {if: "latest_status == 'WATCH'", headline: MODERATE, message: Cap it.}
    - {if: "latest_status == 'GOOD'", headline: GREEN LIGHT, message: Go.}
    - {message: Fill in the remaining tabs}
"""


def oct(day: int) -> dt.datetime:
    return dt.datetime(2026, 10, day)


SESSIONS: list[tuple[dt.datetime, int | None, str | None, str | None]] = [
    (oct(1), None, "GOOD", None),
    (oct(3), None, "WATCH", None),
    (oct(5), None, None, None),  # no status yet
    (oct(7), None, "LOW", None),
    (oct(9), None, "GOOD", None),
    (oct(10), 1, "LOW", None),
    (oct(10), 2, "GOOD", None),  # the day's second session is the latest
    (oct(29), None, "LOW", "2026-10-29 21:00"),  # soft deleted: not counted
    (oct(30), None, "GOOD", None),
]


def test_every_rule_in_the_sheet_as_in_the_app(load: LoadYaml) -> None:
    module = load(DESIGNED)
    with tempfile.TemporaryDirectory(dir=work_dir()) as folder:
        path = Path(folder) / "designed.xlsx"
        sync_module(module, XlsxBackend(path))
        book = openpyxl.load_workbook(path)
        sheet = book["Sessions"]
        columns = {str(c.value): c.column for c in sheet[2] if c.value}
        for r, (day, number, status, deleted) in enumerate(SESSIONS, start=3):
            for key, value in [
                ("date", day),
                ("session_no", number),
                ("status", status),
                ("deleted", deleted),
            ]:
                sheet.cell(r, columns[key]).value = value
        book.save(path)
        days = [oct(d).date() for d in (1, 2, 3, 5, 7, 9, 10, 21, 22, 29, 30)]
        days.insert(0, dt.date(2026, 9, 30))
        wrong, reached = mismatches(module, path, days)
    assert wrong == []
    assert reached == {
        "Log a session to get guidance",
        "RECOVER",
        "MODERATE",
        "GREEN LIGHT",
        "Fill in the remaining tabs",
    }


# The legacy Summary sheet, Recent part (00-current-state §1): latest overall status (B6), LOW
# sessions in the last 14 days (B8) and the suggested next session (B9), with TODAY() as the
# day in column A. Its data rows start at row 5; one session a day (it matches the first row).
LEGACY_B6 = (
    "=IF(COUNT(Session!$A$5:$A$304)=0,\"\",INDEX('Session Feel'!$K$5:$K$304,"
    "MATCH(MAX(Session!$A$5:$A$304),Session!$A$5:$A$304,0)))"
)
LEGACY_B8 = '=COUNTIFS(\'Session Feel\'!$K$5:$K$304,"LOW",Session!$A$5:$A$304,">="&$A{r}-14)'
LEGACY_B9 = (
    '=IF(COUNT(Session!$A$5:$A$304)=0,"Log a session to get guidance",'
    'IF(OR(B{r}="LOW",C{r}>=2),"RECOVER — rest day or easy technical climbing. No limit '
    'attempts or max hangs until your recovery scores come back up.",IF(B{r}="WATCH",'
    '"MODERATE — train as planned but cap the intensity; re-check your pre-session wellbeing '
    'before you start.",IF(B{r}="GOOD","GREEN LIGHT — good window for quality work: project '
    'attempts, limit boulders or strength.","Fill in the remaining tabs for your latest '
    'session to get guidance"))))'
)


@pytest.mark.parametrize("seed", [1, 2, 3, 4])
def test_the_app_recommends_as_the_legacy_summary(seed: int) -> None:
    rng = random.Random(seed)
    start = dt.date(2026, 9, 1)
    days = sorted(rng.sample(range(40), rng.randint(0, 12)))
    sessions = [
        (start + dt.timedelta(days=d), rng.choice(["GOOD", "WATCH", "LOW", None])) for d in days
    ]
    last = sessions[-1][0] if sessions else start
    checks = [last + dt.timedelta(days=k) for k in (0, 1, 7, 13, 14, 15, 30)]

    book = openpyxl.Workbook()
    book.remove(book.active)  # type: ignore[arg-type]
    session, feel, check = (book.create_sheet(n) for n in ("Session", "Session Feel", "Check"))
    for r, (day, status) in enumerate(sessions, start=5):
        session.cell(r, 1).value = day
        feel.cell(r, 11).value = status
    for r, day in enumerate(checks, start=1):
        check.cell(r, 1).value = day
        check.cell(r, 2).value = LEGACY_B6
        check.cell(r, 3).value = LEGACY_B8.format(r=r)
        check.cell(r, 4).value = LEGACY_B9.format(r=r)
    plan = recommendation_plan(JOURNAL)
    assert plan is not None
    rows: list[dict[str, Value]] = [
        {"date": day, "session_no": 1, "overall_status": status} for day, status in sessions
    ]
    with tempfile.TemporaryDirectory(dir=work_dir()) as folder:
        path = Path(folder) / "legacy.xlsx"
        book.save(path)
        legacy = recalculate(command(), path)["Check"]
    for day, row in zip(checks, legacy, strict=True):
        advice = recommend(JOURNAL, plan, {"session": rows}, {}, day)
        assert advice.text == row[3], (day, sessions)
