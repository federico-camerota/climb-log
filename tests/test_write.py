# SPDX-License-Identifier: GPL-3.0-or-later
"""Writing records: add, update, soft delete, across split tabs and children (04 §1, D-049)."""

import datetime as dt
from collections.abc import Sequence
from pathlib import Path

import openpyxl
import pytest
from conftest import LoadYaml
from workbooks import tab, write_workbook

from climblog.expr import Value
from climblog.modules import Module
from climblog.storage import (
    Edit,
    RecordNotFound,
    RecordWriter,
    StorageError,
    WriteConflict,
    WriteError,
    XlsxBackend,
    load_data,
)

# A fictional module: a session split over two tabs, board sets as its children, keyless falls.
MODULE = """\
module: test_write
name: Test write
version: 1
kind: journal
scales:
  s10: {min: 1, max: 10, direction: higher_better}
entities:
  session:
    label: Session
    key: [date, session_no, kind]  # a board and a climbing session may share a day
    anchor_tab: Session
    sections:
      info: {label: Info, tab: Session}
      wb: {label: Wellbeing, tab: Wellbeing}
    fields:
      date: {type: date, required: true, section: "*"}
      session_no: {type: integer, min: 1, default: 1, section: "*"}
      kind: {type: choice, choices: [Climb, Board], section: "*"}
      warm_up: {type: bool, section: info}
      mood: {type: scale, scale: s10, section: wb}
      note: {type: text, section: wb}
  set:
    label: Set
    tab: Sets
    key: [date, session_no, set_no]
    parent: {entity: session, key: [date, session_no], fixed: {kind: Board}}
    fields:
      date: {type: date}
      session_no: {type: integer, default: 1}
      set_no: {type: integer, default: next}  # not filled in by writes yet
      kg: {type: number}
  fall:
    label: Fall
    tab: Falls
    id_prefix: fall
    fields:
      id: {type: text}
      date: {type: date}
"""

D19, D20, D21 = dt.date(2026, 10, 19), dt.date(2026, 10, 20), dt.date(2026, 10, 21)
NOW = dt.datetime(2026, 10, 24, 18, 30, 5)
SESSION = ["date", "session_no", "kind", "warm_up", "auto:status", "updated_at", "deleted"]
WELLBEING = ["date", "session_no", "kind", "mood", "note", "deleted"]  # no updated_at: optional
FALLS = ["id", "date"]
SETS = ["date", "session_no", "set_no", "kg", "updated_at", "deleted"]


def at(day: dt.date) -> dt.datetime:
    return dt.datetime.combine(day, dt.time())


@pytest.fixture
def module(load: LoadYaml) -> Module:
    return load(MODULE)


@pytest.fixture
def path(tmp_path: Path) -> Path:
    """A board session on the 20th with two sets and a mood; a climbing session on the 21st."""
    return write_workbook(
        tmp_path / "journal.xlsx",
        {
            "Session": tab(SESSION, [at(D20), None, "Board", "Yes"], [at(D21), 1, "Climb"]),
            "Wellbeing": tab(WELLBEING, [at(D20), 1, "Board", 7, "Bene"]),
            "Sets": tab(SETS, [at(D20), None, 1, 10], [at(D20), None, 2, 12.5]),
            "Falls": tab(FALLS),
        },
    )


@pytest.fixture
def writer(module: Module, path: Path) -> RecordWriter:
    return RecordWriter(module, XlsxBackend(path), clock=lambda: NOW.replace(microsecond=7))


def records(module: Module, path: Path, entity: str) -> list[dict[str, Value]]:
    data = load_data(module, XlsxBackend(path))
    assert data.issues == []
    return data.entities[entity].records


def rows(path: Path, name: str) -> list[tuple[object, ...]]:
    workbook = openpyxl.load_workbook(path)
    try:
        return [tuple(c.value for c in row) for row in workbook[name].iter_rows(min_row=3)]
    finally:
        workbook.close()


BOARD20: dict[str, Value] = {
    "date": D20,
    "kind": "Board",
}  # the key of the board session (session_no blank = 1)


def test_add_writes_only_the_tabs_with_values(
    writer: RecordWriter, module: Module, path: Path
) -> None:
    writer.add("session", {"date": D19, "kind": "climb", "warm_up": False})
    assert rows(path, "Session")[-1] == (at(D19), 1, "Climb", "No", None, NOW, None)
    assert len(rows(path, "Wellbeing")) == 1  # nothing for the Wellbeing tab
    assert records(module, path, "session")[-1] == {
        "date": D19, "session_no": 1, "kind": "Climb", "warm_up": False
    }  # fmt: skip


def test_add_splits_a_record_across_its_tabs(
    writer: RecordWriter, module: Module, path: Path
) -> None:
    note = "=po' stanco"  # written as text, not as a formula
    writer.add("session", {"date": D19, "session_no": 2, "kind": "Board", "mood": 6, "note": note})
    assert rows(path, "Session")[-1] == (at(D19), 2, "Board", None, None, NOW, None)
    assert rows(path, "Wellbeing")[-1] == (at(D19), 2, "Board", 6, note, None)
    assert records(module, path, "session")[-1] == {
        "date": D19, "session_no": 2, "kind": "Board", "warm_up": None, "mood": 6, "note": note
    }  # fmt: skip


def test_add_needs_a_new_complete_key(writer: RecordWriter, path: Path) -> None:
    before = path.read_bytes()
    with pytest.raises(WriteError, match="Session already has a row for date 2026-10-20, "):
        writer.add("session", BOARD20)  # session_no blank = 1, as on read
    with pytest.raises(WriteError, match="no kind: a session needs one"):
        writer.add("session", {"date": D19})
    with pytest.raises(WriteError, match="no set_no: a set needs one"):
        writer.add("set", {"date": D20, "kg": 5})
    assert path.read_bytes() == before


def test_add_refuses_a_key_held_by_an_orphan_section_row(module: Module, tmp_path: Path) -> None:
    path = write_workbook(
        tmp_path / "orphan.xlsx",
        {
            "Session": tab(SESSION),
            "Wellbeing": tab(WELLBEING, ["oops", 1, "Board"], [at(D19), 1, "Board", 5]),
            "Sets": tab(SETS),
            "Falls": tab(FALLS),
        },
    )  # a bad key cell (a data issue, can't match), then an orphan row
    with pytest.raises(WriteError, match=r"Wellbeing already has a row for .* \(row 4\)"):
        RecordWriter(module, XlsxBackend(path)).add("session", {"date": D19, "kind": "Board"})


def test_deleted_rows_dont_hold_their_key(module: Module, tmp_path: Path) -> None:
    path = write_workbook(
        tmp_path / "deleted.xlsx",
        {
            "Session": tab(SESSION, [at(D19), 1, "Board", "No", None, None, "x"]),
            "Wellbeing": tab(WELLBEING),
            "Sets": tab(SETS),
            "Falls": tab(FALLS),
        },
    )
    RecordWriter(module, XlsxBackend(path)).add("session", {"date": D19, "kind": "Board"})
    assert [r.get("warm_up") for r in records(module, path, "session")] == [None]


def test_every_bad_value_is_reported(writer: RecordWriter) -> None:
    """All the fields at fault at once, so a form can show each (D-079)."""
    with pytest.raises(WriteError) as caught:
        writer.add("session", {"date": D19, "kind": "Yoga", "mood": 11, "readiness": 7})
    assert [field for field, _ in caught.value.problems] == ["readiness", "kind", "mood"]
    assert caught.value.field == "readiness"  # the first; unknown fields' message as before
    assert str(caught.value) == "Session has no field readiness"
    with pytest.raises(WriteError) as caught:
        writer.add("session", {"date": D19, "kind": "Yoga", "mood": 11})
    assert [field for field, _ in caught.value.problems] == ["kind", "mood"]
    assert caught.value.problems[1] == ("mood", "mood: 11 is above the maximum 10")
    assert (caught.value.field, str(caught.value)) == caught.value.problems[0]


def test_values_must_fit_the_module(writer: RecordWriter) -> None:
    with pytest.raises(WriteError, match="Session has no field mood_post, readiness"):
        writer.add("session", {"date": D19, "mood_post": 5, "readiness": 7})
    with pytest.raises(WriteError, match="mood: 11 is above the maximum 10"):
        writer.add("session", {"date": D19, "kind": "Board", "mood": 11})
    with pytest.raises(WriteError, match="no entity 'rest_day' in module test_write"):
        writer.add("rest_day", {"date": D19})
    with pytest.raises(WriteError, match="Fall has generated ids: writing it comes later"):
        writer.add("fall", {"date": D19})


def test_a_child_needs_its_parent(writer: RecordWriter, module: Module, path: Path) -> None:
    writer.add("set", {"date": D20, "set_no": 3, "kg": 15})
    assert records(module, path, "set")[-1] == {
        "date": D20, "session_no": 1, "set_no": 3, "kg": 15
    }  # fmt: skip
    with pytest.raises(WriteError, match="no Session row for date 2026-10-21, session_no 1"):
        writer.add("set", {"date": D21, "set_no": 1})  # the 21st is a Climb session, not Board


def test_update_changes_the_rows_the_record_has(writer: RecordWriter, path: Path) -> None:
    writer.update("session", BOARD20, {"warm_up": False, "mood": None, "date": D20})
    assert rows(path, "Session")[0] == (at(D20), None, "Board", "No", None, NOW, None)
    assert rows(path, "Wellbeing")[0] == (at(D20), 1, "Board", None, "Bene", None)


def test_update_adds_a_row_to_a_tab_the_record_lacks(
    writer: RecordWriter, module: Module, path: Path
) -> None:
    writer.update("session", {"date": D21, "session_no": 1, "kind": "Climb"}, {"mood": 8})
    assert rows(path, "Wellbeing")[-1] == (at(D21), 1, "Climb", 8, None, None)
    assert records(module, path, "session")[1]["mood"] == 8


def test_update_refuses_key_changes_and_unknown_records(writer: RecordWriter, path: Path) -> None:
    before = path.read_bytes()
    with pytest.raises(WriteError, match="kind can't be changed: delete the record and add it"):
        writer.update("session", BOARD20, {"kind": "Climb"})
    with pytest.raises(RecordNotFound, match="no session for date 2026-10-19, session_no 1, "):
        writer.update("session", {"date": D19, "kind": "Board"}, {"mood": 5})
    with pytest.raises(WriteError, match="mood is not part of the session key"):
        writer.update("session", {**BOARD20, "mood": 7}, {"mood": 5})
    writer.update("session", BOARD20, {})  # nothing to write
    assert path.read_bytes() == before


def test_update_of_a_row_edited_by_hand_meanwhile_is_a_conflict(module: Module, path: Path) -> None:
    class HandEdit(XlsxBackend):
        """Someone types in the sheet between the app's read and its write."""

        def apply(self, edits: Sequence[Edit]) -> list[int]:
            workbook = openpyxl.load_workbook(self.path)
            workbook["Session"]["D3"] = "No"
            workbook.save(self.path)
            return super().apply(edits)

    with pytest.raises(WriteConflict, match="tab 'Session', row 3 changed since it was read"):
        RecordWriter(module, HandEdit(path)).update("session", BOARD20, {"warm_up": True})
    assert rows(path, "Session")[0][3] == "No"


def test_delete_hides_the_record_and_its_children(
    writer: RecordWriter, module: Module, path: Path
) -> None:
    writer.delete("session", BOARD20)
    assert rows(path, "Session")[0][5:] == (NOW, NOW)
    assert [row[4:] for row in rows(path, "Sets")] == [(NOW, NOW), (NOW, NOW)]
    assert rows(path, "Wellbeing")[0] == (at(D20), 1, "Board", 7, "Bene", NOW)  # all its tabs
    data = load_data(module, XlsxBackend(path))
    assert [r["date"] for r in data.entities["session"].records] == [D21]
    assert data.entities["set"].records == []


def test_a_child_of_another_kind_is_not_deleted(
    writer: RecordWriter, module: Module, path: Path
) -> None:
    writer.add("session", {"date": D20, "kind": "Climb"})
    writer.delete("session", {"date": D20, "kind": "Climb"})  # the Board sets aren't its children
    assert len(records(module, path, "set")) == 2
    writer.delete("set", {"date": D20, "set_no": 2})
    assert [r["set_no"] for r in records(module, path, "set")] == [1]


def test_delete_needs_a_deleted_column(module: Module, tmp_path: Path) -> None:
    path = write_workbook(
        tmp_path / "old.xlsx",
        {
            "Session": tab(["date", "session_no", "kind"], [at(D19), None, "Board"]),  # older
            "Wellbeing": tab(WELLBEING),
            "Sets": tab(SETS),
            "Falls": tab(FALLS),
        },
    )
    with pytest.raises(StorageError, match="tab 'Session' has no column deleted in row 2"):
        RecordWriter(module, XlsxBackend(path)).delete("session", {"date": D19, "kind": "Board"})
