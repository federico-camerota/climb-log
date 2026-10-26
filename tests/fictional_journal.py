# SPDX-License-Identifier: GPL-3.0-or-later
# ruff: noqa: RUF001  (the legacy texts and choices use en dashes, as in "Yes - fully")
"""Fictional climbing-journal workbooks for the tests (D-034, D-060): every value is invented.

- `legacy_journal`: the layout of the legacy `Climbing_Journal.xlsx` (00-current-state.md §1):
  its tabs, header rows 1-4, and its formulas in rows 5-304 (copied from the real workbook's
  structure, no data), with the date typed on `Session` only and the other session tabs linked
  to it by row. Data rows come from a seed, with the cases hand entry produces.
- `new_journal`: the new layout (two header rows, D-014) with the hand-entry cases of 04 §1:
  blank session numbers, reordered columns, decimal commas, notes starting with `'` or `=`,
  choices typed in another case, orphan section rows, a blank key with data, a duplicate key.

`python tests/fictional_journal.py FOLDER` (`make fixtures`) writes one of each, to open them.
"""

import datetime as dt
import random
import sys
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

FIRST, LAST = 5, 304  # legacy data rows
NOTES = [
    "[ESEMPIO] Un po' stanco dal lavoro",
    "[ESEMPIO] Dormito benissimo, perché no?",
    "[ESEMPIO] Più riposo tra i tentativi",
    "'po' di dolore al dito",  # starts with an apostrophe: kept as typed (04 §3b)
    "=non una formula",  # starts with "=": text, not a formula
    "[ESEMPIO] così così",
]
VENUES = ["[ESEMPIO] Palestra Nord", "[ESEMPIO] Falesia del Lago", "Home wall"]
GRADES = ["6a", "6b+", "6C", "7a/7a+", "7a+", "5", "6", None]  # French, gym, slash grades
SETTINGS = {"LowMax": 4, "WatchMax": 6, "AttnMin": 8, "NoteMin": 6, "TrendDelta": 0.5}

WEEKDAYS = 'CHOOSE(WEEKDAY(A{r},2),"Mon","Tue","Wed","Thu","Fri","Sat","Sun")'
LINK = '=IF(Session!$A{r}="","",Session!$A{r})'
_MISSING = (
    'IF(Wellbeing!M{r}="","Wellbeing, ","")&IF(COUNT(Physical!B{r}:F{r},Physical!H{r}:L{r})<10,'
    '"Physical, ","")&IF(Mental!M{r}="","Mental, ","")&IF(\'Session Feel\'!H{r}="","Feel, ","")'
    '&IF(\'Wins & Reflect\'!M{r}="","Reflect, ","")'
)
_MENTAL_MIN = (
    'MIN(IF(C{r}="",10,C{r}),IF(D{r}="",10,D{r}),IF(G{r}="",10,G{r}),IF(H{r}="",10,H{r}),'
    'IF(J{r}="",10,J{r}),IF(K{r}="",10,K{r}),IF(B{r}="",10,11-B{r}),IF(E{r}="",10,11-E{r}),'
    'IF(I{r}="",10,11-I{r}))'
)
_REST_MIN = (
    'MIN(IF(E{r}="",10,E{r}),IF(F{r}="",10,F{r}),IF(G{r}="",10,G{r}),IF(H{r}="",10,H{r}),'
    'IF(I{r}="",10,11-I{r}))'
)


@dataclass(frozen=True)
class LegacyTab:
    name: str
    title: str
    intro: str
    groups: dict[str, str]  # row 3: column → group label
    headers: list[str]  # row 4, from column A
    formulas: dict[str, str]  # column → template ({r} = the row)


LEGACY: list[LegacyTab] = [
    LegacyTab(
        "Session",
        "Session Log",
        "Start here for every session · type the date ONCE — it flows to every other tab so "
        "rows always line up · yellow = you type, grey = automatic · use the filter arrows to "
        "look back",
        {"A": "SESSION DETAILS", "I": "AUTOMATIC"},
        [
            "Date",
            "Venue",
            "Indoor / outdoor",
            "Session type",
            "Partners",
            "Duration\n(min)",
            "Hardest grade\ntried",
            "Notes",
            "Overall status\n(all tabs)",
            "Log check",
            "Weekday",
            "Month",
            "Days since\nprevious",
        ],
        {
            "I": '=IF(A{r}="","",\'Session Feel\'!K{r})',
            "J": f'=IF(A{{r}}="","",IF(LEN({_MISSING})=0,"Complete","Missing: "&LEFT({_MISSING},'
            f"LEN({_MISSING})-2)))",
            "K": f'=IF(A{{r}}="","",{WEEKDAYS})',
            "L": '=IF(A{r}="","",YEAR(A{r})&"-"&RIGHT("0"&MONTH(A{r}),2))',
            "M": '=IF(A{r}="","",IF(COUNT(A$4:A{p})=0,"",A{r}-MAX(A$4:A{p})))',
        },
    ),
    LegacyTab(
        "Wellbeing",
        "Wellbeing",
        "Sleep, mood and energy before and after each session · scores 1–10, 10 = best "
        "(except ↓ columns, where lower is better)",
        {"A": "PRE-SESSION", "G": "POST-SESSION", "K": "AUTOMATIC"},
        [
            "Date",
            "Sleep\n(hours)",
            "Sleep quality\n(1–10)",
            "Mood\n(1–10)",
            "Energy\n(1–10)",
            "Pre-session note",
            "Post fatigue\n(1–10 ↓)",
            "Post mood\n(1–10)",
            "Recovery feel\n(1–10)",
            "Post-session note",
            "Pre\nstatus",
            "Post\nstatus",
            "Wellbeing\nscore",
            "Wellbeing\nstatus",
        ],
        {
            "A": LINK,
            "K": '=IF(COUNT(C{r}:E{r})<3,"",IF(MIN(C{r}:E{r})<=LowMax,"LOW",'
            'IF(MIN(C{r}:E{r})<=WatchMax,"WATCH","GOOD")))',
            "L": '=IF(COUNT(G{r}:I{r})<3,"",IF(MIN(11-G{r},H{r},I{r})<=LowMax,"LOW",'
            'IF(MIN(11-G{r},H{r},I{r})<=WatchMax,"WATCH","GOOD")))',
            "M": '=IF(COUNT(C{r}:E{r},G{r}:I{r})<6,"",'
            "ROUND(AVERAGE(C{r},D{r},E{r},11-G{r},H{r},I{r}),1))",
            "N": '=IF(OR(K{r}="LOW",L{r}="LOW"),"LOW",IF(OR(K{r}="WATCH",L{r}="WATCH"),"WATCH",'
            'IF(OR(K{r}="GOOD",L{r}="GOOD"),"GOOD","")))',
        },
    ),
    LegacyTab(
        "Physical",
        "Physical",
        "Soreness by body part before and after each session · 1 = fresh, 10 = very sore "
        "(↓ lower is better) · status looks at your worst body part before climbing",
        {"A": "PRE-SESSION", "H": "POST-SESSION", "N": "AUTOMATIC"},
        [
            "Date",
            *(f"{p}\n(1–10 ↓)" for p in ("Fingers", "Forearms", "Shoulders", "Wrists", "Skin")),
            "Pre-session note",
            *(f"{p}\n(1–10 ↓)" for p in ("Fingers", "Forearms", "Shoulders", "Wrists", "Skin")),
            "Post-session note",
            "Pre-session\nreadiness",
            "Physical\nstatus",
        ],
        {
            "A": LINK,
            "N": '=IF(COUNT(B{r}:F{r})=0,"",ROUND(11-AVERAGE(B{r}:F{r}),1))',
            "O": '=IF(COUNT(B{r}:F{r})=0,"",IF(MAX(B{r}:F{r})>=11-LowMax,"LOW",'
            'IF(MAX(B{r}:F{r})>=11-WatchMax,"WATCH","GOOD")))',
        },
    ),
    LegacyTab(
        "Mental",
        "Mental",
        "Confidence, focus and fear response · 1–10, 10 = best (↓ columns: lower is better) · "
        "status flags your weakest mental item",
        {"A": "PRE-SESSION", "G": "DURING / POST-SESSION", "M": "AUTOMATIC"},
        [
            "Date",
            "Life stress\n(1–10 ↓)",
            "Confidence\n(1–10)",
            "Motivation\n(1–10)",
            "Anxiety\n(1–10 ↓)",
            "Pre-session note",
            "Focus during\n(1–10)",
            "Commitment\nto falls (1–10)",
            "Fear response\n(1–10 ↓)",
            "Confidence on\nroutes (1–10)",
            "Approach\nmindset (1–10)",
            "Post-session note",
            "Mental\nscore",
            "Mental\nstatus",
        ],
        {
            "A": LINK,
            "M": '=IF(COUNT(B{r}:E{r},G{r}:K{r})<9,"",'
            "ROUND(AVERAGE(11-B{r},C{r},D{r},11-E{r},G{r},H{r},11-I{r},J{r},K{r}),1))",
            "N": f'=IF(COUNT(B{{r}}:E{{r}},G{{r}}:K{{r}})=0,"",IF({_MENTAL_MIN}<=LowMax,"LOW",'
            f'IF({_MENTAL_MIN}<=WatchMax,"WATCH","GOOD")))',
        },
    ),
    LegacyTab(
        "Session Feel",
        "Session Feel",
        "Overall outcome, plan adherence and satisfaction · the Overall columns combine all "
        "four classes (worst status wins)",
        {"A": "SESSION OUTCOME", "H": "AUTOMATIC"},
        [
            "Date",
            "Session outcome",
            "Stuck to plan?",
            "Enjoyment\n(1–10)",
            "Satisfaction\n(1–10)",
            "Effort / energy\nused (1–10)",
            "Session note",
            "Feel\nscore",
            "Feel\nstatus",
            "Overall\nscore",
            "Overall status\n(all tabs)",
        ],
        {
            "A": LINK,
            "H": '=IF(COUNT(D{r}:E{r})<2,"",ROUND(AVERAGE(D{r}:E{r}),1))',
            "I": '=IF(COUNT(D{r}:E{r})=0,"",IF(MIN(D{r}:E{r})<=LowMax,"LOW",'
            'IF(MIN(D{r}:E{r})<=WatchMax,"WATCH","GOOD")))',
            "J": '=IF(COUNT(Wellbeing!M{r},Physical!N{r},Mental!M{r},H{r})=0,"",'
            "ROUND(AVERAGE(Wellbeing!M{r},Physical!N{r},Mental!M{r},H{r}),1))",
            "K": '=IF(OR(Wellbeing!N{r}="LOW",Physical!O{r}="LOW",Mental!N{r}="LOW",I{r}="LOW"),'
            '"LOW",IF(OR(Wellbeing!N{r}="WATCH",Physical!O{r}="WATCH",Mental!N{r}="WATCH",'
            'I{r}="WATCH"),"WATCH",IF(OR(Wellbeing!N{r}="GOOD",Physical!O{r}="GOOD",'
            'Mental!N{r}="GOOD",I{r}="GOOD"),"GOOD","")))',
        },
    ),
    LegacyTab(
        "Wins & Reflect",
        "Wins & Reflect",
        "The four journal prompts from the video, each scored 1–10 · the Summary surfaces the "
        "strongest and most urgent notes from the last 2 weeks",
        {"A": "JOURNAL PROMPTS — SCORE EACH ONE", "J": "BANK A WIN", "M": "AUTO"},
        [
            "Date",
            "1 · What felt good",
            "Score\n(1–10)",
            "2 · What felt off",
            "Severity\n(1–10 ↓)",
            "3 · What surprised you",
            "Impact\n(1–10)",
            "4 · What would you change next time",
            "Priority\n(1–10)",
            "Overall session\nscore (1–10)",
            "Win size",
            "Win to bank",
            "Reflection\nstatus",
        ],
        {
            "A": LINK,
            "M": '=IF(COUNT(E{r},I{r})=0,"",IF(OR(E{r}>=AttnMin,I{r}>=AttnMin),"NEEDS ATTENTION",'
            'IF(OR(E{r}>=NoteMin,I{r}>=NoteMin),"WORTH NOTING","GOOD")))',
        },
    ),
    LegacyTab(
        "Rest Days",
        "Rest Days",
        "How you feel on days off · one row per rest day · scores 1–10, 10 = best (↓ columns: "
        "lower is better) · yellow = you type, grey = automatic",
        {"A": "REST DAY", "D": "HOW YOU FEEL", "J": "BODY", "P": "AUTOMATIC"},
        [
            "Date",
            "Rest type",
            "Activity",
            "Sleep\n(hours)",
            "Sleep quality\n(1–10)",
            "Mood\n(1–10)",
            "Energy\n(1–10)",
            "Motivation to\nclimb (1–10)",
            "Life stress\n(1–10 ↓)",
            *(f"{p}\n(1–10 ↓)" for p in ("Fingers", "Forearms", "Shoulders", "Wrists", "Skin")),
            "Note",
            "Wellbeing\nscore",
            "Body\nreadiness",
            "Wellbeing\nstatus",
            "Body\nstatus",
            "Rest-day\nstatus",
            "Weekday",
            "Days since\nlast climb",
        ],
        {
            "P": '=IF(COUNT(E{r}:I{r})<5,"",ROUND(AVERAGE(E{r},F{r},G{r},H{r},11-I{r}),1))',
            "Q": '=IF(COUNT(J{r}:N{r})=0,"",ROUND(11-AVERAGE(J{r}:N{r}),1))',
            "R": f'=IF(COUNT(E{{r}}:I{{r}})=0,"",IF({_REST_MIN}<=LowMax,"LOW",'
            f'IF({_REST_MIN}<=WatchMax,"WATCH","GOOD")))',
            "S": '=IF(COUNT(J{r}:N{r})=0,"",IF(MAX(J{r}:N{r})>=11-LowMax,"LOW",'
            'IF(MAX(J{r}:N{r})>=11-WatchMax,"WATCH","GOOD")))',
            "T": '=IF(OR(R{r}="LOW",S{r}="LOW"),"LOW",IF(OR(R{r}="WATCH",S{r}="WATCH"),"WATCH",'
            'IF(OR(R{r}="GOOD",S{r}="GOOD"),"GOOD","")))',
            "U": f'=IF(A{{r}}="","",{WEEKDAYS})',
            # The real file has MAXIFS without the `_xlfn.` an .xlsx needs for it (LibreOffice
            # shows #NAME? then); the fixture writes it, so its values can be computed.
            "V": '=IF(A{r}="","",IF(COUNTIFS(Session!$A$5:$A1000,"<="&A{r})=0,"",'
            'A{r}-_xlfn.MAXIFS(Session!$A$5:$A1000,Session!$A$5:$A1000,"<="&A{r})))',
        },
    ),
]

SESSION_TYPES = ["Climbing", "Hangboard", "Strength", "Endurance / ARC", "Active rest"]
SETTINGS_ROWS = range(46, 51)  # 'Start Here'!F46:F50, the named settings


@dataclass
class LegacyJournal:
    """What `legacy_journal` wrote, to compare with what an importer reads."""

    path: Path
    sessions: list[dict[str, Any]] = field(default_factory=list)  # legacy column → value, per row
    rest_days: list[dict[str, Any]] = field(default_factory=list)


def legacy_journal(path: Path, seed: int, sessions: int = 30, rest_days: int = 8) -> LegacyJournal:
    """A fictional workbook in the legacy layout, with its formulas and seeded data rows."""
    rng = random.Random(seed)
    book = openpyxl.Workbook()
    start = book.active
    assert isinstance(start, Worksheet)
    start.title = "Start Here"
    start["A1"] = "[ESEMPIO] Climbing journal: start here"
    for row, (name, value) in zip(SETTINGS_ROWS, SETTINGS.items(), strict=True):
        start.cell(row, 5).value = name
        start.cell(row, 6).value = value
        book.defined_names[name] = DefinedName(name, attr_text=f"'Start Here'!$F${row}")
    sheets = {}
    for spec in LEGACY:
        sheet = book.create_sheet(spec.name)
        sheet["A1"], sheet["A2"] = spec.title, spec.intro
        for column, label in spec.groups.items():
            sheet[f"{column}3"] = label
        for i, header in enumerate(spec.headers, start=1):
            sheet.cell(4, i).value = header
        for row in range(FIRST, LAST + 1):
            for column, template in spec.formulas.items():
                sheet[f"{column}{row}"] = template.format(r=row, p=row - 1)
        sheet.freeze_panes = "A5" if spec.name == "Rest Days" else "B5"
        sheets[spec.name] = sheet
    written = LegacyJournal(path)
    day = dt.date(2026, 6, 1)
    for i in range(sessions):
        row = FIRST + i
        day += dt.timedelta(days=0 if rng.random() < 0.1 else rng.choice([1, 2, 3]))
        written.sessions.append(_legacy_session(sheets, row, day, rng))
    day = dt.date(2026, 6, 2)
    for i in range(rest_days):
        day += dt.timedelta(days=rng.choice([3, 5, 8]))
        written.rest_days.append(_legacy_rest_day(sheets["Rest Days"], FIRST + i, day, rng))
    book.save(path)
    return written


def _maybe(rng: random.Random, value: Any, blank: float = 0.15) -> Any:
    return None if rng.random() < blank else value


def _score(rng: random.Random) -> int | None:
    return None if rng.random() < 0.15 else rng.randint(1, 10)


def _put(sheet: Worksheet, row: int, values: Sequence[Any], start: int = 1) -> dict[str, Any]:
    """Write values from column `start`; the values by header (row 4)."""
    out = {}
    for i, value in enumerate(values, start=start):
        _text_cell(sheet.cell(row, i), value)
        out[str(sheet.cell(4, i).value)] = value
    return out


def _text_cell(cell: Any, value: Any) -> None:
    """A value as typed: a note starting with "=" is text, as the user meant it, not a formula."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _legacy_session(
    sheets: dict[str, Worksheet], row: int, day: dt.date, rng: random.Random
) -> dict[str, Any]:
    when = dt.datetime.combine(day, dt.time())
    values: dict[str, Any] = {}
    values |= _put(
        sheets["Session"],
        row,
        [
            when,
            _maybe(rng, rng.choice(VENUES)),
            _maybe(rng, rng.choice(["Indoor", "Outdoor", "Home wall"])),
            _maybe(rng, rng.choice(SESSION_TYPES), 0.05),
            _maybe(rng, "[ESEMPIO] Anna, Luca", 0.6),
            _maybe(rng, rng.choice([60, 90, 120, 150])),
            _maybe(rng, rng.choice(GRADES), 0.2),
            _maybe(rng, rng.choice(NOTES), 0.5),
        ],
    )
    sleep: Any = round(rng.uniform(5, 9) * 2) / 2
    if rng.random() < 0.15:
        sleep = str(sleep).replace(".", ",")  # typed with a decimal comma: text in the cell
    values |= {
        f"wb:{k}": v
        for k, v in _put(
            sheets["Wellbeing"],
            row,
            [
                _maybe(rng, sleep, 0.1),
                _score(rng),
                _score(rng),
                _score(rng),
                _maybe(rng, rng.choice(NOTES), 0.6),
                _score(rng),
                _score(rng),
                _score(rng),
                _maybe(rng, rng.choice(NOTES), 0.6),
            ],
            start=2,
        ).items()
    }
    values |= {
        f"ph:{k}": v
        for k, v in _put(
            sheets["Physical"],
            row,
            [
                *(_score(rng) for _ in range(5)),
                _maybe(rng, rng.choice(NOTES), 0.7),
                *(_score(rng) for _ in range(5)),
                _maybe(rng, rng.choice(NOTES), 0.7),
            ],
            start=2,
        ).items()
    }
    values |= {
        f"mn:{k}": v
        for k, v in _put(
            sheets["Mental"],
            row,
            [
                *(_score(rng) for _ in range(4)),
                _maybe(rng, rng.choice(NOTES), 0.7),
                *(_score(rng) for _ in range(5)),
                _maybe(rng, rng.choice(NOTES), 0.7),
            ],
            start=2,
        ).items()
    }
    values |= {
        f"fe:{k}": v
        for k, v in _put(
            sheets["Session Feel"],
            row,
            [
                _maybe(
                    rng,
                    rng.choice(
                        [
                            "Excellent – exceeded plan",
                            "Good – as planned",
                            "Ok – mixed",
                            "Poor – struggled",
                            "Incomplete – cut short",
                        ]
                    ),
                ),
                _maybe(
                    rng, rng.choice(["Yes – fully", "Mostly", "Partially", "No – changed a lot"])
                ),
                _score(rng),
                _score(rng),
                _score(rng),
                _maybe(rng, rng.choice(NOTES), 0.6),
            ],
            start=2,
        ).items()
    }
    values |= {
        f"rf:{k}": v
        for k, v in _put(
            sheets["Wins & Reflect"],
            row,
            [
                _maybe(rng, rng.choice(NOTES), 0.4),
                _score(rng),
                _maybe(rng, rng.choice(NOTES), 0.4),
                _score(rng),
                _maybe(rng, rng.choice(NOTES), 0.5),
                _score(rng),
                _maybe(rng, rng.choice(NOTES), 0.5),
                _score(rng),
                _score(rng),
                _maybe(
                    rng,
                    rng.choice(
                        [
                            "Major send",
                            "Solid progress",
                            "Small win",
                            "Insight / learning",
                            "Effort win",
                        ]
                    ),
                    0.4,
                ),
                _maybe(rng, rng.choice(NOTES), 0.6),
            ],
            start=2,
        ).items()
    }
    return values


def _legacy_rest_day(
    sheet: Worksheet, row: int, day: dt.date, rng: random.Random
) -> dict[str, Any]:
    return _put(
        sheet,
        row,
        [
            dt.datetime.combine(day, dt.time()),
            _maybe(rng, rng.choice(["Full rest", "Active recovery"])),
            _maybe(
                rng,
                rng.choice(
                    [
                        "None",
                        "Mobility / stretching",
                        "Walk / hike",
                        "Easy cardio",
                        "Yoga",
                        "Massage / self-massage",
                        "Other",
                    ]
                ),
            ),
            _maybe(rng, round(rng.uniform(6, 9) * 2) / 2),
            *(_score(rng) for _ in range(5)),
            *(_score(rng) for _ in range(5)),
            _maybe(rng, rng.choice(NOTES), 0.5),
        ],
    )


# ---- The new layout ----


@dataclass
class NewJournal:
    path: Path
    issues: list[tuple[str, int | None, str | None]]  # (tab, row, field) the loader must report


def new_journal(path: Path) -> NewJournal:
    """A fictional workbook in the new layout, typed by hand: one case of each kind (04 §1)."""
    day = dt.datetime(2026, 10, 20)
    tabs: dict[str, list[list[Any]]] = {
        "Session": [
            ["date", "session_no", "session_type", "setting", "duration_min", "notes"],
            [day, None, "climbing", "INDOOR", 90, "'po' stanco"],  # blank #, choices in any case
            [day, 2, "Hangboard", "Home wall", "45", "=non una formula"],  # a 2nd session
            [day + dt.timedelta(days=2), None, "Climbing", "Outdoor", 120, "[ESEMPIO] così"],
            [day + dt.timedelta(days=2), 1, "Strength", None, 30, None],  # duplicate of the above
            [None, None, "Climbing", None, 60, "[ESEMPIO] senza data"],  # a blank key with data
            [None, None, None, None, None, None],  # an empty buffer row
        ],
        # Reordered columns: the app matches on row 2 (D-014).
        "Wellbeing": [
            ["wb_mood_pre", "date", "wb_sleep_hours", "session_no", "wb_sleep_quality"],
            [7, day, "7,5", None, 8],  # a decimal comma
            [6, day + dt.timedelta(days=9), 6, None, 5],  # no Session row: an orphan
        ],
    }
    book = openpyxl.Workbook()
    book.remove(book.active)  # type: ignore[arg-type]
    for name, rows in tabs.items():
        sheet = book.create_sheet(name)
        for r, row in enumerate([[str(k).upper() for k in rows[0]], *rows], start=1):
            for c, value in enumerate(row, start=1):
                _text_cell(sheet.cell(r, c), value)
    book.save(path)
    return NewJournal(
        path,
        [
            ("Session", 6, None),  # a duplicate key (date, session 1)
            ("Session", 7, "date"),  # a blank key with data
            ("Wellbeing", 4, None),  # an orphan: no Session row
        ],
    )


def main(argv: list[str]) -> None:
    folder = Path(argv[1] if len(argv) > 1 else "tests/fixtures")
    folder.mkdir(parents=True, exist_ok=True)
    legacy_journal(folder / "legacy-journal.xlsx", seed=1)
    new_journal(folder / "new-journal.xlsx")
    print(f"wrote {folder}/legacy-journal.xlsx and {folder}/new-journal.xlsx (fictional)")


if __name__ == "__main__":
    main(sys.argv)
