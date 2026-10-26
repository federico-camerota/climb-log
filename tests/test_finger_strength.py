# SPDX-License-Identifier: GPL-3.0-or-later
"""The finger strength journal end to end on xlsx: sync, write, read, derive (D-019, D-059).

On a temporary workbook with fictional sessions: schema sync creates it, the app writes sessions
and sets (hangs and curls on the same day, a set number left blank, an update, a delete), and
reading it back gives the records and derived values the rules say, with no data issue, no
schema difference and no drift. Formula parity of the same module: `test_formula_parity.py`.
"""

import datetime as dt
from pathlib import Path

import openpyxl
import pytest

from climblog.expr import Value
from climblog.modules import Module, load_module
from climblog.modules.compute import compute_module
from climblog.storage import (
    Backups,
    RecordWriter,
    XlsxBackend,
    check_schema,
    find_drift,
    load_data,
    sync_module,
)

FINGERS = Path(__file__).parent.parent / "modules" / "finger-strength.yaml"
DAY = dt.date(2026, 10, 19)


@pytest.fixture(scope="module")
def module() -> Module:
    return load_module(FINGERS)


@pytest.fixture
def workbook(module: Module, tmp_path: Path) -> tuple[Path, XlsxBackend]:
    path = tmp_path / "finger-strength.xlsx"
    repo = XlsxBackend(path, backups=Backups(tmp_path / "backups"))
    assert sync_module(module, repo)[0] == "created the workbook"
    writer = RecordWriter(module, repo)
    # A morning hang session and an evening curl session on the same day (D-032).
    writer.add(
        "session",
        {
            "date": DAY,
            "session_type": "Max hangs",
            "start_time": dt.time(8, 0),
            **{"pre_fingers": 2, "pre_forearms": 2, "pre_shoulders": 3},
            "session_rpe": 9,
            "notes": "[ESEMPIO] Mattina, ottime sensazioni",
        },
    )
    writer.add(
        "session",
        {
            "date": DAY,
            "session_type": "Finger curls",
            "start_time": dt.time(20, 30),
            **{"pre_fingers": 4, "pre_forearms": 6, "pre_shoulders": 3},
        },
    )
    for n, kg in enumerate([12.5, 15], start=1):
        writer.add(
            "hang_set",
            {
                "date": DAY,
                "set_no": n,
                "edge_mm": 20,
                "grip": "Half crimp",
                "hang_s": 10,
                "added_kg": kg,
                "rest_s": 180,
                "rpe": 8,
            },
        )
    writer.add(
        "curl_set",
        {
            "date": DAY,
            "set_no": 1,
            "load_left_kg": 30,
            "reps_left": 6,
            "load_right_kg": 32.5,
            "reps_right": 6,
        },
    )
    writer.add(
        "curl_set",
        {
            "date": DAY,
            "set_no": 2,
            "load_left_kg": 30,
            "reps_left": 5,
            "load_right_kg": 32.5,
            "reps_right": 6,
            "completed": False,
        },
    )
    writer.update("session", {"date": DAY, "session_type": "Finger curls"}, {"session_rpe": 8})
    return path, repo


def derived(module: Module, repo: XlsxBackend) -> dict[str, list[dict[str, Value]]]:
    data = load_data(module, repo)
    assert data.issues == []
    tables = {name: e.records for name, e in data.entities.items()}
    return compute_module(module, tables, data.settings)


def test_sessions_and_sets_read_back_with_their_values(
    module: Module, workbook: tuple[Path, XlsxBackend]
) -> None:
    _, repo = workbook
    data = load_data(module, repo)
    sessions = data.entities["session"].records
    assert [(s["session_type"], s["start_time"], s["session_rpe"]) for s in sessions] == [
        ("Max hangs", dt.time(8, 0), 9),
        ("Finger curls", dt.time(20, 30), 8),  # the update
    ]
    assert sessions[0]["notes"] == "[ESEMPIO] Mattina, ottime sensazioni"
    hangs = data.entities["hang_set"].records
    assert [(h["set_no"], h["added_kg"], h["arms"], h["completed"]) for h in hangs] == [
        (1, 12.5, "Two arms", True),  # defaults: two arms, completed
        (2, 15, "Two arms", True),
    ]
    curls = data.entities["curl_set"].records
    assert [c["completed"] for c in curls] == [True, False]


def test_derived_values_follow_the_rules(
    module: Module, workbook: tuple[Path, XlsxBackend]
) -> None:
    _, repo = workbook
    result = derived(module, repo)
    hangs_session, curls_session = result["session"]
    # Readiness 11 - mean(2, 2, 3) = 8.7, worst 3 → 11 - 3 = 8 → GOOD
    assert (hangs_session["readiness"], hangs_session["readiness_status"]) == (8.7, "GOOD")
    # Evening: 11 - mean(4, 6, 3) = 6.7, worst 6 → 5 → WATCH (rated apart from the morning)
    assert (curls_session["readiness"], curls_session["readiness_status"]) == (6.7, "WATCH")
    # Each session counts only its own type's sets (the type is implied by the tab).
    assert (hangs_session["n_sets"], hangs_session["top_added_kg"]) == (2, 15)
    assert curls_session["n_sets"] == 2
    assert curls_session["top_added_kg"] is None
    # Epley: 30 * (1 + 6/30) = 36, 32.5 * 1.2 = 39 → imbalance 100 * (36 - 39) / 39 = -7.7
    assert (curls_session["est_1rm_left"], curls_session["est_1rm_right"]) == (36, 39)
    assert curls_session["lr_imbalance_pct"] == -7.7
    assert curls_session["curl_volume"] == 30 * 6 + 32.5 * 6 + 30 * 5 + 32.5 * 6
    assert result["curl_set"][1]["volume_kg"] == 30 * 5 + 32.5 * 6


def test_the_workbook_is_in_sync_and_its_formulas_intact(
    module: Module, workbook: tuple[Path, XlsxBackend]
) -> None:
    path, repo = workbook
    assert check_schema(module, repo) == []
    assert find_drift(module, repo) == []
    sheet = openpyxl.load_workbook(path)["Sessions"]
    keys = {str(c.value): int(c.column) for c in sheet[2] if c.value}
    assert str(sheet.cell(3, keys["auto:readiness"]).value).startswith("=IF(A3=")
    assert sheet.cell(202, keys["auto:readiness"]).value  # the 200-row buffer is there
    assert (path.parent / "backups").exists()  # the writes took a backup first


def test_deleting_a_session_deletes_its_sets(
    module: Module, workbook: tuple[Path, XlsxBackend]
) -> None:
    _, repo = workbook
    RecordWriter(module, repo).delete("session", {"date": DAY, "session_type": "Max hangs"})
    data = load_data(module, repo)
    assert [s["session_type"] for s in data.entities["session"].records] == ["Finger curls"]
    assert data.entities["hang_set"].records == []  # cascaded (D-049)
    assert len(data.entities["curl_set"].records) == 2  # the other session's sets stay
    assert data.issues == []


def test_a_set_typed_by_hand_with_a_blank_number_takes_the_next(
    module: Module, workbook: tuple[Path, XlsxBackend]
) -> None:
    path, _ = workbook
    book = openpyxl.load_workbook(path)
    sheet = book["Hang sets"]
    keys = {str(c.value): int(c.column) for c in sheet[2] if c.value}
    row = 5  # the first free buffer row, after the app's two sets
    sheet.cell(row, keys["date"]).value = dt.datetime.combine(DAY, dt.time())
    sheet.cell(row, keys["added_kg"]).value = 15
    book.save(path)
    hangs = load_data(module, XlsxBackend(path)).entities["hang_set"].records
    assert [h["set_no"] for h in hangs] == [1, 2, 3]  # blank = next (D-050)
    assert derived(module, XlsxBackend(path))["session"][0]["n_sets"] == 3
