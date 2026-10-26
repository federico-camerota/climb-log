# SPDX-License-Identifier: GPL-3.0-or-later
"""Formula columns in xlsx: written by sync into a buffer, topped up, checked for drift (D-054)."""

import datetime as dt
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from conftest import LoadYaml
from openpyxl.worksheet.worksheet import Worksheet

from climblog.modules import Module
from climblog.modules.formulas import normalize
from climblog.storage import Append, Table, XlsxBackend, find_drift, sync_module
from climblog.storage.drift import Drift, drift_differences, same

# A fictional module: sessions with a score and a status, and their sets.
MODULE = """\
module: test_columns
name: Test columns
version: 1
kind: journal
status_levels: [GOOD, WATCH, LOW]
settings:
  low_max: {type: number, default: 4}
scales:
  s10: {min: 1, max: 10, direction: higher_better}
entities:
  session:
    label: Session
    tab: Sessions
    key: [date]
    fields:
      date: {type: date}
      mood: {type: scale, scale: s10}
      energy: {type: scale, scale: s10}
    derived:
      score: "round(mean(mood, energy), 1)"
      status:
        bands_of: "min(mood, energy)"
        bands:
          - {le: settings.low_max, label: LOW}
          - {le: 6, label: WATCH}
          - {else: GOOD}
      sets: "count(children(set))"
  set:
    label: Set
    tab: Sets
    key: [date, set_no]
    parent: {entity: session, key: [date]}
    fields:
      date: {type: date}
      set_no: {type: integer}
      kg: {type: number}
"""

D20 = dt.datetime(2026, 10, 20)
SCORE = '=IF(A{0}="","",IF(COUNT(B{0},C{0})=0,"",ROUND(AVERAGE(B{0},C{0}),1)))'


@pytest.fixture
def module(load: LoadYaml) -> Module:
    return load(MODULE)


@pytest.fixture
def path(module: Module, tmp_path: Path) -> Path:
    path = tmp_path / "test.xlsx"
    sync_module(module, XlsxBackend(path))
    return path


def sheet(path: Path, name: str = "Sessions") -> Worksheet:
    book = openpyxl.load_workbook(path)
    result = book[name]
    assert isinstance(result, Worksheet)
    return result


def keys(ws: Worksheet) -> dict[str, int]:
    return {str(c.value): int(c.column) for c in ws[2] if c.value}


def type_rows(path: Path, rows: list[tuple[Any, ...]], start: int = 3) -> None:
    """Rows typed by hand into Sessions: date, mood, energy."""
    book = openpyxl.load_workbook(path)
    ws = book["Sessions"]
    for i, row in enumerate(rows):
        for j, value in enumerate(row, start=1):
            ws.cell(start + i, j).value = value
    book.save(path)


# ---- Sync writes the columns ----


def test_sync_writes_the_formula_columns(path: Path) -> None:
    ws = sheet(path)
    columns = keys(ws)
    assert list(columns) == [
        *("date", "mood", "energy", "auto:score", "auto:status", "auto:sets", "auto:_key"),
        *("updated_at", "deleted"),
    ]
    assert ws.cell(3, columns["auto:score"]).value == SCORE.format(3)
    assert ws.cell(202, columns["auto:score"]).value == SCORE.format(202)  # a 200-row buffer
    assert ws.cell(203, columns["auto:score"]).value is None
    assert ws.cell(3, columns["auto:score"]).number_format == "0.0"
    assert ws.cell(3, columns["auto:sets"]).number_format == "0"
    assert ws.cell(1, columns["auto:score"]).fill.fgColor.rgb == "00BFBFBF"  # AUTOMATIC grey
    assert ws.column_dimensions["G"].hidden  # auto:_key
    assert ws.cell(3, columns["auto:_key"]).value == '=IF(A3="","",INT(A3)*10)'


def test_statuses_are_coloured(path: Path) -> None:
    ws = sheet(path)
    rules = {
        (str(f.sqref), r.formula[0]): r.dxf.fill.fgColor.rgb
        for f in ws.conditional_formatting
        for r in f.rules
        if r.type == "cellIs"
    }
    assert rules == {
        ("E3:E1048576", '"GOOD"'): "00C6EFCE",
        ("E3:E1048576", '"WATCH"'): "00FFE699",
        ("E3:E1048576", '"LOW"'): "00F8CBAD",
    }


def test_the_buffer_grows_with_the_data(module: Module, path: Path) -> None:
    type_rows(path, [(D20, 5, 6)], start=190)  # 12 free rows left
    assert sync_module(module, XlsxBackend(path)) == []
    ws = sheet(path)
    score = keys(ws)["auto:score"]
    assert ws.cell(390, score).value == SCORE.format(390)
    assert ws.cell(391, score).value is None


def test_app_writes_top_the_buffer_up(module: Module, path: Path) -> None:
    repo = XlsxBackend(path)
    repo.apply([Append("Sessions", {"date": D20 + dt.timedelta(days=i)}) for i in range(160)])
    ws = sheet(path)
    score = keys(ws)["auto:score"]
    assert ws.cell(362, score).value == SCORE.format(362)  # 200 past row 162
    assert find_drift(module, repo) == []  # the copied formulas are the compiled ones


def test_a_tab_without_formula_columns_is_left_alone(load: LoadYaml, tmp_path: Path) -> None:
    module = load(
        """\
        module: test_columns
        name: Test columns
        version: 1
        kind: journal
        entities:
          rest:
            label: Rest
            tab: Rest
            key: [date]
            fields:
              date: {type: date}
        """
    )
    path = tmp_path / "test.xlsx"
    sync_module(module, XlsxBackend(path))
    XlsxBackend(path).apply([Append("Rest", {"date": D20})])
    assert sheet(path, "Rest").max_row == 3


# ---- Drift ----


def test_a_formula_typed_over_is_drift_and_sync_puts_it_back(module: Module, path: Path) -> None:
    type_rows(path, [(D20, 5, 6), (D20 + dt.timedelta(days=1), 7, 8)])
    book = openpyxl.load_workbook(path)
    score = keys(book["Sessions"])["auto:score"]
    book["Sessions"].cell(4, score).value = 9  # typed over by hand
    book["Sessions"].cell(3, keys(book["Sessions"])["auto:status"]).value = None  # deleted
    book["Sessions"].cell(50, score).value = "=1+1"  # in an empty buffer row
    book["Sessions"].cell(60, score).value = None  # deleted where no data needs it
    book.save(path)
    repo = XlsxBackend(path)
    drifts = find_drift(module, repo)
    assert [(d.kind, d.row, d.column, d.sheet) for d in drifts] == [
        ("formula", 3, "auto:status", None),
        ("formula", 4, "auto:score", 9),
        ("formula", 50, "auto:score", "=1+1"),
    ]
    assert str(drifts[1]) == "Sessions, row 4, auto:score: the formula was changed"
    assert sync_module(module, repo) == [
        "Sessions: rewrote 2 changed cells of 'auto:score'",
        "Sessions: rewrote 1 changed cells of 'auto:status'",
    ]
    assert find_drift(module, repo) == []


def test_formulas_are_compared_as_spreadsheets_rewrite_them() -> None:
    written = "=IF(A3=\"\",\"\",_xlfn.MAXIFS('Hang sets'!$G:$G,'Sessions'!$A:$A,TRUE))"
    saved = '=if(A3="","", MAXIFS(\'Hang sets\'!$G:$G,Sessions!$A:$A,TRUE()))'
    assert normalize(written) == normalize(saved)
    assert normalize('=A3="Max hangs"') != normalize('=A3="MAX HANGS"')  # strings kept


class Shown(XlsxBackend):
    """The workbook with the values a spreadsheet would have computed (openpyxl can't)."""

    def __init__(self, path: Path, values: dict[tuple[int, str], Any]) -> None:
        super().__init__(path)
        self.values = values

    def read_table(self, name: str) -> Table:
        table = super().read_table(name)
        if name != "Sessions":
            return table
        rows = []
        for row in table.rows:
            cells = dict(row.cells)
            for (number, key), value in self.values.items():
                if number == row.number:
                    cells[key] = value
            rows.append(type(row)(row.number, cells))
        return Table(table.name, table.columns, tuple(rows), table.numbers)


def test_values_other_than_the_app_are_drift(module: Module, path: Path) -> None:
    type_rows(path, [(D20, 5, 6), (D20 + dt.timedelta(days=1), 7, 8)])
    repo = Shown(path, {(3, "auto:score"): 5.5, (4, "auto:score"): 7.0, (4, "auto:status"): "GOOD"})
    drifts = find_drift(module, repo)
    assert [(d.kind, d.row, d.column, d.sheet, d.app) for d in drifts] == [
        ("value", 4, "auto:score", 7.0, 7.5),
    ]
    assert str(drifts[0]) == (
        "Sessions, row 4, auto:score: the sheet shows 7.0, the app computes 7.5"
    )


def test_drift_as_check_lines() -> None:
    drifts = [Drift("formula", "Sessions", n, "auto:score", 1, "=x") for n in range(3, 10)]
    drifts.append(Drift("value", "Sets", 4, "auto:sets", 2, 3))
    assert [f"{d.level}: {d}" for d in drift_differences(drifts)] == [
        "sync: Sessions: 'auto:score': formula changed in rows 3, 4, 5, 6, 7, …",
        "info: Sets: 'auto:sets': the sheet shows other values than the app in rows 4 "
        "(row 4: 2, app 3)",
    ]


@pytest.mark.parametrize(
    ("app", "sheet", "equal"),
    [
        (None, None, True),
        (None, "", True),
        (None, 0, False),
        (True, True, True),
        (True, 1, False),
        (1.0, 1.0000000000001, True),
        (1.0, 1.1, False),
        (dt.date(2026, 10, 20), D20, True),
        (dt.time(18, 30), dt.datetime(1899, 12, 30, 18, 30), True),
        ("GOOD", "GOOD", True),
        ("GOOD", "LOW", False),
    ],
)
def test_same(app: Any, sheet: Any, equal: bool) -> None:
    assert same(app, sheet) is equal


def test_missing_tabs_are_left_to_check(module: Module, tmp_path: Path) -> None:
    path = tmp_path / "test.xlsx"
    sync_module(module, XlsxBackend(path))
    book = openpyxl.load_workbook(path)
    del book["Sets"]
    book.save(path)
    assert find_drift(module, XlsxBackend(path)) == []


def test_top_up_copies_each_column_from_its_last_formula(module: Module, path: Path) -> None:
    book = openpyxl.load_workbook(path)
    sessions = book["Sessions"]
    assert isinstance(sessions, Worksheet)
    sessions.cell(202, keys(sessions)["auto:status"]).value = None  # the column's last formula
    book.save(path)
    XlsxBackend(path).apply(
        [Append("Sessions", {"date": D20 + dt.timedelta(days=i)}) for i in range(160)]
    )
    ws = sheet(path)
    assert ws.cell(362, keys(ws)["auto:score"]).value == SCORE.format(362)
    assert ws.cell(362, keys(ws)["auto:status"]).value is None  # nothing to copy from
