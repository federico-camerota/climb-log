# SPDX-License-Identifier: GPL-3.0-or-later
"""Writing a local workbook: what is particular to .xlsx (D-049): number formats, text cells,
merged cells, locks, atomic save, backups.

What every backend must do is in `test_backend_write.py`.
"""

import datetime as dt
from pathlib import Path

import openpyxl
import pytest
from workbooks import tab, write_workbook

from climblog.storage import (
    Append,
    Backups,
    StorageError,
    Update,
    WorkbookLocked,
    XlsxBackend,
)

D20, D21 = dt.datetime(2026, 10, 20), dt.datetime(2026, 10, 21)
KEYS = ["date", "kg", "note", "auto:double", "updated_at"]


def formula(row: int) -> str:
    return f'=IF($A{row}="","",B{row}*2)'


@pytest.fixture
def path(tmp_path: Path) -> Path:
    """Two records, a blank-key row with hand data, then two buffer rows with formulas."""
    return write_workbook(
        tmp_path / "fingers.xlsx",
        {
            "Sets": tab(
                KEYS,
                [D20, 10, None, formula(3)],
                [D21, 12.5, "Più riposo", formula(4)],
                [None, 15, None, formula(5)],  # no date: a data issue, never overwritten
                [None, None, None, formula(6)],
                [None, None, None, formula(7)],
            ),
            "_settings": [["key", "value"], ["low_max", 4]],
            "Labels": [["Date"]],  # no key row yet
        },
    )


@pytest.fixture
def book(path: Path) -> XlsxBackend:
    return XlsxBackend(path)


def test_strings_are_text_cells(book: XlsxBackend, path: Path) -> None:
    notes = ["=1+1", "'po' stanco", "+3 kg", "-2", "perché 💪", "riga 1\nriga 2", "  "]
    book.apply([Append("Sets", {"date": D20, "note": note}) for note in notes])
    workbook = openpyxl.load_workbook(path)
    cells = [workbook["Sets"].cell(row, 3) for row in range(6, 6 + len(notes))]
    assert [(c.value, c.data_type, c.number_format) for c in cells] == [
        (note, "s", "@") for note in notes
    ]
    workbook.close()


def test_dates_and_times_get_a_format(book: XlsxBackend, path: Path) -> None:
    stamp = dt.datetime(2026, 10, 24, 18, 30, 5)
    book.apply(
        [
            Append("Sets", {"date": dt.date(2026, 10, 24), "kg": dt.time(18, 30)}),
            Append("Sets", {"date": D20, "kg": dt.timedelta(minutes=90), "updated_at": stamp}),
        ]
    )
    workbook = openpyxl.load_workbook(path)
    formats = [[c.number_format for c in workbook["Sets"][row]] for row in (6, 7)]
    workbook.close()
    assert formats[0][:2] == ["DD/MM/YYYY", "HH:MM"]
    assert [formats[1][i] for i in (1, 4)] == ["[H]:MM", "DD/MM/YYYY HH:MM:SS"]


def test_a_format_chosen_in_the_sheet_is_kept(book: XlsxBackend, path: Path) -> None:
    row = book.read_table("Sets").rows[0]
    book.apply([Update("Sets", 3, {"date": dt.date(2026, 10, 19)}, row.cells)])
    workbook = openpyxl.load_workbook(path)
    assert (
        workbook["Sets"]["A3"].number_format == "yyyy-mm-dd h:mm:ss"
    )  # as `write_workbook` left it
    workbook.close()


def test_a_number_replacing_text_drops_the_text_format(book: XlsxBackend, path: Path) -> None:
    row = book.read_table("Sets").rows[1]
    book.apply([Update("Sets", 4, {"note": 7}, row.cells)])
    workbook = openpyxl.load_workbook(path)
    assert workbook["Sets"]["C4"].number_format == "General"
    workbook.close()


def test_a_refused_edit_leaves_the_file_as_it_was(book: XlsxBackend, path: Path) -> None:
    before = path.read_bytes()
    with pytest.raises(StorageError, match="tab 'Sets' has no column weight in row 2"):
        book.apply([Append("Sets", {"date": D20}), Append("Sets", {"weight": 70})])
    assert path.read_bytes() == before


def test_merged_cells_are_refused(path: Path) -> None:
    workbook = openpyxl.load_workbook(path)
    workbook["Sets"].merge_cells("B6:C6")
    workbook.save(path)
    with pytest.raises(StorageError, match="C6 is inside merged cells"):
        XlsxBackend(path).apply([Append("Sets", {"date": D20, "note": "x"})])


def test_a_failed_save_leaves_the_workbook_untouched(
    book: XlsxBackend, path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    before = path.read_bytes()

    def broken(*_: object) -> None:
        raise OSError("disk full")

    monkeypatch.setattr("climblog.storage.xlsx.os.replace", broken)
    with pytest.raises(OSError, match="disk full"):
        book.apply([Append("Sets", {"date": D20})])
    assert path.read_bytes() == before
    assert sorted(p.name for p in path.parent.iterdir()) == ["fingers.xlsx"]


def test_the_saved_file_keeps_its_permissions(book: XlsxBackend, path: Path) -> None:
    path.chmod(0o640)
    book.apply([Append("Sets", {"date": D20})])
    assert path.stat().st_mode & 0o777 == 0o640


@pytest.mark.parametrize(
    ("lock", "program"), [("~$fingers.xlsx", "Excel"), (".~lock.fingers.xlsx#", "LibreOffice")]
)
def test_an_open_workbook_is_not_written(
    book: XlsxBackend, path: Path, lock: str, program: str
) -> None:
    (path.parent / lock).write_text("")
    with pytest.raises(WorkbookLocked, match=f"fingers.xlsx is open in {program}: close it"):
        book.apply([Append("Sets", {"date": D20})])
    with pytest.raises(WorkbookLocked):
        book.write_kv("_settings", {"low_max": 5})
    assert len(book.read_table("Sets").rows) == 3  # reading still works


def test_a_missing_workbook_takes_no_backup(tmp_path: Path) -> None:
    with pytest.raises(StorageError, match=r"cannot read nothing\.xlsx"):
        XlsxBackend(tmp_path / "nothing.xlsx", backups=Backups(tmp_path / "b")).apply([])
    assert not (tmp_path / "b").exists()


class Clock:
    def __init__(self) -> None:
        self.now = dt.datetime(2026, 10, 24, 9, 0, 0)

    def __call__(self) -> dt.datetime:
        return self.now


def test_backups_before_the_first_write_then_every_interval(path: Path, tmp_path: Path) -> None:
    clock = Clock()
    folder = tmp_path / "backups" / "fingers"
    book = XlsxBackend(path, backups=Backups(folder, every=dt.timedelta(minutes=30), clock=clock))
    original = path.read_bytes()

    book.apply([Append("Sets", {"date": D20})])
    assert [p.name for p in folder.iterdir()] == ["20261024-090000.xlsx"]
    assert (folder / "20261024-090000.xlsx").read_bytes() == original  # taken before the write

    clock.now += dt.timedelta(minutes=29)
    book.apply([Append("Sets", {"date": D20})])
    clock.now += dt.timedelta(minutes=1)
    book.write_kv("_settings", {"low_max": 5})
    assert sorted(p.name for p in folder.iterdir()) == [
        "20261024-090000.xlsx",
        "20261024-093000.xlsx",
    ]


def test_backups_keep_the_newest(path: Path, tmp_path: Path) -> None:
    clock = Clock()
    folder = tmp_path / "backups"
    folder.mkdir()
    (folder / "notes.xlsx").write_text("")  # not a backup: left alone
    backups = Backups(folder, every=dt.timedelta(0), keep=2, clock=clock)
    for _ in range(4):
        backups.before_write(path)
        clock.now += dt.timedelta(hours=1)
    assert sorted(p.name for p in folder.iterdir()) == [
        "20261024-110000.xlsx",
        "20261024-120000.xlsx",
        "notes.xlsx",
    ]
    with pytest.raises(ValueError, match="keep at least one"):
        Backups(folder, keep=0)
