# SPDX-License-Identifier: GPL-3.0-or-later
"""Local backups of a Google module: Drive exports as `.xlsx` (04-storage.md §3, D-068).

On the in-memory Google, whose export is a real `.xlsx` of the sheet's values.
"""

import datetime as dt
import logging
import stat
from pathlib import Path

import openpyxl
import pytest
from backends import LIVE, GSheets
from conftest import LoadYaml
from fake_sheets import FakeSheetsClient, cell_data, spreadsheet_body
from workbooks import SYNC_MODULE, tab

from climblog.config import config_path, load_config
from climblog.storage import Append, Backups, GSheetsBackend, StorageError, sync_module
from climblog.storage.google_client import RestSheetsClient

START = dt.datetime(2026, 10, 30, 9, 0)


class Clock:
    def __init__(self) -> None:
        self.now = START

    def __call__(self) -> dt.datetime:
        return self.now


@pytest.fixture
def clock() -> Clock:
    return Clock()


@pytest.fixture
def client() -> FakeSheetsClient:
    return FakeSheetsClient()


def backed_up(client: FakeSheetsClient, folder: Path, clock: Clock) -> GSheetsBackend:
    spreadsheet_id = client.create(spreadsheet_body("Journal", {"T": tab(["a"], [1])}))
    backups = Backups(folder, clock=clock)
    return GSheetsBackend(spreadsheet_id, client, clock=clock, backups=backups)


def copies(folder: Path) -> list[str]:
    return sorted(p.name for p in folder.glob("*.xlsx")) if folder.exists() else []


def hand_edit(client: FakeSheetsClient, repo: GSheetsBackend, value: int) -> None:
    assert repo.spreadsheet_id is not None
    update = {"rows": [{"values": [cell_data(value)]}], "fields": "userEnteredValue",
              "start": {"sheetId": 1000, "rowIndex": 2, "columnIndex": 0}}  # fmt: skip
    client.batch_update(repo.spreadsheet_id, [{"updateCells": update}])


def test_a_write_is_backed_up_first(client: FakeSheetsClient, tmp_path: Path, clock: Clock) -> None:
    folder = tmp_path / "backups" / "journal"
    repo = backed_up(client, folder, clock)
    repo.apply([Append("T", {"a": 2})])
    assert copies(folder) == ["20261030-090000.xlsx"]
    copy = folder / "20261030-090000.xlsx"
    assert stat.S_IMODE(copy.stat().st_mode) == 0o600
    assert stat.S_IMODE(folder.stat().st_mode) == 0o700
    sheet = openpyxl.load_workbook(copy)["T"]
    assert [row[0].value for row in sheet.iter_rows(min_row=3)] == [1]  # before the write
    clock.now += dt.timedelta(minutes=29)
    repo.apply([Append("T", {"a": 3})])
    assert len(copies(folder)) == 1  # at most one every 30 minutes
    clock.now += dt.timedelta(minutes=1)
    repo.apply([Append("T", {"a": 4})])
    assert len(copies(folder)) == 2


def test_no_backup_no_write(client: FakeSheetsClient, tmp_path: Path, clock: Clock) -> None:
    repo = backed_up(client, tmp_path, clock)

    def offline(spreadsheet_id: str) -> bytes:
        raise StorageError("cannot reach Google")

    client.export_xlsx = offline  # type: ignore[method-assign]
    with pytest.raises(StorageError, match="cannot reach Google"):
        repo.apply([Append("T", {"a": 2})])
    assert "batch_update" not in client.calls
    assert [r.cells for r in repo.read_table("T").rows] == [{"a": 1}]


def test_a_change_seen_on_reading_is_backed_up(
    client: FakeSheetsClient, tmp_path: Path, clock: Clock
) -> None:
    repo = backed_up(client, tmp_path, clock)
    repo.read_table("T")  # the first read: nothing to compare with
    assert copies(tmp_path) == []
    hand_edit(client, repo, 2)  # on the phone
    clock.now += dt.timedelta(minutes=1)
    assert repo.read_table("T").rows[0].cells == {"a": 2}
    assert copies(tmp_path) == ["20261030-090100.xlsx"]
    hand_edit(client, repo, 3)
    clock.now += dt.timedelta(minutes=1)
    repo.read_table("T")
    assert len(copies(tmp_path)) == 1  # the same 30 minutes


def test_a_failed_backup_on_reading_is_only_logged(
    client: FakeSheetsClient, tmp_path: Path, clock: Clock, caplog: pytest.LogCaptureFixture
) -> None:
    repo = backed_up(client, tmp_path, clock)
    repo.read_table("T")

    def offline(spreadsheet_id: str) -> bytes:
        raise StorageError("cannot reach Google")

    client.export_xlsx = offline  # type: ignore[method-assign]
    hand_edit(client, repo, 2)
    clock.now += dt.timedelta(minutes=1)
    with caplog.at_level(logging.WARNING):
        assert repo.read_table("T").rows[0].cells == {"a": 2}
    assert caplog.messages == ["no backup of Google Sheet 'Journal': cannot reach Google"]


def test_sync_backs_up_a_sheet_that_exists(
    client: FakeSheetsClient, tmp_path: Path, clock: Clock, load: LoadYaml
) -> None:
    module = load(SYNC_MODULE)
    new = GSheetsBackend(None, client, clock=clock, backups=Backups(tmp_path, clock=clock))
    sync_module(module, new)
    assert copies(tmp_path) == []  # a new sheet: nothing to back up
    assert new.spreadsheet_id is not None
    again = GSheetsBackend(
        new.spreadsheet_id, client, clock=clock, backups=Backups(tmp_path, clock=clock)
    )
    sync_module(module, again)
    assert copies(tmp_path) == ["20261030-090000.xlsx"]


def test_no_backups_configured(client: FakeSheetsClient) -> None:
    spreadsheet_id = client.create(spreadsheet_body("Journal", {"T": tab(["a"], [1])}))
    repo = GSheetsBackend(spreadsheet_id, client)
    assert repo.backup() is None
    repo.apply([Append("T", {"a": 2})])
    assert "export_xlsx" not in client.calls


@pytest.mark.skipif(not LIVE, reason="set CLIMBLOG_GOOGLE_LIVE=1 to run on Google")
def test_a_live_backup_is_an_xlsx(tmp_path: Path) -> None:  # pragma: no cover - live only
    harness = GSheets(RestSheetsClient(load_config(config_path()).google_session), live=True)
    try:
        made = harness.make({"T": tab(["a", "note"], [1, "perché"])})
        repo = GSheetsBackend(made.spreadsheet_id, harness.client, backups=Backups(tmp_path))
        repo.apply([Append("T", {"a": 2})])
        (copy,) = tmp_path.glob("*.xlsx")
        rows = list(openpyxl.load_workbook(copy)["T"].iter_rows(min_row=3, values_only=True))
        assert rows[0][:2] == (1, "perché")  # the sheet before the write
    finally:
        harness.close()
