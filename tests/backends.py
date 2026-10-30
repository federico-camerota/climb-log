# SPDX-License-Identifier: GPL-3.0-or-later
"""The storage backends the contract tests run on (`test_backend_*.py`, 07-roadmap.md).

A contract test uses a backend only through the `Repository` protocol, plus this harness for
what the protocol doesn't offer: building a workbook from rows, reading a tab as stored
(formulas as text) and changing a cell by hand. Each backend has one harness here; the
`backend` fixture runs every contract test on each of them, and skips the contracts a backend
doesn't meet yet (`pending`).

`gsheets` is the Google backend on an in-memory Google (`fake_sheets.py`). With
`CLIMBLOG_GOOGLE_LIVE=1`, `gsheets-live` runs it on the real API with your login
(`climblog google login`): each workbook is a new sheet named `climblog-test-…`, moved to the
Drive trash when the test ends. Never in CI.
"""

import datetime as dt
import os
import uuid
from collections.abc import Mapping
from pathlib import Path
from typing import Any, Protocol

import openpyxl
from fake_sheets import FakeSheetsClient, cell_data, spreadsheet_body
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from openpyxl.worksheet.worksheet import Worksheet
from workbooks import Rows, write_workbook

from climblog.config import config_path, load_config
from climblog.storage import GSheetsBackend, Repository, XlsxBackend
from climblog.storage.google_client import RestSheetsClient, SheetsClient
from climblog.storage.gsheets import MAX_AGE, grid_rows

Cells = list[tuple[Any, ...]]


class Backend(Protocol):
    name: str
    pending: Mapping[str, str]  # contract → the issue that brings it to this backend

    def make(self, tabs: Mapping[str, Rows]) -> Repository:
        """A new workbook with these tabs, in order (row 1 first; formulas as `=…` text)."""
        ...

    def new(self) -> Repository:
        """A workbook that doesn't exist yet (schema sync creates it)."""
        ...

    def cells(self, repo: Repository, tab: str) -> Cells:
        """The tab as stored: formulas as text, rows up to the last one with a value, each as
        wide as the widest."""
        ...

    def edit(self, repo: Repository, tab: str, cell: str, value: Any) -> None:
        """Change a cell (`B4`) outside the app, as a hand edit would."""
        ...

    def label(self, repo: Repository) -> str:
        """How messages name the workbook."""
        ...

    def close(self) -> None:
        """Clean up after the test."""
        ...


class Xlsx:
    name = "xlsx"
    pending: Mapping[str, str] = {}

    def __init__(self, folder: Path) -> None:
        self.folder = folder
        self.count = 0

    def _path(self) -> Path:
        self.count += 1
        return self.folder / f"book{self.count}.xlsx"

    def make(self, tabs: Mapping[str, Rows]) -> XlsxBackend:
        return XlsxBackend(write_workbook(self._path(), tabs))

    def new(self) -> XlsxBackend:
        return XlsxBackend(self._path())

    def cells(self, repo: Repository, tab: str) -> Cells:
        assert isinstance(repo, XlsxBackend)
        workbook = openpyxl.load_workbook(repo.path)
        try:
            sheet = workbook[tab]
            assert isinstance(sheet, Worksheet)
            rows = [tuple(c.value for c in row) for row in sheet.iter_rows()]
        finally:
            workbook.close()
        while rows and all(v is None for v in rows[-1]):
            rows.pop()
        return rows

    def edit(self, repo: Repository, tab: str, cell: str, value: Any) -> None:
        assert isinstance(repo, XlsxBackend)
        workbook = openpyxl.load_workbook(repo.path)
        workbook[tab][cell] = value
        workbook.save(repo.path)
        later = repo.path.stat().st_mtime + 5
        os.utime(repo.path, (later, later))  # a clearly newer mtime, whatever the clock resolution

    def label(self, repo: Repository) -> str:
        assert isinstance(repo, XlsxBackend)
        return repo.path.name

    def close(self) -> None:
        pass


class GSheets:
    """The Google backend on `client`: the fake, or the real API (`live`)."""

    pending: Mapping[str, str] = {"sync": "#53"}

    def __init__(self, client: SheetsClient, *, live: bool) -> None:
        self.client, self.live = client, live
        self.name = "gsheets-live" if live else "gsheets"
        self.created: list[str] = []

    def make(self, tabs: Mapping[str, Rows]) -> GSheetsBackend:
        title = f"climblog-test-{uuid.uuid4().hex[:8]}"
        spreadsheet_id = self.client.create(spreadsheet_body(title, tabs))
        self.created.append(spreadsheet_id)
        return self._backend(spreadsheet_id)

    def new(self) -> GSheetsBackend:
        return self._backend(f"climblog-test-missing-{uuid.uuid4().hex}")

    def _backend(self, spreadsheet_id: str) -> GSheetsBackend:
        """Live, every read goes to Google: Drive would show an edit only minutes later."""
        max_age = dt.timedelta(0) if self.live else MAX_AGE
        return GSheetsBackend(spreadsheet_id, self.client, max_age=max_age)

    def cells(self, repo: Repository, tab: str) -> Cells:
        assert isinstance(repo, GSheetsBackend)
        locale = self.client.get_grid(repo.spreadsheet_id)["properties"].get("locale", "en_US")
        rows = grid_rows(self._sheet(repo, tab), formulas=True, locale=locale)
        while rows and all(v is None for v in rows[-1]):
            rows.pop()
        width = max((len(row) for row in rows), default=0)  # Google leaves out trailing blanks
        return [row + (None,) * (width - len(row)) for row in rows]

    def edit(self, repo: Repository, tab: str, cell: str, value: Any) -> None:
        assert isinstance(repo, GSheetsBackend)
        letters, row = coordinate_from_string(cell)
        start = {
            "sheetId": self._sheet(repo, tab)["properties"]["sheetId"],
            "rowIndex": row - 1,
            "columnIndex": column_index_from_string(letters) - 1,
        }
        rows = [{"values": [cell_data(value)]}]
        update = {"rows": rows, "start": start, "fields": "userEnteredValue,userEnteredFormat"}
        self.client.batch_update(repo.spreadsheet_id, [{"updateCells": update}])

    def label(self, repo: Repository) -> str:
        assert isinstance(repo, GSheetsBackend)
        return f"Google Sheet '{self.client.title(repo.spreadsheet_id)}'"

    def close(self) -> None:
        for spreadsheet_id in self.created:
            self.client.trash(spreadsheet_id)

    def _sheet(self, repo: Repository, tab: str) -> dict[str, Any]:
        assert isinstance(repo, GSheetsBackend)
        sheets = self.client.get_grid(repo.spreadsheet_id)["sheets"]
        return next(s for s in sheets if s["properties"]["title"] == tab)


LIVE = os.environ.get("CLIMBLOG_GOOGLE_LIVE") == "1"
BACKENDS = ["xlsx", "gsheets", *(["gsheets-live"] if LIVE else [])]


def harness(name: str, folder: Path) -> Backend:
    if name == "xlsx":
        return Xlsx(folder)
    if name == "gsheets":
        return GSheets(FakeSheetsClient(), live=False)
    if name == "gsheets-live":  # pragma: no cover - only with CLIMBLOG_GOOGLE_LIVE=1
        config = load_config(config_path())
        return GSheets(RestSheetsClient(config.google_session), live=True)
    raise ValueError(f"no harness for backend '{name}'")
