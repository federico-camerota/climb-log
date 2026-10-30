# SPDX-License-Identifier: GPL-3.0-or-later
import json
import logging
import os
import shutil
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from conftest import WriteYaml
from test_google_auth import CLIENT, SECRETS, FakeConsent, FakeGoogle
from typer.testing import CliRunner

from climblog import cli, google_auth
from climblog.google_auth import SCOPES
from climblog.modules import load_module
from climblog.storage import XlsxBackend, sync_module

runner = CliRunner()


@pytest.fixture
def uvicorn_calls(monkeypatch: pytest.MonkeyPatch) -> list[dict[str, Any]]:
    """Replace uvicorn.run with a stub that records its arguments instead of starting a server."""
    calls: list[dict[str, Any]] = []

    def fake_run(app: str, **kwargs: Any) -> None:
        calls.append({"app": app, **kwargs})

    monkeypatch.setattr("uvicorn.run", fake_run)
    return calls


@pytest.mark.parametrize("args", [["--help"], ["serve", "--help"]])
def test_help_exits_cleanly(args: list[str]) -> None:
    assert runner.invoke(cli.app, args).exit_code == 0


def test_serve_binds_localhost_on_default_port(uvicorn_calls: list[dict[str, Any]]) -> None:
    result = runner.invoke(cli.app, ["serve"])
    assert result.exit_code == 0
    assert uvicorn_calls == [
        {
            "app": "climblog.main:serve_app",
            "factory": True,
            "host": "127.0.0.1",
            "port": 8000,
            "reload": False,
        }
    ]


def test_serve_port_and_reload_options(uvicorn_calls: list[dict[str, Any]]) -> None:
    result = runner.invoke(cli.app, ["serve", "--port", "8123", "--reload"])
    assert result.exit_code == 0
    assert uvicorn_calls[0]["port"] == 8123
    assert uvicorn_calls[0]["reload"] is True
    assert uvicorn_calls[0]["host"] == "127.0.0.1"


def test_serve_has_no_host_option(uvicorn_calls: list[dict[str, Any]]) -> None:
    # Binding beyond localhost must wait for LAN mode with a PIN (D-010).
    result = runner.invoke(cli.app, ["serve", "--host", "0.0.0.0"])
    assert result.exit_code != 0
    assert uvicorn_calls == []


# ---- climblog check ----

FINGERS = Path("modules/finger-strength.yaml").absolute()
MOCKUP = Path("design/mockups/mockup-finger-strength.xlsx")


@pytest.fixture
def config(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """A config with the fictional finger-strength mock-up, synced, named by CLIMBLOG_CONFIG."""
    shutil.copy(MOCKUP, tmp_path / "fingers.xlsx")
    sync_module(load_module(FINGERS), XlsxBackend(tmp_path / "fingers.xlsx"))
    file = tmp_path / "config.toml"
    file.write_text(
        f'[modules.finger_strength]\ndefinition = "{FINGERS}"\nworkbook = "fingers.xlsx"\n',
        encoding="utf-8",
    )
    monkeypatch.setenv("CLIMBLOG_CONFIG", str(file))
    return file


def add_module(config: Path, name: str, definition: Path, workbook: str) -> None:
    with config.open("a", encoding="utf-8") as file:
        file.write(f'[modules.{name}]\ndefinition = "{definition}"\nworkbook = "{workbook}"\n')


def test_check_in_sync_changes_nothing(config: Path) -> None:
    workbook = config.parent / "fingers.xlsx"
    before = workbook.stat().st_mtime_ns
    result = runner.invoke(cli.app, ["check"])
    assert result.exit_code == 0
    assert result.output == f"finger_strength ({workbook})\n  ok\n"
    assert workbook.stat().st_mtime_ns == before
    assert not (config.parent / ".local").exists()  # no backup folder either


def test_check_lists_differences_and_fails(config: Path) -> None:
    workbook = openpyxl.load_workbook(config.parent / "fingers.xlsx")
    del workbook["Curl sets"]
    workbook["_meta"]["B3"] = 2
    workbook.save(config.parent / "fingers.xlsx")
    result = runner.invoke(cli.app, ["check", "finger_strength"])
    assert result.exit_code == 1
    assert result.output.splitlines()[1:] == [
        "  sync: _meta: synced with v2, the YAML is v3",
        "  sync: Curl sets: the tab is missing",
    ]


def test_check_info_alone_passes(config: Path) -> None:
    workbook = openpyxl.load_workbook(config.parent / "fingers.xlsx")
    workbook["_settings"].delete_rows(2)
    workbook.save(config.parent / "fingers.xlsx")
    result = runner.invoke(cli.app, ["check"])
    assert result.exit_code == 0
    assert "  info: _settings: no row for " in result.output


def test_check_reports_each_module(config: Path, write: WriteYaml) -> None:
    add_module(config, "broken", write("module: broken\n"), "fingers.xlsx")
    add_module(config, "fingers", FINGERS, "fingers.xlsx")
    result = runner.invoke(cli.app, ["check", "finger_strength", "broken", "fingers"])
    assert result.exit_code == 1
    lines = result.output.splitlines()
    assert lines[1] == "  ok"
    assert lines[2].startswith("broken (") and "test.yaml" in lines[3]
    assert lines[-1] == f"  error: {FINGERS} defines module 'finger_strength', not 'fingers'"


def test_check_missing_workbook(config: Path) -> None:
    (config.parent / "fingers.xlsx").unlink()
    result = runner.invoke(cli.app, ["check"])
    assert result.exit_code == 1
    assert result.output.splitlines()[1].startswith("  error: workbook not found")


def test_check_option_wins_over_environment(config: Path, tmp_path: Path) -> None:
    result = runner.invoke(cli.app, ["check", "--config", str(tmp_path / "other.toml")])
    assert result.exit_code == 1
    assert (
        result.output == f"{tmp_path / 'other.toml'}: no config file here (see README, Configure)\n"
    )


def test_check_unknown_or_no_modules(config: Path) -> None:
    result = runner.invoke(cli.app, ["check", "journal", "shoulder"])
    assert (result.exit_code, result.output) == (1, "not in the config: journal, shoulder\n")
    config.write_text("", encoding="utf-8")
    result = runner.invoke(cli.app, ["check"])
    assert (result.exit_code, result.output) == (1, f"{config}: no modules configured\n")


def test_check_apply_creates_the_workbook(config: Path) -> None:
    workbook = config.parent / "fingers.xlsx"
    workbook.unlink()
    with config.open("a", encoding="utf-8") as file:
        file.write('[backups]\nfolder = "backups"\n')
    result = runner.invoke(cli.app, ["check", "--apply"])
    assert result.exit_code == 0
    assert result.output.splitlines()[1:] == [
        "  synced: created the workbook",
        "  synced: created tab 'Sessions'",
        "  synced: created tab 'Hang sets'",
        "  synced: created tab 'Curl sets'",
        "  synced: created tab '_settings'",
        "  synced: created tab '_meta'",
        "  ok",
    ]
    result = runner.invoke(cli.app, ["check", "--apply"])
    assert result.output.splitlines()[1:] == [
        "  synced: nothing to add (dropdowns and _meta refreshed)",
        "  ok",
    ]
    assert len(list((config.parent / "backups" / "finger_strength").iterdir())) == 1


def test_check_apply_refuses_errors(config: Path) -> None:
    workbook = openpyxl.load_workbook(config.parent / "fingers.xlsx")
    workbook["_meta"]["B2"] = "journal"
    workbook.save(config.parent / "fingers.xlsx")
    result = runner.invoke(cli.app, ["check", "--apply"])
    assert result.exit_code == 1
    assert result.output.splitlines()[1].startswith("  error: not synced, fix these first: _meta")


def test_serve_passes_the_config_to_the_app(
    uvicorn_calls: list[dict[str, Any]], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delenv("CLIMBLOG_CONFIG", raising=False)
    result = runner.invoke(cli.app, ["serve", "--config", str(tmp_path / "c.toml")])
    assert result.exit_code == 0
    assert os.environ["CLIMBLOG_CONFIG"] == str((tmp_path / "c.toml").resolve())
    monkeypatch.delenv("CLIMBLOG_CONFIG")


# ---- climblog google (fictional credentials, a fake Google: test_google_auth.py) ----


@pytest.fixture
def google_config(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """A config with a fictional desktop client, the token in the temp folder; no real Google."""
    (tmp_path / "client.json").write_text(json.dumps(CLIENT), encoding="utf-8")
    file = tmp_path / "google.toml"
    file.write_text('[google]\nclient_secret = "client.json"\ntoken = "token.json"\n', "utf-8")
    monkeypatch.setenv("CLIMBLOG_CONFIG", str(file))
    monkeypatch.setattr(google_auth, "browser_consent", FakeConsent())
    return file


def use_google(monkeypatch: pytest.MonkeyPatch, google: FakeGoogle) -> None:
    monkeypatch.setattr(google_auth, "HttpRequest", lambda: google)


def no_secrets(output: str) -> bool:
    return not any(secret in output for secret in SECRETS)


def test_google_login_status_logout(
    google_config: Path, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    caplog.set_level(logging.DEBUG)
    token = google_config.parent / "token.json"
    login = runner.invoke(cli.app, ["google", "login", "--no-browser"])
    assert login.exit_code == 0, login.output
    assert login.output == f"logged in; token saved in {token} (readable by you only)\n"
    assert google_auth.browser_consent.asked[0][2] is False  # type: ignore[attr-defined]

    google = FakeGoogle()
    use_google(monkeypatch, google)
    status = runner.invoke(cli.app, ["google", "status"])
    assert status.exit_code == 0, status.output
    assert status.output.startswith(f"logged in ({token})\n  scopes: {', '.join(SCOPES)}\n")
    assert "access token valid until" in status.output
    assert len(google.calls) == 1  # status checks with Google

    logout = runner.invoke(cli.app, ["google", "logout"])
    assert logout.output == f"access revoked; {token} deleted\n"
    assert not token.exists()
    again = runner.invoke(cli.app, ["google", "logout"])
    assert again.output == "not logged in: nothing to do\n"
    outputs = [login.output, status.output, logout.output, caplog.text]
    assert all(no_secrets(output) for output in outputs)


def test_google_status_without_a_login(google_config: Path) -> None:
    result = runner.invoke(cli.app, ["google", "status"])
    assert result.exit_code == 1
    assert "token.json: not logged in to Google: run `climblog google login`" in result.output


def test_google_logout_when_google_is_unreachable(
    google_config: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runner.invoke(cli.app, ["google", "login"])
    use_google(monkeypatch, FakeGoogle(offline=True))
    result = runner.invoke(cli.app, ["google", "logout"])
    assert result.exit_code == 0
    assert "deleted, but Google did not confirm the revoke" in result.output
    assert "myaccount.google.com/permissions" in result.output


def test_google_login_errors(google_config: Path, tmp_path: Path) -> None:
    (tmp_path / "client.json").write_text('{"web": {}}', encoding="utf-8")
    result = runner.invoke(cli.app, ["google", "login"])
    assert result.exit_code == 1
    assert "error: " in result.output and "a web client" in result.output
    bad = tmp_path / "bad.toml"
    bad.write_text("[google]\nwhat = 1\n", encoding="utf-8")
    result = runner.invoke(cli.app, ["google", "login", "--config", str(bad)])
    assert (result.exit_code, result.output) == (1, f"{bad}: google.what: unknown key\n")


def test_check_names_a_google_module_and_needs_the_login(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("HOME", str(tmp_path / "home"))  # no real login is ever used here
    monkeypatch.delenv("XDG_CONFIG_HOME", raising=False)
    sheet = "1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789_-fict"
    file = tmp_path / "google-module.toml"
    file.write_text(
        f'[modules.finger_strength]\ndefinition = "{FINGERS}"\ngoogle_sheet = "{sheet}"\n',
        encoding="utf-8",
    )
    result = runner.invoke(cli.app, ["check", "--config", str(file)])
    assert result.exit_code == 1
    assert result.output.splitlines() == [
        f"finger_strength (Google Sheet {sheet})",
        "  error: not logged in to Google: run `climblog google login`",
    ]
