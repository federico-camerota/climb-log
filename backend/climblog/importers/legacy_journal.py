# SPDX-License-Identifier: GPL-3.0-or-later
# ruff: noqa: RUF001  (the legacy header labels use en dashes, as in "Mood (1-10)")
"""Import the legacy `Climbing_Journal.xlsx` into the journal module (00-current-state.md §1,
04-storage.md §5, D-061).

The legacy workbook is only read. Its layout is checked first (the header labels in row 4 of
every column read), then each `Session` row with a date becomes a session, with the values of
the same row in the other session tabs (they are linked to it by row), and each `Rest Days` row
a rest day. A second session on the same day gets `session_no` 2, and so on. The settings come
from the named ranges on `Start Here`.

Cells are read as the new module reads them (`parse_cell`); a value that can't be read is left
out and listed in the report (its place, never its content), and a row with data but no date
is skipped and listed. The parity report compares the values the legacy sheets computed with the
app's, column by column.
"""

import datetime as dt
import zipfile
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import column_index_from_string
from openpyxl.utils.datetime import to_excel
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from climblog.expr import Value
from climblog.modules import Module
from climblog.modules.compute import compute_module
from climblog.modules.schema import BaseField
from climblog.storage import (
    CellError,
    RecordWriter,
    Repository,
    StorageError,
    load_data,
    parse_cell,
    sync_module,
)
from climblog.storage.drift import same
from climblog.storage.load import SETTINGS_TAB

FIRST_ROW = 5  # legacy data rows start under four header rows

# Inputs: tab → (column, its row-4 label, the journal field). Labels as in the legacy file.
PARTS = ("fingers", "forearms", "shoulders", "wrists", "skin")
_SORE = [f"{p.capitalize()} (1–10 ↓)" for p in PARTS]
INPUTS: dict[str, list[tuple[str, str, str]]] = {
    "Session": [
        ("A", "Date", "date"),
        ("B", "Venue", "venue"),
        ("C", "Indoor / outdoor", "setting"),
        ("D", "Session type", "session_type"),
        ("E", "Partners", "partners"),
        ("F", "Duration (min)", "duration_min"),
        ("G", "Hardest grade tried", "hardest_grade"),
        ("H", "Notes", "notes"),
    ],
    "Wellbeing": [
        ("B", "Sleep (hours)", "wb_sleep_hours"),
        ("C", "Sleep quality (1–10)", "wb_sleep_quality"),
        ("D", "Mood (1–10)", "wb_mood_pre"),
        ("E", "Energy (1–10)", "wb_energy_pre"),
        ("F", "Pre-session note", "wb_note_pre"),
        ("G", "Post fatigue (1–10 ↓)", "wb_fatigue_post"),
        ("H", "Post mood (1–10)", "wb_mood_post"),
        ("I", "Recovery feel (1–10)", "wb_recovery_feel"),
        ("J", "Post-session note", "wb_note_post"),
    ],
    "Physical": [
        *((c, label, f"ph_pre_{p}") for c, label, p in zip("BCDEF", _SORE, PARTS, strict=True)),
        ("G", "Pre-session note", "ph_note_pre"),
        *((c, label, f"ph_post_{p}") for c, label, p in zip("HIJKL", _SORE, PARTS, strict=True)),
        ("M", "Post-session note", "ph_note_post"),
    ],
    "Mental": [
        ("B", "Life stress (1–10 ↓)", "mn_life_stress"),
        ("C", "Confidence (1–10)", "mn_confidence"),
        ("D", "Motivation (1–10)", "mn_motivation"),
        ("E", "Anxiety (1–10 ↓)", "mn_anxiety"),
        ("F", "Pre-session note", "mn_note_pre"),
        ("G", "Focus during (1–10)", "mn_focus"),
        ("H", "Commitment to falls (1–10)", "mn_fall_commitment"),
        ("I", "Fear response (1–10 ↓)", "mn_fear_response"),
        ("J", "Confidence on routes (1–10)", "mn_route_confidence"),
        ("K", "Approach mindset (1–10)", "mn_approach"),
        ("L", "Post-session note", "mn_note_post"),
    ],
    "Session Feel": [
        ("B", "Session outcome", "feel_outcome"),
        ("C", "Stuck to plan?", "feel_plan"),
        ("D", "Enjoyment (1–10)", "feel_enjoyment"),
        ("E", "Satisfaction (1–10)", "feel_satisfaction"),
        ("F", "Effort / energy used (1–10)", "feel_effort"),
        ("G", "Session note", "feel_note"),
    ],
    "Wins & Reflect": [
        ("B", "1 · What felt good", "rf_good"),
        ("C", "Score (1–10)", "rf_good_score"),
        ("D", "2 · What felt off", "rf_off"),
        ("E", "Severity (1–10 ↓)", "rf_off_severity"),
        ("F", "3 · What surprised you", "rf_surprise"),
        ("G", "Impact (1–10)", "rf_surprise_impact"),
        ("H", "4 · What would you change next time", "rf_change"),
        ("I", "Priority (1–10)", "rf_change_priority"),
        ("J", "Overall session score (1–10)", "rf_session_score"),
        ("K", "Win size", "rf_win_size"),
        ("L", "Win to bank", "rf_win"),
    ],
}
REST_TAB = "Rest Days"
REST_INPUTS: list[tuple[str, str, str]] = [
    ("A", "Date", "date"),
    ("B", "Rest type", "rest_type"),
    ("C", "Activity", "activity"),
    ("D", "Sleep (hours)", "sleep_hours"),
    ("E", "Sleep quality (1–10)", "sleep_quality"),
    ("F", "Mood (1–10)", "mood"),
    ("G", "Energy (1–10)", "energy"),
    ("H", "Motivation to climb (1–10)", "motivation"),
    ("I", "Life stress (1–10 ↓)", "life_stress"),
    *((c, label, f"body_{p}") for c, label, p in zip("JKLMN", _SORE, PARTS, strict=True)),
    ("O", "Note", "note"),
]
# What the legacy sheets computed: tab → (column, label, the journal's derived field).
COMPUTED: dict[str, list[tuple[str, str, str]]] = {
    "Session": [
        ("I", "Overall status (all tabs)", "overall_status"),
        ("J", "Log check", "log_check"),
        ("K", "Weekday", "weekday"),
        ("L", "Month", "month"),
        ("M", "Days since previous", "days_since_prev"),
    ],
    "Wellbeing": [
        ("K", "Pre status", "wb_pre_status"),
        ("L", "Post status", "wb_post_status"),
        ("M", "Wellbeing score", "wb_score"),
        ("N", "Wellbeing status", "wb_status"),
    ],
    "Physical": [
        ("N", "Pre-session readiness", "ph_readiness"),
        ("O", "Physical status", "ph_status"),
    ],
    "Mental": [("M", "Mental score", "mn_score"), ("N", "Mental status", "mn_status")],
    "Session Feel": [
        ("H", "Feel score", "feel_score"),
        ("I", "Feel status", "feel_status"),
        ("J", "Overall score", "overall_score"),
        ("K", "Overall status (all tabs)", "overall_status"),
    ],
    "Wins & Reflect": [("M", "Reflection status", "reflect_status")],
}
REST_COMPUTED: list[tuple[str, str, str]] = [
    ("P", "Wellbeing score", "wellbeing_score"),
    ("Q", "Body readiness", "body_readiness"),
    ("R", "Wellbeing status", "wellbeing_status"),
    ("S", "Body status", "body_status"),
    ("T", "Rest-day status", "status"),
    ("U", "Weekday", "weekday"),
    ("V", "Days since last climb", "days_since_climb"),
]
SETTINGS = {
    "LowMax": "low_max",
    "WatchMax": "watch_max",
    "AttnMin": "attn_min",
    "NoteMin": "note_min",
    "TrendDelta": "trend_delta",
}
# Legacy rules that differ on purpose: shown in the parity report, not counted as mismatches.
KNOWN = {
    ("rest_day", "days_since_climb"): "the legacy sheet counts a session on the same day, "
    "the app only earlier days (D-045)",
    ("session", "days_since_prev"): "the legacy sheet looks at the rows above, the app at "
    "earlier dates: they differ where rows are not in date order",
}


class LegacyLayoutError(StorageError):
    """The workbook doesn't have the legacy journal's layout."""


class ImportRefused(StorageError):
    """The target workbook already holds records: an import never mixes with them."""


RawValue = Any  # a cell as the legacy sheet holds it


@dataclass
class LegacyRecord:
    entity: str
    values: dict[str, Value]
    rows: dict[str, int]  # legacy tab → its row
    computed: dict[str, RawValue] = field(default_factory=dict)  # derived field → legacy value


@dataclass
class LegacyData:
    records: list[LegacyRecord]
    settings: dict[str, Value]
    skipped: list[str]  # rows left out, by place
    unread: list[str]  # values that couldn't be read, by place (never their content)

    def count(self, entity: str) -> int:
        return sum(1 for r in self.records if r.entity == entity)

    @property
    def computed(self) -> bool:
        """Whether the workbook holds the values its formulas computed. A file written by a
        program that doesn't compute formulas has none until opened and saved in a spreadsheet."""
        return any(v not in (None, "") for r in self.records for v in r.computed.values())


def read_legacy(path: Path, module: Module) -> LegacyData:
    """Every record of a legacy journal workbook, and the values its sheets computed."""
    try:
        book = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except (OSError, KeyError, ValueError, zipfile.BadZipFile, InvalidFileException) as error:
        raise StorageError(f"cannot read {path.name}: {error}") from None
    sheets = {name: _sheet(book, name) for name in [*INPUTS, REST_TAB]}
    for name, columns in [*INPUTS.items(), (REST_TAB, REST_INPUTS)]:
        _check_layout(sheets[name], [*columns, *COMPUTED.get(name, []), *_rest_computed(name)])
    data = LegacyData([], _settings(book, module), [], [])
    sessions = module.entities["session"]
    counts: dict[dt.date, int] = {}
    last = max(sheet.max_row for name, sheet in sheets.items() if name != REST_TAB)
    for row in range(FIRST_ROW, last + 1):
        day = _day(sheets["Session"].cell(row, 1).value)
        values: dict[str, Value] = {}
        has_data = False
        for tab, columns in INPUTS.items():
            for column, _, name in columns:
                if name == "date":
                    continue
                raw = sheets[tab][f"{column}{row}"].value
                has_data = has_data or raw not in (None, "")
                value = _read(module, sessions.fields[name], raw, f"{tab}!{column}{row}", data)
                if value is not None:
                    values[name] = value
        if day is None:
            if has_data:
                data.skipped.append(f"Session row {row}: data in the session tabs but no date")
            continue
        counts[day] = counts.get(day, 0) + 1
        record = LegacyRecord(
            "session",
            {"date": day, "session_no": counts[day], **values},
            dict.fromkeys(INPUTS, row),
        )
        for tab, columns in COMPUTED.items():
            for column, _, name in columns:
                record.computed[name] = sheets[tab][f"{column}{row}"].value
        data.records.append(record)
    rest = sheets[REST_TAB]
    rest_days = module.entities["rest_day"]
    for row in range(FIRST_ROW, rest.max_row + 1):
        day = _day(rest.cell(row, 1).value)
        values = {}
        has_data = False
        for column, _, name in REST_INPUTS[1:]:
            raw = rest[f"{column}{row}"].value
            has_data = has_data or raw not in (None, "")
            value = _read(module, rest_days.fields[name], raw, f"{REST_TAB}!{column}{row}", data)
            if value is not None:
                values[name] = value
        if day is None:
            if has_data:
                data.skipped.append(f"Rest Days row {row}: data but no date")
            continue
        record = LegacyRecord("rest_day", {"date": day, **values}, {REST_TAB: row})
        for column, _, name in REST_COMPUTED:
            record.computed[name] = rest[f"{column}{row}"].value
        data.records.append(record)
    return data


@dataclass
class ImportReport:
    counts: dict[str, int]
    settings: dict[str, Value]
    skipped: list[str]
    unread: list[str]


def import_legacy(path: Path, module: Module, repo: Repository) -> ImportReport:
    """Write a legacy journal into an empty (or new) journal workbook, in one save."""
    if repo.exists():
        found = load_data(module, repo)
        held = sum(len(e.records) for e in found.entities.values())
        if held:
            raise ImportRefused(
                f"the journal workbook already holds {held} records: import into a new one"
            )
    legacy = read_legacy(path, module)
    sync_module(module, repo)
    RecordWriter(module, repo).add_many([(r.entity, r.values) for r in legacy.records])
    repo.write_kv(SETTINGS_TAB, legacy.settings)
    return ImportReport(
        {name: legacy.count(name) for name in ("session", "rest_day")},
        legacy.settings,
        legacy.skipped,
        legacy.unread,
    )


@dataclass
class ParityColumn:
    entity: str
    field: str
    compared: int  # records where the legacy sheet had a computed value
    mismatches: list[tuple[dt.date, RawValue, Value]]  # (date, legacy, app)
    known: str | None = None  # a difference that is on purpose

    @property
    def matching(self) -> int:
        return self.compared - len(self.mismatches)


def parity(legacy: LegacyData, module: Module) -> list[ParityColumn]:
    """The legacy computed values against the app's, per derived field."""
    tables: dict[str, list[Mapping[str, Value]]] = {}
    for record in legacy.records:
        tables.setdefault(record.entity, []).append(record.values)
    derived = compute_module(module, tables, _with_defaults(module, legacy.settings))
    columns: dict[tuple[str, str], ParityColumn] = {}
    position = dict.fromkeys(tables, 0)
    for record in legacy.records:
        app = derived[record.entity][position[record.entity]]
        position[record.entity] += 1
        day = record.values["date"]
        assert isinstance(day, dt.date)
        for name, sheet_value in record.computed.items():
            column = columns.setdefault(
                (record.entity, name),
                ParityColumn(record.entity, name, 0, [], KNOWN.get((record.entity, name))),
            )
            if sheet_value in (None, "") and app.get(name) is None:
                continue
            column.compared += 1
            sheet_value = _as_number(sheet_value, app.get(name))
            if not same(app.get(name), sheet_value):
                column.mismatches.append((day, sheet_value, app.get(name)))
    return list(columns.values())


# ---- Reading cells ----


def _sheet(book: Any, name: str) -> Worksheet:
    if name not in book.sheetnames:
        raise LegacyLayoutError(f"no tab '{name}': not a legacy climbing journal")
    sheet = book[name]
    assert isinstance(sheet, Worksheet)
    return sheet


def _check_layout(sheet: Worksheet, columns: Sequence[tuple[str, str, str]]) -> None:
    for column, label, _ in columns:
        found = sheet.cell(4, column_index_from_string(column)).value
        if _label(found) != _label(label):
            raise LegacyLayoutError(
                f"{sheet.title}!{column}4 is {found!r}, expected {label!r}: "
                "the layout differs from the legacy journal"
            )


def _rest_computed(name: str) -> list[tuple[str, str, str]]:
    return REST_COMPUTED if name == REST_TAB else []


def _label(value: object) -> str:
    return " ".join(str(value or "").split())


def _day(value: object) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    return value if isinstance(value, dt.date) else None


def _read(module: Module, field: Any, raw: object, place: str, data: LegacyData) -> Value:
    assert isinstance(field, BaseField)
    try:
        return parse_cell(module, field, raw)  # type: ignore[arg-type]
    except CellError:
        data.unread.append(f"{place}: not a valid value for this column, left out")
        return None


def _settings(book: Any, module: Module) -> dict[str, Value]:
    """The named settings' values; a missing or unreadable one keeps the module default."""
    values: dict[str, Value] = {}
    for legacy, key in SETTINGS.items():
        defined = book.defined_names.get(legacy)
        if defined is None:
            continue
        for title, cell in defined.destinations:
            value = book[title][cell.replace("$", "")].value
            if isinstance(value, int | float) and not isinstance(value, bool):
                values[key] = value
    return values


def _as_number(sheet: RawValue, app: Value) -> RawValue:
    """A number the legacy sheet shows in a date or duration format (days since, …) as the
    number it is, when the app's value is a number."""
    if not isinstance(app, int | float) or isinstance(app, bool):
        return sheet
    if isinstance(sheet, dt.timedelta):
        return sheet.total_seconds() / 86400
    if isinstance(sheet, dt.datetime | dt.date | dt.time):
        return to_excel(sheet)  # type: ignore[no-untyped-call]
    return sheet


def _with_defaults(module: Module, values: Mapping[str, Value]) -> dict[str, Value]:
    defaults: dict[str, Value] = {k: s.default for k, s in module.settings.items()}
    return defaults | dict(values)
