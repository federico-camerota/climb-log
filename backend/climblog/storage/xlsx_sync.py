# SPDX-License-Identifier: GPL-3.0-or-later
"""Schema sync for `.xlsx`: apply a `WorkbookLayout` to an open workbook (04-storage.md §1, D-052).

Additive only: missing tabs and columns are created (new columns at the right, so the `auto:`
formulas' references stay valid), pending renames change the key row, and nothing is deleted.
Existing columns keep their labels, widths and colours, which may have been changed by hand.
Data validation and scale colours are rewritten on every sync for the module's columns, so the
dropdowns always match the YAML; rules on other columns are kept.
"""

import datetime as dt
from collections.abc import Mapping

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.comments import Comment
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
from openpyxl.worksheet.datavalidation import DataValidation, DataValidationList
from openpyxl.worksheet.worksheet import Worksheet

from climblog.modules.formulas import XLSX
from climblog.storage.base import AUTO, FIRST_DATA_ROW, StorageError
from climblog.storage.cells import RawCell
from climblog.storage.grid import formula_rows_end
from climblog.storage.layout import (
    EARLIEST_DATE,
    GROUP_COLOURS,
    META_TAB,
    SETTINGS_TAB,
    ColumnSpec,
    DateRule,
    ListRule,
    RangeRule,
    TabLayout,
    WorkbookLayout,
)

LAST_ROW = 1_048_576  # validation and colours cover the whole column below the headers
KEY_NOTE = (
    "Row 2 holds the field keys the app matches on. Do not edit this row; "
    "labels in row 1 can be reworded freely."
)
ERROR_TITLE = "climblog"
ERROR_TEXT = "Value outside the allowed range/choices"
# Scale cells from red (worst) through yellow to green (best), as in the mock-ups.
RED, YELLOW, GREEN = "F8696B", "FFEB84", "63BE7B"
_GOOD, _WATCH, _LOW = "C6EFCE", "FFE699", "F8CBAD"  # status cells (mock-ups)

_LABEL_FONT = Font(bold=True, color="1F1F1F")
_KEY_FONT = Font(name="Roboto Mono", size=9, color="404040")
_CENTRED = Alignment(horizontal="center", vertical="center", wrap_text=True)
_SEPARATOR = Border(bottom=Side(style="medium", color="404040"))


def apply_layout(
    workbook: Workbook, layout: WorkbookLayout, meta: Mapping[str, RawCell]
) -> list[str]:
    """Make the workbook match the layout; what was done, one line per change."""
    done: list[str] = []
    for tab in layout.tabs:
        if tab.name in workbook.sheetnames:
            done += _update_tab(_sheet(workbook, tab.name), tab)
        else:
            _create_tab(workbook, tab)
            done.append(f"created tab '{tab.name}'")
        _rewrite_rules(_sheet(workbook, tab.name), tab)
    done += _settings(workbook, layout)
    done += _write_formulas(workbook, layout)
    if META_TAB not in workbook.sheetnames:
        done.append(f"created tab '{META_TAB}'")
    _key_values(workbook, META_TAB, meta, overwrite=True)
    return done


# ---- Data tabs ----


def _create_tab(workbook: Workbook, tab: TabLayout) -> None:
    sheet = workbook.create_sheet(tab.name)
    columns = [spec for spec in tab.columns if not spec.retired]
    for i, spec in enumerate(columns, start=1):
        _header(sheet, i, spec)
    sheet.row_dimensions[1].height = 48
    sheet.row_dimensions[2].height = 18
    sheet.freeze_panes = sheet.cell(FIRST_DATA_ROW, 3)  # headers and the first two columns
    sheet.cell(2, 1).comment = Comment(KEY_NOTE, ERROR_TITLE)


def _update_tab(sheet: Worksheet, tab: TabLayout) -> list[str]:
    done = []
    keys = _keys(sheet)
    for spec in tab.columns:
        if spec.key in keys:
            continue
        old = spec.renamed_from
        if old is not None and old in keys:
            i = keys[old]
            sheet.cell(2, i).value = spec.key
            label = sheet.cell(1, i)
            if label.value in (spec.old_label, old):  # reworded labels are kept
                label.value = spec.label
            keys[spec.key] = i
            done.append(f"{sheet.title}: renamed column '{old}' to '{spec.key}'")
        elif not spec.retired:
            i = sheet.max_column + 1 if keys else 1  # at the right (or a tab left empty)
            _header(sheet, i, spec)
            keys[spec.key] = i
            done.append(f"{sheet.title}: added column '{spec.key}'")
    return done


def _header(sheet: Worksheet, column: int, spec: ColumnSpec) -> None:
    strong, pale = GROUP_COLOURS[spec.group]
    label, key = sheet.cell(1, column), sheet.cell(2, column)
    label.value, key.value = spec.label, spec.key
    label.font, key.font = _LABEL_FONT, _KEY_FONT
    label.fill = PatternFill("solid", fgColor=strong)
    key.fill = PatternFill("solid", fgColor=pale)
    label.alignment = key.alignment = _CENTRED
    key.border = _SEPARATOR
    sheet.column_dimensions[get_column_letter(column)].width = spec.width
    if spec.hidden:
        sheet.column_dimensions[get_column_letter(column)].hidden = True
    if spec.number_format and not spec.formula:  # formula cells get theirs when written
        # Rows typed later take the column's format; the cells already there get it too.
        sheet.column_dimensions[get_column_letter(column)].number_format = spec.number_format
        for row in range(FIRST_DATA_ROW, sheet.max_row + 1):
            cell = sheet.cell(row, column)
            if cell.value is None:
                cell.number_format = spec.number_format


def _rewrite_rules(sheet: Worksheet, tab: TabLayout) -> None:
    """Validation and scale colours for the module's columns, replacing the old ones there."""
    keys = _keys(sheet)
    columns = {keys[spec.key]: spec for spec in tab.columns if spec.key in keys}
    letters = {get_column_letter(i) for i in columns}

    kept = DataValidationList()
    for rule in sheet.data_validations.dataValidation:
        ranges = _outside(rule.sqref, letters)
        if ranges:
            rule.sqref = ranges
            kept.append(rule)
    sheet.data_validations = kept

    old = sheet.conditional_formatting
    sheet.conditional_formatting = ConditionalFormattingList()
    for formatting in old:
        ranges = _outside(formatting.sqref, letters)
        for rule in formatting.rules:
            if ranges:
                sheet.conditional_formatting.add(str(ranges), rule)

    for i, spec in sorted(columns.items()):
        cells = f"{get_column_letter(i)}{FIRST_DATA_ROW}:{get_column_letter(i)}{LAST_ROW}"
        validation = _validation(spec)
        if validation is not None:
            validation.add(cells)
            sheet.add_data_validation(validation)
        if spec.levels:
            for level, fill in zip(spec.levels, _level_fills(len(spec.levels)), strict=True):
                rule = CellIsRule(  # type: ignore[no-untyped-call]
                    operator="equal",
                    formula=[_quoted(level)],
                    fill=PatternFill("solid", fgColor=fill),
                )
                sheet.conditional_formatting.add(cells, rule)
        if (scale := spec.colour_scale) is not None:
            low, high = (GREEN, RED) if scale.lower_better else (RED, GREEN)
            sheet.conditional_formatting.add(
                cells,
                ColorScaleRule(  # type: ignore[no-untyped-call]
                    start_type="num",
                    start_value=scale.min,
                    start_color=low,
                    mid_type="num",
                    mid_value=(scale.min + scale.max) / 2,
                    mid_color=YELLOW,
                    end_type="num",
                    end_value=scale.max,
                    end_color=high,
                ),
            )


def _validation(spec: ColumnSpec) -> DataValidation | None:
    rule = spec.validation
    match rule:
        case ListRule():
            values = ",".join(rule.values)
            if len(values) > 255:  # Excel's limit for a list typed in the rule
                return None
            validation = DataValidation("list", formula1=f'"{values}"')
            if not rule.strict:
                validation.errorStyle = "warning"  # other values are accepted (`allow_other`)
        case RangeRule(min=low, max=high):
            if low is not None and high is not None:
                validation = DataValidation(
                    rule.kind, operator="between", formula1=_n(low), formula2=_n(high)
                )
            elif low is not None:
                validation = DataValidation(
                    rule.kind, operator="greaterThanOrEqual", formula1=_n(low)
                )
            else:
                assert high is not None
                validation = DataValidation(
                    rule.kind, operator="lessThanOrEqual", formula1=_n(high)
                )
        case DateRule(kind="date"):
            start = EARLIEST_DATE
            validation = DataValidation(
                "date",
                operator="greaterThanOrEqual",
                formula1=f"DATE({start.year},{start.month},{start.day})",
            )
        case DateRule(kind="time"):
            validation = DataValidation(
                "time", operator="between", formula1="0", formula2="0.999988"
            )
        case _:
            return None
    validation.allow_blank = True
    validation.showErrorMessage = True
    validation.errorTitle, validation.error = ERROR_TITLE, ERROR_TEXT
    return validation


def _outside(ranges: MultiCellRange, letters: set[str]) -> MultiCellRange:
    """The parts of `ranges` not on those columns (a rule spanning several columns is split)."""
    kept = MultiCellRange()
    for cells in ranges.ranges:
        columns = range(cells.min_col, cells.max_col + 1)
        if not any(get_column_letter(c) in letters for c in columns):
            kept.add(cells)
            continue
        for c in columns:
            if get_column_letter(c) not in letters:
                kept.add(
                    CellRange(min_col=c, min_row=cells.min_row, max_col=c, max_row=cells.max_row)
                )
    return kept


# ---- Formula columns (D-053) ----


def _write_formulas(workbook: Workbook, layout: WorkbookLayout) -> list[str]:
    """Write every `auto:` column's formula into each row of the tab's formula buffer.

    Rewritten on every sync: a formula changed by hand is put back (D-053), and only those
    cells are reported. The buffer runs from row 3 to at least row 202, and to 200 rows past
    the last data row when fewer than 50 are left.
    """
    columns = {sheet.title: _keys(sheet) for sheet in workbook.worksheets}
    templates = layout.formulas(columns, XLSX)
    done = []
    for tab in layout.tabs:
        sheet = _sheet(workbook, tab.name)
        keys = columns[tab.name]
        specs = {spec.key: spec for spec in tab.columns if spec.formula and spec.key in keys}
        if not specs:
            continue
        end = buffer_end(sheet, keys)
        data = [i for k, i in keys.items() if not k.startswith(AUTO)]
        with_data = {
            row
            for row in range(FIRST_DATA_ROW, end + 1)
            if any(sheet.cell(row, i).value not in (None, "") for i in data)
        }
        for key, spec in specs.items():
            template = templates.get((tab.name, key))
            if template is None:
                continue  # pragma: no cover - sync has just made every column it needs
            filled = differs = deleted = 0
            for row in range(FIRST_DATA_ROW, end + 1):
                cell = sheet.cell(row, keys[key])
                formula = template.at(row)
                if cell.value is not None:
                    filled += 1
                    differs += cell.value != formula
                elif row in with_data:  # deleted where it is needed
                    deleted += 1
                cell.value = formula
                cell.number_format = spec.number_format or "General"
            if not filled:
                continue  # a new column: reported as added
            if differs == filled:  # every cell: the YAML or the compiler changed the formula
                done.append(f"{tab.name}: new formula in '{key}'")
            elif differs + deleted:  # cells changed by hand (drift): put back
                done.append(f"{tab.name}: rewrote {differs + deleted} changed cells of '{key}'")
    return done


def buffer_end(sheet: Worksheet, keys: Mapping[str, int]) -> int:
    """The last row the formula buffer should reach (`formula_rows_end`)."""
    data = [i for key, i in keys.items() if not key.startswith(AUTO)]
    autos = [i for key, i in keys.items() if key.startswith(AUTO)]
    last_data = FIRST_DATA_ROW - 1
    last_formula = FIRST_DATA_ROW - 1
    for row in range(FIRST_DATA_ROW, sheet.max_row + 1):
        if any(sheet.cell(row, i).value not in (None, "") for i in data):
            last_data = row
        if any(str(sheet.cell(row, i).value or "").startswith("=") for i in autos):
            last_formula = row
    return formula_rows_end(last_data, last_formula)


def top_up(sheet: Worksheet, keys: Mapping[str, int]) -> None:
    """Extend the formula buffer after the app wrote rows, by copying the last formula row
    (relative references follow the row; whole-column ones stay)."""
    autos = [i for key, i in keys.items() if key.startswith(AUTO)]
    if not autos:
        return
    last = max(
        (
            r
            for r in range(FIRST_DATA_ROW, sheet.max_row + 1)
            if any(str(sheet.cell(r, i).value or "").startswith("=") for i in autos)
        ),
        default=None,
    )
    if last is None:
        return
    end = buffer_end(sheet, keys)
    for i in autos:
        source = sheet.cell(last, i)
        if not str(source.value or "").startswith("="):
            continue
        for row in range(last + 1, end + 1):
            target = sheet.cell(row, i)
            target.value = Translator(
                str(source.value), origin=source.coordinate
            ).translate_formula(target.coordinate)
            target.number_format = source.number_format


def _level_fills(n: int) -> list[str]:
    """Status colours, best first: green … amber … red (the mock-ups' GOOD/WATCH/LOW)."""
    return [_GOOD if i == 0 else _LOW if i == n - 1 else _WATCH for i in range(n)]


def _quoted(text: str) -> str:
    return '"' + text.replace('"', '""') + '"'


# ---- _settings and _meta ----


def _settings(workbook: Workbook, layout: WorkbookLayout) -> list[str]:
    """A row per setting (existing values kept) and a defined name on each value cell."""
    created = SETTINGS_TAB not in workbook.sheetnames
    values = {row.key: row.default for row in layout.settings}
    labels = {row.key: row.label for row in layout.settings}
    added = _key_values(workbook, SETTINGS_TAB, values, overwrite=False, labels=labels)
    sheet = _sheet(workbook, SETTINGS_TAB)
    rows = _kv_rows(sheet)
    value_column = get_column_letter(_kv_columns(sheet)[1])
    for key in values:
        reference = f"{quote_sheetname(SETTINGS_TAB)}!${value_column}${rows[key]}"
        workbook.defined_names[key] = DefinedName(key, attr_text=reference)
    if created:
        return [f"created tab '{SETTINGS_TAB}'"]
    return [f"{SETTINGS_TAB}: added '{key}'" for key in added]


def _key_values(
    workbook: Workbook,
    tab: str,
    values: Mapping[str, RawCell],
    *,
    overwrite: bool,
    labels: Mapping[str, str] | None = None,
) -> list[str]:
    """Set key/value rows, creating the tab (with a `label` column if given); keys added."""
    if tab not in workbook.sheetnames:
        sheet = workbook.create_sheet(tab)
        headers = ["key", "value", *(["label"] if labels is not None else [])]
        for i, header in enumerate(headers, start=1):
            sheet.cell(1, i).value = header
            sheet.cell(1, i).font = _LABEL_FONT
        sheet.column_dimensions["A"].width = 16
        sheet.column_dimensions["B"].width = 20
        sheet.column_dimensions["C"].width = 50
    sheet = _sheet(workbook, tab)
    k, v = _kv_columns(sheet)
    label_column = _header_index(sheet, "label")
    rows = _kv_rows(sheet)
    added = []
    for key, value in values.items():
        if key in rows and not overwrite:
            continue
        if key not in rows:
            rows[key] = max(sheet.max_row, 1) + 1
            sheet.cell(rows[key], k).value = key
            added.append(key)
            if labels is not None and label_column is not None:
                sheet.cell(rows[key], label_column).value = labels.get(key) or None
        _value(sheet.cell(rows[key], v), value)
    return added


def _value(cell: Cell | MergedCell, value: RawCell) -> None:
    """A setting or `_meta` value: Yes/No for booleans, text as text, dates formatted."""
    if isinstance(value, bool):
        value = "Yes" if value else "No"
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
        cell.number_format = "@"
    elif isinstance(value, dt.date):
        cell.number_format = "DD/MM/YYYY"


def _kv_columns(sheet: Worksheet) -> tuple[int, int]:
    """1-based `key` and `value` columns from row 1."""
    k, v = _header_index(sheet, "key"), _header_index(sheet, "value")
    if k is None or v is None:
        raise StorageError(f"tab '{sheet.title}' needs the headers 'key' and 'value' in row 1")
    return k, v


def _kv_rows(sheet: Worksheet) -> dict[str, int]:
    """Key → row; the upper row wins, as on read."""
    k = _kv_columns(sheet)[0]
    rows: dict[str, int] = {}
    for row in range(2, sheet.max_row + 1):
        key = sheet.cell(row, k).value
        if key is not None and str(key).strip():
            rows.setdefault(str(key).strip(), row)
    return rows


def _header_index(sheet: Worksheet, name: str) -> int | None:
    for cell in sheet[1]:
        if isinstance(cell.value, str) and cell.value.strip().casefold() == name:
            return int(cell.column)
    return None


# ---- Helpers ----


def _sheet(workbook: Workbook, name: str) -> Worksheet:
    sheet = workbook[name]
    assert isinstance(sheet, Worksheet)
    return sheet


def _keys(sheet: Worksheet) -> dict[str, int]:
    """Row-2 keys → 1-based column."""
    keys: dict[str, int] = {}
    for cell in sheet[2] if sheet.max_row >= 2 else ():
        if isinstance(cell.value, str) and cell.value.strip():
            keys.setdefault(cell.value.strip(), int(cell.column))
    return keys


def _n(value: float) -> str:
    return f"{value:g}"
