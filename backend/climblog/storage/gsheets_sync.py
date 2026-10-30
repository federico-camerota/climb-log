# SPDX-License-Identifier: GPL-3.0-or-later
"""Schema sync on a Google Sheet (04-storage.md §1, §3, D-052, D-054, D-065).

The same changes as `xlsx_sync.py`, as `batchUpdate` requests built from one read of the sheet
(`SheetsClient.get_layout`) and sent in one call, which Google applies all or none: tabs with
their two header rows (frozen, a note on A2), columns added at the right, renames, dropdowns and
ranges, scale colours and status colours rewritten on the module's columns, the `auto:` formulas
(Google dialect) in the formula buffer, `_settings` with a named range per setting, `_meta`.
Google adds what `.xlsx` can't: the header rows and the `auto:` columns get a *warning only*
protection, so an edit there asks first (04 §1).
"""

import datetime as dt
from collections.abc import Iterator, Mapping
from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter

from climblog.modules.formulas import GOOGLE, normalize
from climblog.storage.base import AUTO, FIRST_DATA_ROW
from climblog.storage.cells import RawCell
from climblog.storage.google_client import Json
from climblog.storage.grid import blank, formula_rows_end, key_columns, kv_columns
from climblog.storage.gsheets import (
    DATE_FORMAT,
    comma_decimal,
    grid_rows,
    local_formula,
    serial,
)
from climblog.storage.layout import (
    EARLIEST_DATE,
    GROUP_COLOURS,
    META_TAB,
    SETTINGS_TAB,
    ColumnSpec,
    DateRule,
    ListRule,
    RangeRule,
    TabLayout,
    WorkbookLayout,
)
from climblog.storage.xlsx_sync import GREEN, KEY_NOTE, RED, YELLOW, level_fills

PROTECTED = "climb-log: "  # the description of the app's protected ranges, to find them again
NEW_ROWS, NEW_COLUMNS = 1000, 26  # a new tab's grid, as Google makes it
_PIXELS_PER_CHARACTER = 7  # column widths are given in characters, as for .xlsx
_INK = "1F1F1F"
_KEY_INK = "404040"


def sync_requests(
    book: Json, layout: WorkbookLayout, meta: Mapping[str, RawCell], *, locale: str | None = None
) -> tuple[list[Json], list[str]]:
    """The requests that make the sheet (a `get_layout` response) match the layout, and what
    they do, one line per change (as `xlsx_sync.apply_layout` reports it). `locale` is the
    sheet's once the requests apply, if they change it."""
    sync = _Sync(book, locale)
    done: list[str] = []
    for tab in layout.tabs:
        done += sync.data_tab(tab)
    done += sync.settings(layout)
    done += sync.formulas(layout)
    if META_TAB not in sync.tabs:
        done.append(f"created tab '{META_TAB}'")
    sync.key_values(META_TAB, meta, overwrite=True)
    return sync.requests, done


@dataclass
class _Tab:
    name: str
    sheet_id: int
    rows: list[tuple[RawCell, ...]]  # values as read; formulas in en_US syntax in `formulas`
    formulas: list[tuple[RawCell, ...]]
    row_count: int
    column_count: int
    conditional: list[Json] = field(default_factory=list)
    protected: list[Json] = field(default_factory=list)

    @property
    def width(self) -> int:
        return max((len(row) for row in self.rows), default=0)

    def keys(self) -> dict[str, int]:
        """Row-2 keys → 0-based column."""
        return key_columns(self.name, self.rows[1] if len(self.rows) > 1 else ())

    def cell(self, row: int, column: int) -> RawCell:
        values = self.rows[row - 1] if row <= len(self.rows) else ()
        return values[column] if column < len(values) else None

    def set(self, row: int, column: int, value: RawCell) -> None:
        """Keep the model in step with a write (later steps read it)."""
        while len(self.rows) < row:
            self.rows.append(())
            self.formulas.append(())
        for rows in (self.rows, self.formulas):
            values = list(rows[row - 1])
            values += [None] * (column + 1 - len(values))
            values[column] = value
            rows[row - 1] = tuple(values)


class _Sync:
    def __init__(self, book: Json, target: str | None) -> None:
        locale = str(book.get("properties", {}).get("locale", "en_US"))
        # What `batchUpdate` sends is read in the sheet's locale, cell formulas and rule values
        # alike: 5,5 and `;` in it_IT (measured on 2026-10-30).
        self.comma = comma_decimal(target or locale)
        self.tabs: dict[str, _Tab] = {}
        for sheet in book.get("sheets", []):
            props = sheet["properties"]
            grid = props.get("gridProperties", {})
            name = str(props["title"])
            self.tabs[name] = _Tab(
                name,
                props.get("sheetId", 0),
                grid_rows(sheet),
                grid_rows(sheet, formulas=True, locale=locale),
                grid.get("rowCount", NEW_ROWS),
                grid.get("columnCount", NEW_COLUMNS),
                list(sheet.get("conditionalFormats", [])),
                list(sheet.get("protectedRanges", [])),
            )
        self.named = {str(n["name"]): str(n["namedRangeId"]) for n in book.get("namedRanges", [])}
        self.requests: list[Json] = []
        self._next_id = max((t.sheet_id for t in self.tabs.values()), default=0) + 1

    # ---- Data tabs ----

    def data_tab(self, layout: TabLayout) -> list[str]:
        done = []
        if layout.name in self.tabs:
            tab = self.tabs[layout.name]
            done += self._update_tab(tab, layout)
        else:
            columns = [spec for spec in layout.columns if not spec.retired]
            tab = self._add_tab(layout.name, max(NEW_COLUMNS, len(columns)), frozen=True)
            for i, spec in enumerate(columns):
                self._header(tab, i, spec)
            self._heights(tab)
            self._update(tab, 2, 0, {"note": KEY_NOTE}, "note")
            done.append(f"created tab '{layout.name}'")
        self._rewrite_rules(tab, layout)
        self._protect(tab, layout)
        return done

    def _update_tab(self, tab: _Tab, layout: TabLayout) -> list[str]:
        done = []
        keys = tab.keys()
        width = tab.width if keys else 0  # at the right (or from A in a tab left empty)
        for spec in layout.columns:
            if spec.key in keys:
                continue
            old = spec.renamed_from
            if old is not None and old in keys:
                i = keys[old]
                self._value(tab, 2, i, spec.key)
                if tab.cell(1, i) in (spec.old_label, old):  # reworded labels are kept
                    self._value(tab, 1, i, spec.label)
                keys[spec.key] = i
                done.append(f"{tab.name}: renamed column '{old}' to '{spec.key}'")
            elif not spec.retired:
                if width >= tab.column_count:
                    self._append(tab, "COLUMNS", width + 1 - tab.column_count)
                self._header(tab, width, spec)
                keys[spec.key] = width
                width += 1
                done.append(f"{tab.name}: added column '{spec.key}'")
        return done

    def _add_tab(self, name: str, columns: int, *, frozen: bool) -> _Tab:
        tab = _Tab(name, self._next_id, [], [], NEW_ROWS, columns)
        self._next_id += 1
        grid: Json = {"rowCount": NEW_ROWS, "columnCount": columns}
        if frozen:
            grid |= {"frozenRowCount": FIRST_DATA_ROW - 1, "frozenColumnCount": 2}
        props = {"sheetId": tab.sheet_id, "title": name, "gridProperties": grid}
        self.requests.append({"addSheet": {"properties": props}})
        self.tabs[name] = tab
        return tab

    def _header(self, tab: _Tab, column: int, spec: ColumnSpec) -> None:
        strong, pale = GROUP_COLOURS[spec.group]
        centred = {
            "horizontalAlignment": "CENTER",
            "verticalAlignment": "MIDDLE",
            "wrapStrategy": "WRAP",
        }
        label = {
            "backgroundColorStyle": _colour(strong),
            "textFormat": {"bold": True, "foregroundColorStyle": _colour(_INK)},
            **centred,
        }
        key = {
            "backgroundColorStyle": _colour(pale),
            "textFormat": {
                "fontFamily": "Roboto Mono",
                "fontSize": 9,
                "foregroundColorStyle": _colour(_KEY_INK),
            },
            "borders": {"bottom": {"style": "SOLID_MEDIUM", "colorStyle": _colour(_KEY_INK)}},
            **centred,
        }
        fields = "userEnteredValue,userEnteredFormat"
        self._update(tab, 1, column, {**_text(spec.label), "userEnteredFormat": label}, fields)
        self._update(tab, 2, column, {**_text(spec.key), "userEnteredFormat": key}, fields)
        tab.set(1, column, spec.label)
        tab.set(2, column, spec.key)
        dimension: Json = {"pixelSize": round(spec.width * _PIXELS_PER_CHARACTER)}
        if spec.hidden:
            dimension["hiddenByUser"] = True
        self.requests.append(
            {
                "updateDimensionProperties": {
                    "range": {
                        "sheetId": tab.sheet_id,
                        "dimension": "COLUMNS",
                        "startIndex": column,
                        "endIndex": column + 1,
                    },
                    "properties": dimension,
                    "fields": ",".join(dimension),
                }
            }
        )
        form = google_format(spec.number_format)
        if form is not None and not spec.formula:  # formula cells get theirs when written
            self.requests.append(
                {
                    "repeatCell": {
                        "range": _columns(tab, column, column + 1),
                        "cell": {"userEnteredFormat": {"numberFormat": form}},
                        "fields": "userEnteredFormat.numberFormat",
                    }
                }
            )

    def _heights(self, tab: _Tab) -> None:
        for row, pixels in ((0, 64), (1, 24)):
            self.requests.append(
                {
                    "updateDimensionProperties": {
                        "range": {
                            "sheetId": tab.sheet_id,
                            "dimension": "ROWS",
                            "startIndex": row,
                            "endIndex": row + 1,
                        },
                        "properties": {"pixelSize": pixels},
                        "fields": "pixelSize",
                    }
                }
            )

    def _rewrite_rules(self, tab: _Tab, layout: TabLayout) -> None:
        """Validation and colours on the module's columns, replacing what was there; rules on
        other columns are kept (a rule spanning both is cut down to the others)."""
        keys = tab.keys()
        specs = {keys[spec.key]: spec for spec in layout.columns if spec.key in keys}
        for index in reversed(range(len(tab.conditional))):  # later indexes first
            rule = tab.conditional[index]
            ranges = rule.get("ranges", [])
            if not any(_touches(cells, set(specs)) for cells in ranges):
                continue
            kept = [part for cells in ranges for part in _outside(cells, set(specs))]
            self.requests.append(
                {"deleteConditionalFormatRule": {"sheetId": tab.sheet_id, "index": index}}
            )
            if kept:
                self.requests.append(
                    {
                        "addConditionalFormatRule": {
                            "rule": {**rule, "ranges": kept},
                            "index": index,
                        }
                    }
                )
        for i, spec in sorted(specs.items()):
            cells = _columns(tab, i, i + 1)
            validation = _validation(spec, i, self.comma)
            body: Json = {"range": cells}
            if validation is not None:
                body["rule"] = validation
            self.requests.append({"setDataValidation": body})
            levels = spec.levels or ()
            for level, fill in zip(levels, level_fills(len(levels)), strict=True):
                self._conditional(
                    cells,
                    {
                        "booleanRule": {
                            "condition": {
                                "type": "TEXT_EQ",
                                "values": [{"userEnteredValue": level}],
                            },
                            "format": {"backgroundColorStyle": _colour(fill)},
                        }
                    },
                )
            if (scale := spec.colour_scale) is not None:
                low, high = (GREEN, RED) if scale.lower_better else (RED, GREEN)
                middle = (scale.min + scale.max) / 2
                self._conditional(
                    cells,
                    {
                        "gradientRule": {
                            "minpoint": _point(scale.min, low, self.comma),
                            "midpoint": _point(middle, YELLOW, self.comma),
                            "maxpoint": _point(scale.max, high, self.comma),
                        }
                    },
                )

    def _conditional(self, cells: Json, rule: Json) -> None:
        self.requests.append(
            {
                "addConditionalFormatRule": {
                    "rule": {"ranges": [cells], **rule},
                    "index": 0,
                }
            }
        )

    def _protect(self, tab: _Tab, layout: TabLayout) -> None:
        """Warning-only protection on the header rows and the `auto:` columns, made again on
        every sync."""
        for protected in tab.protected:
            if str(protected.get("description", "")).startswith(PROTECTED):
                rid = protected["protectedRangeId"]
                self.requests.append({"deleteProtectedRange": {"protectedRangeId": rid}})
        keys = tab.keys()
        ranges = [
            (
                "header rows",
                {"sheetId": tab.sheet_id, "startRowIndex": 0, "endRowIndex": FIRST_DATA_ROW - 1},
            )
        ]
        ranges += [
            (spec.key, _columns(tab, keys[spec.key], keys[spec.key] + 1))
            for spec in layout.columns
            if spec.formula and spec.key in keys
        ]
        for what, cells in ranges:
            self.requests.append(
                {
                    "addProtectedRange": {
                        "protectedRange": {
                            "range": cells,
                            "description": f"{PROTECTED}{what}",
                            "warningOnly": True,
                        }
                    }
                }
            )

    # ---- Formula columns (D-053, D-054) ----

    def formulas(self, layout: WorkbookLayout) -> list[str]:
        """Every `auto:` formula into each row of the formula buffer, as `xlsx_sync` does."""
        columns = {
            name: {k: i + 1 for k, i in tab.keys().items()} for name, tab in self.tabs.items()
        }
        templates = layout.formulas(columns, GOOGLE)
        done = []
        for layout_tab in layout.tabs:
            tab = self.tabs[layout_tab.name]
            keys = tab.keys()
            specs = {s.key: s for s in layout_tab.columns if s.formula and s.key in keys}
            if not specs:
                continue
            data = [i for k, i in keys.items() if not k.startswith(AUTO)]
            autos = [i for k, i in keys.items() if k.startswith(AUTO)]
            with_data = {
                n
                for n in range(FIRST_DATA_ROW, len(tab.rows) + 1)
                if any(not _empty(tab.cell(n, i)) for i in data)
            }
            last_formula = max(
                (
                    n
                    for n in range(FIRST_DATA_ROW, len(tab.formulas) + 1)
                    for i in autos
                    if _is_formula(tab.formulas[n - 1], i)
                ),
                default=FIRST_DATA_ROW - 1,
            )
            end = formula_rows_end(max(with_data, default=FIRST_DATA_ROW - 1), last_formula)
            if end > tab.row_count:
                self._append(tab, "ROWS", end - tab.row_count)
            for key, spec in specs.items():
                template = templates.get((layout_tab.name, key))
                if template is None:
                    continue  # pragma: no cover - sync has just made every column it needs
                i = keys[key]
                filled = differs = deleted = 0
                cells = []
                for n in range(FIRST_DATA_ROW, end + 1):
                    old = (
                        tab.formulas[n - 1][i]
                        if n <= len(tab.formulas) and i < len(tab.formulas[n - 1])
                        else None
                    )
                    formula = template.at(n)
                    if old is not None:
                        filled += 1
                        differs += not isinstance(old, str) or normalize(old) != normalize(formula)
                    elif n in with_data:  # deleted where it is needed
                        deleted += 1
                    written = local_formula(formula) if self.comma else formula
                    cell: Json = {"userEnteredValue": {"formulaValue": written}}
                    if (form := google_format(spec.number_format)) is not None:
                        cell["userEnteredFormat"] = {"numberFormat": form}
                    cells.append({"values": [cell]})
                self.requests.append(
                    {
                        "updateCells": {
                            "rows": cells,
                            "start": {
                                "sheetId": tab.sheet_id,
                                "rowIndex": FIRST_DATA_ROW - 1,
                                "columnIndex": i,
                            },
                            "fields": "userEnteredValue,userEnteredFormat.numberFormat",
                        }
                    }
                )
                if not filled:
                    continue  # a new column: reported as added
                if differs == filled:  # every cell: the YAML or the compiler changed it
                    done.append(f"{layout_tab.name}: new formula in '{key}'")
                elif differs + deleted:  # cells changed by hand (drift): put back
                    done.append(
                        f"{layout_tab.name}: rewrote {differs + deleted} changed cells of '{key}'"
                    )
        return done

    # ---- _settings and _meta ----

    def settings(self, layout: WorkbookLayout) -> list[str]:
        """A row per setting (values there are kept) and a named range on each value cell."""
        created = SETTINGS_TAB not in self.tabs
        values = {row.key: row.default for row in layout.settings}
        labels = {row.key: row.label for row in layout.settings}
        added = self.key_values(SETTINGS_TAB, values, overwrite=False, labels=labels)
        tab = self.tabs[SETTINGS_TAB]
        _, v = kv_columns(SETTINGS_TAB, tab.rows[0] if tab.rows else ())
        rows = _kv_rows(tab)
        for key in values:
            if key in self.named:
                self.requests.append({"deleteNamedRange": {"namedRangeId": self.named[key]}})
            cell = {
                "sheetId": tab.sheet_id,
                "startRowIndex": rows[key] - 1,
                "endRowIndex": rows[key],
                "startColumnIndex": v,
                "endColumnIndex": v + 1,
            }
            self.requests.append({"addNamedRange": {"namedRange": {"name": key, "range": cell}}})
        if created:
            return [f"created tab '{SETTINGS_TAB}'"]
        return [f"{SETTINGS_TAB}: added '{key}'" for key in added]

    def key_values(
        self,
        name: str,
        values: Mapping[str, RawCell],
        *,
        overwrite: bool,
        labels: Mapping[str, str] | None = None,
    ) -> list[str]:
        """Set key/value rows, making the tab (with a `label` column if given); keys added."""
        if name not in self.tabs:
            tab = self._add_tab(name, NEW_COLUMNS, frozen=False)
            headers = ["key", "value", *(["label"] if labels is not None else [])]
            bold = {"textFormat": {"bold": True, "foregroundColorStyle": _colour(_INK)}}
            for i, header in enumerate(headers):
                cell = {**_text(header), "userEnteredFormat": bold}
                self._update(tab, 1, i, cell, "userEnteredValue,userEnteredFormat")
                tab.set(1, i, header)
            for i, characters in enumerate((16, 20, 50)):
                self.requests.append(
                    {
                        "updateDimensionProperties": {
                            "range": {
                                "sheetId": tab.sheet_id,
                                "dimension": "COLUMNS",
                                "startIndex": i,
                                "endIndex": i + 1,
                            },
                            "properties": {"pixelSize": characters * _PIXELS_PER_CHARACTER},
                            "fields": "pixelSize",
                        }
                    }
                )
        tab = self.tabs[name]
        k, v = kv_columns(name, tab.rows[0] if tab.rows else ())
        headers = [str(h).strip().casefold() if not blank(h) else "" for h in tab.rows[0]]
        label_column = headers.index("label") if "label" in headers else None
        rows = _kv_rows(tab)
        added = []
        for key, value in values.items():
            if key in rows and not overwrite:
                continue
            if key not in rows:
                rows[key] = max(len(tab.rows), 1) + 1
                self._value(tab, rows[key], k, key)
                added.append(key)
                if labels is not None and label_column is not None and labels.get(key):
                    self._value(tab, rows[key], label_column, labels[key])
            self._value(tab, rows[key], v, value)
        return added

    # ---- Requests ----

    def _value(self, tab: _Tab, row: int, column: int, value: RawCell) -> None:
        """A label, key or key/value cell: text as text, Yes/No for booleans, dates formatted."""
        if isinstance(value, bool):
            value = "Yes" if value else "No"
        if isinstance(value, str):
            cell: Json = _text(value)
        elif isinstance(value, dt.date):
            number, _ = serial(value)
            cell = {
                "userEnteredValue": {"numberValue": number},
                "userEnteredFormat": {"numberFormat": DATE_FORMAT},
            }
        elif value is None:
            cell = {}
        else:
            assert isinstance(value, int | float)
            cell = {"userEnteredValue": {"numberValue": value}}
        self._update(tab, row, column, cell, "userEnteredValue,userEnteredFormat.numberFormat")
        tab.set(row, column, value)

    def _update(self, tab: _Tab, row: int, column: int, cell: Json, fields: str) -> None:
        if row > tab.row_count:
            self._append(tab, "ROWS", row - tab.row_count)
        self.requests.append(
            {
                "updateCells": {
                    "rows": [{"values": [cell]}],
                    "start": {"sheetId": tab.sheet_id, "rowIndex": row - 1, "columnIndex": column},
                    "fields": fields,
                }
            }
        )

    def _append(self, tab: _Tab, dimension: str, length: int) -> None:
        self.requests.append(
            {
                "appendDimension": {
                    "sheetId": tab.sheet_id,
                    "dimension": dimension,
                    "length": length,
                }
            }
        )
        if dimension == "ROWS":
            tab.row_count += length
        else:
            tab.column_count += length


def new_sheet_requests(book: Json, locale: str, time_zone: str) -> tuple[list[Json], list[Json]]:
    """For a spreadsheet just made empty: its locale and time zone first, and last the removal
    of the tab Google put in it (only once the module's tabs exist)."""
    first = [
        {
            "updateSpreadsheetProperties": {
                "properties": {"locale": locale, "timeZone": time_zone},
                "fields": "locale,timeZone",
            }
        }
    ]
    last = [
        {"deleteSheet": {"sheetId": s["properties"].get("sheetId", 0)}}
        for s in book.get("sheets", [])
    ]
    return first, last


def google_format(number_format: str | None) -> Json | None:
    """An `.xlsx` number format of the layout as the API's `NumberFormat` (None: automatic)."""
    if number_format in (None, "General"):
        return None
    assert number_format is not None
    if number_format == "@":
        return {"type": "TEXT"}
    pattern = number_format.lower()
    if "y" in pattern:
        return {"type": "DATE_TIME" if "h" in pattern else "DATE", "pattern": pattern}
    if "h" in pattern:
        return {"type": "TIME", "pattern": pattern}
    return {"type": "NUMBER", "pattern": number_format}


def _validation(spec: ColumnSpec, column: int, comma: bool) -> Json | None:
    """The API's rule for a column; numbers and formulas written as the sheet's locale reads
    them (`comma`: decimal commas and `;`)."""
    rule = spec.validation
    cell = f"{get_column_letter(column + 1)}{FIRST_DATA_ROW}"
    match rule:
        case ListRule():
            values = [{"userEnteredValue": value} for value in rule.values]
            condition = {"type": "ONE_OF_LIST", "values": values}
            return {"condition": condition, "strict": rule.strict, "showCustomUi": True}
        case RangeRule(min=low, max=high):
            if rule.kind == "whole":  # Google has no whole-number rule: a formula
                tests = [f"{cell}=INT({cell})"]
                tests += [f"{cell}>={_n(low)}"] if low is not None else []
                tests += [f"{cell}<={_n(high)}"] if high is not None else []
                formula = f"=AND(ISNUMBER({cell}),{','.join(tests)})"
                formula = local_formula(formula) if comma else formula
                condition = {"type": "CUSTOM_FORMULA", "values": [{"userEnteredValue": formula}]}
            elif low is not None and high is not None:
                condition = {"type": "NUMBER_BETWEEN", "values": [_v(low, comma), _v(high, comma)]}
            elif low is not None:
                condition = {"type": "NUMBER_GREATER_THAN_EQ", "values": [_v(low, comma)]}
            else:
                assert high is not None
                condition = {"type": "NUMBER_LESS_THAN_EQ", "values": [_v(high, comma)]}
        case DateRule(kind="date"):
            start = EARLIEST_DATE.isoformat()
            condition = {"type": "DATE_ON_OR_AFTER", "values": [{"userEnteredValue": start}]}
        case DateRule(kind="time"):
            condition = {"type": "NUMBER_BETWEEN", "values": [_v(0, comma), _v(0.999988, comma)]}
        case _:
            return None
    return {"condition": condition, "strict": True}


def _columns(tab: _Tab, first: int, end: int) -> Json:
    """Columns `first` to `end` (0-based, end excluded), from row 3 down."""
    return {
        "sheetId": tab.sheet_id,
        "startRowIndex": FIRST_DATA_ROW - 1,
        "startColumnIndex": first,
        "endColumnIndex": end,
    }


def _touches(cells: Json, columns: set[int]) -> bool:
    first, end = cells.get("startColumnIndex", 0), cells.get("endColumnIndex")
    return end is not None and any(c in columns for c in range(first, end))


def _outside(cells: Json, columns: set[int]) -> Iterator[Json]:
    """The parts of a range not on those columns (one range per column when cut)."""
    if not _touches(cells, columns):  # whole-row ranges are left as they are
        yield cells
        return
    for c in range(cells.get("startColumnIndex", 0), cells["endColumnIndex"]):
        if c not in columns:
            yield {**cells, "startColumnIndex": c, "endColumnIndex": c + 1}


def _kv_rows(tab: _Tab) -> dict[str, int]:
    """Key → row (1-based); the upper row wins, as on read."""
    k, _ = kv_columns("", tab.rows[0] if tab.rows else ())
    rows: dict[str, int] = {}
    for number in range(2, len(tab.rows) + 1):
        key = tab.cell(number, k)
        if not blank(key):
            rows.setdefault(str(key).strip(), number)
    return rows


def _point(value: float, colour: str, comma: bool) -> Json:
    return {"type": "NUMBER", "value": _n(value, comma), "colorStyle": _colour(colour)}


def _colour(hex_rgb: str) -> Json:
    red, green, blue = (int(hex_rgb[i : i + 2], 16) / 255 for i in (0, 2, 4))
    return {"rgbColor": {"red": red, "green": green, "blue": blue}}


def _text(value: str) -> Json:
    return {"userEnteredValue": {"stringValue": value}}


def _v(value: float, comma: bool) -> Json:
    return {"userEnteredValue": _n(value, comma)}


def _n(value: float | None, comma: bool = False) -> str:
    text = f"{value:g}"
    return text.replace(".", ",") if comma else text


def _empty(value: RawCell) -> bool:
    return value in (None, "")


def _is_formula(values: tuple[RawCell, ...], i: int) -> bool:
    value = values[i] if i < len(values) else None
    return isinstance(value, str) and value.startswith("=")
