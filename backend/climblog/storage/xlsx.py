# SPDX-License-Identifier: GPL-3.0-or-later
"""The local `.xlsx` backend (04-storage.md §2, D-048, D-049).

The workbook is read in one pass, read-only, as the sheet shows it (`data_only`: a formula
gives its last computed value), and kept until the file changes.

A write loads the whole workbook with its formulas, applies the edits and saves it atomically
(a temp file in the same folder, then `os.replace`). It is refused while Excel or LibreOffice
has the file open, and a backup is taken first when one is due. Strings are written as text
cells, so a note starting with `=` or `'` stays as typed (04 §3b).
"""

import datetime as dt
import os
import shutil
import tempfile
import zipfile
from collections.abc import Mapping, Sequence
from decimal import Decimal
from pathlib import Path
from typing import Protocol

import openpyxl
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from climblog.storage.backup import Backups
from climblog.storage.base import (
    AUTO,
    FIRST_DATA_ROW,
    Append,
    Edit,
    SheetRow,
    StorageError,
    Table,
    TableNotFound,
    Update,
    WorkbookLocked,
    WriteConflict,
)
from climblog.storage.cells import RawCell
from climblog.storage.layout import WorkbookLayout
from climblog.storage.xlsx_sync import apply_layout, top_up

Sheet = list[tuple[RawCell, ...]]

# Formats for cells written into General (or text-formatted) cells, as in the mock-ups.
_FORMATS: list[tuple[type, str]] = [
    (dt.datetime, "DD/MM/YYYY HH:MM:SS"),
    (dt.date, "DD/MM/YYYY"),
    (dt.time, "HH:MM"),
    (dt.timedelta, "[H]:MM"),
]


class XlsxBackend:
    """A `Repository` on a local workbook."""

    def __init__(self, path: Path, *, backups: Backups | None = None) -> None:
        self.path = path
        self.backups = backups
        self._cache: tuple[int, dict[str, Sheet]] | None = None

    def list_tables(self) -> list[str]:
        return list(self._sheets())

    def read_table(self, name: str) -> Table:
        return _table(name, self._sheet(name))

    def read_formulas(self, name: str) -> Table:
        workbook = self._load(read_only=True, data_only=False)
        try:
            if name not in workbook.sheetnames:
                raise TableNotFound(f"no tab '{name}' in {self.path.name}")
            sheet = workbook[name]
            rows = [tuple(_raw(v) for v in row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
        return _table(name, rows)

    def read_kv(self, name: str) -> dict[str, RawCell]:
        rows = self._sheet(name)
        k, v = _kv_columns(name, rows[0] if rows else ())
        pairs: dict[str, RawCell] = {}
        for row in rows[1:]:
            key = row[k] if k < len(row) else None
            if not _blank(key):
                pairs.setdefault(str(key).strip(), row[v] if v < len(row) else None)
        return pairs

    def last_modified(self) -> dt.datetime:
        return dt.datetime.fromtimestamp(self._stat_mtime_ns() / 1e9)

    def apply(self, edits: Sequence[Edit]) -> list[int]:
        def write(workbook: Workbook) -> list[int]:
            free: dict[str, _FreeRows] = {}
            numbers = []
            for edit in edits:
                sheet = _worksheet(workbook, edit.tab)
                columns = _columns(edit.tab, _row_values(sheet, FIRST_DATA_ROW - 1))
                _check_columns(edit.tab, columns, edit.cells)
                if isinstance(edit, Append):
                    if edit.tab not in free:
                        free[edit.tab] = _FreeRows(sheet, columns)
                    number = free[edit.tab].take()
                else:
                    number = edit.row
                    _check_unchanged(sheet, columns, edit)
                for key, value in edit.cells.items():
                    _write(sheet.cell(number, columns[key] + 1), value)
                numbers.append(number)
            for tab in free:  # rows were taken from the formula buffer: keep 50 free (D-054)
                sheet = _worksheet(workbook, tab)
                columns = _columns(tab, _row_values(sheet, FIRST_DATA_ROW - 1))
                top_up(sheet, {key: i + 1 for key, i in columns.items()})
            return numbers

        return self._write(write)

    def write_kv(self, name: str, values: Mapping[str, RawCell]) -> None:
        def write(workbook: Workbook) -> None:
            sheet = _worksheet(workbook, name)
            k, v = _kv_columns(name, _row_values(sheet, 1))
            rows = {}
            for number in range(sheet.max_row, 1, -1):  # bottom up: the first of a key wins
                key = _raw(sheet.cell(number, k + 1).value)
                if not _blank(key):
                    rows[str(key).strip()] = number
            for key, value in values.items():
                number = rows.get(key) or sheet.max_row + 1
                rows[key] = number
                _write(sheet.cell(number, k + 1), key)
                _write(sheet.cell(number, v + 1), value)

        self._write(write)

    def exists(self) -> bool:
        return self.path.exists()

    def sync_schema(self, layout: WorkbookLayout, meta: Mapping[str, RawCell]) -> list[str]:
        if self.exists():
            return self._write(lambda workbook: apply_layout(workbook, layout, meta))
        self._check_unlocked()
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)  # type: ignore[arg-type]
        done = apply_layout(workbook, layout, meta)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._save(workbook)
        self._cache = None
        return ["created the workbook", *done]

    # ---- Writing ----

    def _write[T](self, change: "_Change[T]") -> T:
        """Load the workbook with its formulas, change it, save it atomically."""
        self._check_unlocked()
        workbook = self._load(read_only=False, data_only=False)
        try:
            if self.backups is not None:
                self.backups.before_write(self.path)
            result = change(workbook)
            self._save(workbook)
        finally:
            workbook.close()
        self._cache = None
        return result

    def _check_unlocked(self) -> None:
        folder, name = self.path.parent, self.path.name
        for lock, program in ((f"~${name}", "Excel"), (f".~lock.{name}#", "LibreOffice")):
            if (folder / lock).exists():
                raise WorkbookLocked(f"{name} is open in {program}: close it to save")

    def _save(self, workbook: Workbook) -> None:
        handle, temp = tempfile.mkstemp(
            dir=self.path.parent, prefix=f".{self.path.stem}-", suffix=self.path.suffix
        )
        os.close(handle)
        try:
            workbook.save(temp)
            with open(temp, "rb+") as file:
                os.fsync(file.fileno())
            if self.path.exists():
                shutil.copymode(self.path, temp)  # a new workbook keeps mkstemp's 0600
            os.replace(temp, self.path)
        except BaseException:
            Path(temp).unlink(missing_ok=True)
            raise

    # ---- Reading ----

    def _sheet(self, name: str) -> Sheet:
        sheets = self._sheets()
        if name not in sheets:
            raise TableNotFound(f"no tab '{name}' in {self.path.name}")
        return sheets[name]

    def _sheets(self) -> dict[str, Sheet]:
        mtime = self._stat_mtime_ns()
        if self._cache is not None and self._cache[0] == mtime:
            return self._cache[1]
        workbook = self._load(read_only=True, data_only=True)
        try:
            sheets = {
                sheet.title: [
                    tuple(_raw(v) for v in row) for row in sheet.iter_rows(values_only=True)
                ]
                for sheet in workbook.worksheets
            }
        finally:
            workbook.close()
        self._cache = (mtime, sheets)
        return sheets

    def _load(self, *, read_only: bool, data_only: bool) -> Workbook:
        try:
            return openpyxl.load_workbook(self.path, read_only=read_only, data_only=data_only)
        except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as error:
            raise StorageError(f"cannot read {self.path.name}: {error}") from None

    def _stat_mtime_ns(self) -> int:
        try:
            return self.path.stat().st_mtime_ns
        except OSError:
            raise StorageError(f"workbook not found: {self.path}") from None


class _Change[T](Protocol):
    def __call__(self, workbook: Workbook, /) -> T: ...


class _FreeRows:
    """The rows an `Append` may fill, top down: from row 3, those whose columns are blank but
    for `auto:` ones (the formula buffer, 04 §1), then the rows after the last one."""

    def __init__(self, sheet: Worksheet, columns: Mapping[str, int]) -> None:
        self.sheet = sheet
        self.data = [i for key, i in columns.items() if not key.startswith(AUTO)]
        self.next = FIRST_DATA_ROW

    def take(self) -> int:
        while self.next <= self.sheet.max_row and not self._free(self.next):
            self.next += 1
        number = self.next
        self.next += 1
        return number

    def _free(self, number: int) -> bool:
        return all(self.sheet.cell(number, i + 1).value in (None, "") for i in self.data)


def _table(name: str, rows: Sheet) -> Table:
    """Keys from row 2, rows from row 3 (those with a non-blank cell)."""
    columns = _columns(name, rows[FIRST_DATA_ROW - 2] if len(rows) >= FIRST_DATA_ROW - 1 else ())
    data = []
    for number, row in enumerate(rows[FIRST_DATA_ROW - 1 :], start=FIRST_DATA_ROW):
        cells = {k: row[i] for k, i in columns.items() if i < len(row) and row[i] not in (None, "")}
        if any(not _blank(value) for value in cells.values()):  # kept as written (04 §3b)
            data.append(SheetRow(number, cells))
    numbers = {key: i + 1 for key, i in columns.items()}
    return Table(name, tuple(columns), tuple(data), numbers)


def _columns(tab: str, keys: Sequence[object]) -> dict[str, int]:
    """Column key (row 2) → 0-based column index. A key twice is an error."""
    columns: dict[str, int] = {}
    for i, key in enumerate(keys):
        if _blank(key):
            continue
        text = str(key).strip()
        if text in columns:
            first, second = (get_column_letter(c + 1) for c in (columns[text], i))
            raise StorageError(
                f"tab '{tab}': column key '{text}' is in row 2 twice ({first}2 and {second}2)"
            )
        columns[text] = i
    return columns


def _kv_columns(tab: str, headers: Sequence[object]) -> tuple[int, int]:
    """The `key` and `value` columns of a key/value tab, from row 1."""
    names = [str(h).strip().casefold() if not _blank(h) else "" for h in headers]
    if "key" not in names or "value" not in names:
        raise StorageError(f"tab '{tab}' needs the headers 'key' and 'value' in row 1")
    return names.index("key"), names.index("value")


def _worksheet(workbook: Workbook, name: str) -> Worksheet:
    if name not in workbook.sheetnames:
        raise TableNotFound(f"no tab '{name}' in the workbook")
    sheet = workbook[name]
    assert isinstance(sheet, Worksheet)
    return sheet


def _row_values(sheet: Worksheet, number: int) -> list[object]:
    if number > sheet.max_row:
        return []
    return [cell.value for cell in sheet[number]]


def _check_columns(tab: str, columns: Mapping[str, int], cells: Mapping[str, RawCell]) -> None:
    missing = [key for key in cells if key not in columns]
    if missing:
        raise StorageError(f"tab '{tab}' has no column {', '.join(missing)} in row 2")
    formulas = [key for key in cells if key.startswith(AUTO)]
    if formulas:
        raise StorageError(f"tab '{tab}': {', '.join(formulas)} is app-managed, not written")


def _check_unchanged(sheet: Worksheet, columns: Mapping[str, int], edit: Update) -> None:
    """The row still holds what was read, `auto:` columns aside (blank = absent)."""
    if edit.row < FIRST_DATA_ROW:
        raise StorageError(f"tab '{edit.tab}': row {edit.row} is a header row")
    for key, i in columns.items():
        if key.startswith(AUTO):
            continue
        now = _raw(sheet.cell(edit.row, i + 1).value)
        if _cell_value(now) != _cell_value(edit.expected.get(key)):
            raise WriteConflict(
                f"tab '{edit.tab}', row {edit.row} changed since it was read: reload and retry"
            )


def _write(cell: Cell | MergedCell, value: RawCell) -> None:
    """Write a value; a string is always a text cell, never a formula (04 §3b)."""
    if isinstance(cell, MergedCell):
        raise StorageError(f"cell {cell.coordinate} is inside merged cells: unmerge them to save")
    before = cell.number_format  # a format chosen in the sheet is kept
    cell.value = value  # openpyxl gives dates and times its own format here
    if isinstance(value, str):
        cell.data_type = "s"
        cell.number_format = "@"
    elif before not in ("General", "@"):
        cell.number_format = before
    else:
        cell.number_format = next((f for t, f in _FORMATS if isinstance(value, t)), "General")


def _cell_value(value: RawCell) -> RawCell:
    return None if value == "" else value


def _raw(value: object) -> RawCell:
    """openpyxl's cell value as one of ours: a Decimal is a float, rich text its plain text."""
    if value is None or isinstance(value, str | int | float | dt.date | dt.time | dt.timedelta):
        return value
    if isinstance(value, Decimal):
        return float(value)
    return str(value)  # rich text, or a formula object (not given with data_only)


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())
