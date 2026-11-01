# SPDX-License-Identifier: GPL-3.0-or-later
"""Derived fields compiled to per-row sheet formulas (03-module-framework.md §2.2b, D-018, D-053).

The sheets stay informative without the app: every derived field gets an `auto:` column whose
formula gives what the app computes, with the app's blank rules (D-043): a blank input makes a
blank result, ordering comparisons with a blank are false, aggregates skip blanks. The target is
the subset common to Google Sheets, Excel and LibreOffice: per-row formulas, no LET, LAMBDA or
array formulas. `Dialect` holds what differs: `.xlsx` needs `_xlfn.` before newer functions.

Records are matched across rows and tabs through a hidden numeric key column, `auto:_key` =
day serial x 10 + the code of the second key part (D-031): `prev`, `parent.x`, `children(...)`
and fields on another section tab of a split entity look rows up on it. An entry with no
translation (plugins, M5 names, `prev(x, 2)`, …) is app-only: it gets no column, and
`climblog module validate` lists it.

`formula_columns` says which `auto:` columns each tab gets (placement, label, number format);
once the backend knows where they are, `compile_formulas` writes their formulas.
"""

import dataclasses
import re
from collections.abc import Callable, Iterator, Mapping, Sequence
from dataclasses import dataclass

from openpyxl.utils import get_column_letter

from climblog.expr import Checked, Kind, NameInfo, Value
from climblog.expr.nodes import (
    Binary,
    BoolOp,
    Call,
    Children,
    Compare,
    Cond,
    Fields,
    Last,
    Literal,
    Name,
    Node,
    ParentRef,
    Prev,
    Setting,
    Unary,
    Window,
)
from climblog.modules.checks import derived_plan, recommendation_plan
from climblog.modules.derived import BandsStep, ExprStep, MapStep, Step, WhenStep
from climblog.modules.schema import (
    NEXT,
    BandsDerived,
    BaseField,
    BoolField,
    ChoiceField,
    ExprDerived,
    MapDerived,
    Module,
    NumberField,
    OrdinalScale,
    PluginDerived,
    WhenDerived,
)

AUTO = "auto:"
KEY_COLUMN = "auto:_key"
DELETED = "deleted"  # soft delete (04 §1): deleted rows are left out, as on read
_ROW = "\x00"  # the row number in a template
_YES = ('"yes"', '"y"', '"true"', '"si"', '"sì"', '"1"', "TRUE", "1")  # as `cells.py` reads
_NUMBER = '">-1E+307"'  # a criterion matching any number (formula blanks are text)
_HUGE = "1E+307"
_WEEKDAYS = ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")  # as `weekday()` gives
_CELL_LIKE = re.compile(r"^([a-z]{1,3}[0-9]+|r|c|r[0-9]+c[0-9]+)$")  # not valid as a name


@dataclass(frozen=True)
class Dialect:
    name: str
    prefix: str  # before functions newer than Excel 2007

    def fn(self, name: str) -> str:
        return f"{self.prefix}{name}" if name in ("MAXIFS", "MINIFS") else name


XLSX = Dialect("xlsx", "_xlfn.")
GOOGLE = Dialect("google", "")


@dataclass(frozen=True)
class Template:
    """A formula for any row: `at(5)` is the formula of sheet row 5."""

    text: str  # without the leading "="

    def at(self, row: int) -> str:
        return "=" + self.text.replace(_ROW, str(row))


@dataclass(frozen=True)
class FormulaColumn:
    tab: str
    key: str  # `auto:<derived field>`, or `auto:_key`
    label: str
    number_format: str
    levels: tuple[str, ...] | None = None  # a status: its levels, best first (for colours, D-054)
    hidden: bool = False
    hand_written: bool = False  # a `sheet_formula: "=..."`: not compared with the app


@dataclass(frozen=True)
class AppOnly:
    """A derived field without a sheet column, and why."""

    entity: str
    key: str
    reason: str

    def __str__(self) -> str:
        return f"{self.entity}.{self.key}: app only ({self.reason})"


ColumnMap = Mapping[str, Mapping[str, int]]
"""Tab → column key (row 2) → 1-based column: where the columns are in the workbook."""


class Unsupported(Exception):
    """An expression with no formula translation (yet)."""


def formula_columns(module: Module) -> tuple[list[FormulaColumn], list[AppOnly]]:
    """The `auto:` columns of every tab, in YAML order, and the derived fields that get none."""
    plans = _Plans(module)
    columns: list[FormulaColumn] = []
    app_only: list[AppOnly] = []
    for e_name in module.entities:
        placed, missing = _Entity(module, e_name, plans).placement()
        columns += placed
        app_only += missing
    return columns, app_only


def compile_formulas(
    module: Module, columns: ColumnMap, dialect: Dialect = XLSX
) -> dict[tuple[str, str], Template]:
    """The formula of each `auto:` column of `formula_columns`, as (tab, key) → template.

    `columns` gives where every column is, `auto:` ones included; a column it lacks makes the
    formulas that need it app-only for now (left out).
    """
    plans = _Plans(module)
    result: dict[tuple[str, str], Template] = {}
    for e_name in module.entities:
        result.update(_Entity(module, e_name, plans, columns, dialect).formulas())
    return result


def default_columns(module: Module) -> dict[str, dict[str, int]]:
    """Where the columns are in a workbook made by schema sync (`storage/layout.py`): the tab's
    fields, its `auto:` columns with the hidden key last, then `updated_at` and `deleted`."""
    placed, _ = formula_columns(module)
    tabs: dict[str, dict[str, int]] = {}
    for entity in module.entities.values():
        for tab, names in entity.tab_fields().items():
            autos = [c for c in placed if c.tab == tab]
            keys = [n for n in names if isinstance(entity.fields[n], BaseField)]
            keys += [c.key for c in autos if not c.hidden] + [c.key for c in autos if c.hidden]
            keys += ["updated_at", DELETED]
            tabs[tab] = {key: i for i, key in enumerate(keys, start=1)}
    return tabs


def recommendation_formulas(
    module: Module, columns: ColumnMap, refs: Mapping[str, str], dialect: Dialect = XLSX
) -> dict[str, str]:
    """The sheet formulas of the module's recommendation (D-076): one per value, then
    `headline` and `text` (as the legacy sheet shows it, "RECOVER — …"; blank when no rule
    holds). `refs` gives the cell of the day (`date`, e.g. `TODAY()` or `A3`) and of each value,
    where a value's formula reads the earlier ones. Raises `Unsupported` for what has no
    formula."""
    plan = recommendation_plan(module)
    if plan is None:
        return {}
    names = {"date": NameInfo("date"), **plan.values.names}
    compiler = _DayCompiler(module, _Plans(module), columns, refs, names, dialect)
    result = {step.key: "=" + compiler.value_of(step) for step in plan.values.steps}
    headline, text = '""', '""'
    for condition, head, message in reversed(plan.rules):
        shown = f"{head} — {message}" if head is not None else message
        if condition is None:
            headline, text = _text(head), _text(shown)
        else:
            test = compiler.condition(condition)
            headline = f"IF({test},{_text(head)},{headline})"
            text = f"IF({test},{_text(shown)},{text})"
    result["headline"], result["text"] = "=" + headline, "=" + text
    return result


# ---- Per entity ----


class _Plans:
    """Each entity's checked derived plan, computed once, and the column map of the dry run:
    every tab with its fields, system columns, `auto:_key` and every derived field of its entity
    (to tell what compiles before the workbook has the columns)."""

    def __init__(self, module: Module) -> None:
        self.plans = {name: derived_plan(module, name) for name in module.entities}
        steps = [step for plan in self.plans.values() for step in plan.steps]
        recommendation = recommendation_plan(module)
        trees = [node for step in steps for node in _walk(step)]
        if recommendation is not None:
            trees += [node for step in recommendation.values.steps for node in _walk(step)]
        self.read_by_last = {node.entity for node in trees if isinstance(node, Last)}
        self.dry: dict[str, dict[str, int]] = {}
        for entity in module.entities.values():
            autos = [AUTO + key for key in entity.derived]
            for tab, names in entity.tab_fields().items():
                keys = [*names, "updated_at", DELETED, KEY_COLUMN, *autos]
                self.dry[tab] = {key: i for i, key in enumerate(keys, start=1)}


@dataclass(frozen=True)
class _Placed:
    tab: str
    step: Step | None  # None: a hand-written `sheet_formula`
    override: str | None = None


class _Entity:
    """Placement and compilation of one entity's `auto:` columns."""

    def __init__(
        self,
        module: Module,
        name: str,
        plans: _Plans,
        columns: ColumnMap | None = None,
        dialect: Dialect = XLSX,
    ) -> None:
        self.module = module
        self.name = name
        self.entity = module.entities[name]
        self.plans = plans
        self.plan = plans.plans[name]
        self.dialect = dialect
        self.tabs = self.entity.tab_fields()
        self.main = self.entity.main_tab
        self.field_tab = {
            f: tab
            for tab, names in self.tabs.items()
            for f in names
            if f not in (self.entity.key or [])
        }
        self.key_code = _KeyCode(module, name)
        self._columns = columns

    # ---- Placement ----

    def needs_key(self) -> bool:
        """Whether the entity's tabs get `auto:_key`: rows are matched across tabs or rows."""
        entity = self.entity
        is_parent = any(
            other.parent is not None and other.parent.entity == self.name
            for other in self.module.entities.values()
        )
        uses_prev = any(_has(step, Prev) for step in self.plan.steps)
        # last() takes the latest row of a day by its key: needed when a day can have several.
        read_by_last = len(entity.key or []) > 1 and self.name in self.plans.read_by_last
        return (
            entity.is_split or is_parent or entity.parent is not None or uses_prev or read_by_last
        )

    def placement(self) -> tuple[list[FormulaColumn], list[AppOnly]]:
        placed, app_only = self._placed()
        columns = []
        if self.needs_key() and self.key_code.reason is None:
            for tab in self.tabs:
                columns.append(
                    FormulaColumn(tab, KEY_COLUMN, "key (auto, hidden)", "0", None, True)
                )
        for key, where in placed.items():
            step = where.step
            levels = step.kind.levels if step is not None else None
            fmt = _number_format(step) if step is not None else "General"
            hand = where.override is not None
            columns.append(
                FormulaColumn(where.tab, AUTO + key, _label(key), fmt, levels, hand_written=hand)
            )
        return columns, app_only

    def _placed(self) -> tuple[dict[str, _Placed], list[AppOnly]]:
        """Each derived field's tab, and those that get no column, both in YAML order.

        Entries are tried in dependency order, so one reading an app-only value is app-only too.
        """
        tab_of: dict[str, str] = {}
        found: dict[str, _Placed] = {}
        reasons: dict[str, str] = {}
        for step in self.plan.steps:  # what an entry reads comes first
            key, option = step.key, _sheet_formula(self.entity.derived[step.key])
            tab = self._tab_for(step, tab_of)
            if option is False:
                reasons[key] = "sheet_formula: false"
                continue
            compiler = _Compiler(self, self.plans.dry, tab, {**tab_of, key: tab})
            try:
                compiler.step(step) if option is True else compiler.override(str(option))
            except Unsupported as error:
                reasons[key] = str(error)
                continue
            found[key] = _Placed(tab, step, option if isinstance(option, str) else None)
            tab_of[key] = tab
        placed: dict[str, _Placed] = {}
        app_only = []
        for key, entry in self.entity.derived.items():
            option = _sheet_formula(entry)
            if key in found:
                placed[key] = found[key]
            elif isinstance(option, str) and self._writable(option):  # deferred, hand-written
                placed[key] = _Placed(self.main, None, option)
            else:
                reason = reasons.get(key) or self._reason(option, entry)
                app_only.append(AppOnly(self.name, key, reason))
        return placed, app_only

    def _writable(self, text: str) -> bool:
        try:
            _Compiler(self, self.plans.dry, self.main, {}).override(text)
        except Unsupported:
            return False
        return True

    def _reason(self, option: bool | str, entry: object) -> str:
        if option is False:
            return "sheet_formula: false"
        if isinstance(option, str):
            return "a placeholder of its sheet_formula is not a field with a column"
        return _deferred_reason(entry)

    def _tab_for(self, step: Step, tab_of: Mapping[str, str]) -> str:
        """The tab of everything the field reads on this record, or the anchor tab."""
        tabs = set()
        for name in _step_names(step):
            if name in self.field_tab:
                tabs.add(self.field_tab[name])
            elif name in tab_of:
                tabs.add(tab_of[name])
        return tabs.pop() if len(tabs) == 1 else self.main

    # ---- Formulas ----

    def formulas(self) -> dict[tuple[str, str], Template]:
        assert self._columns is not None
        placed, _ = self._placed()
        tab_of = {key: where.tab for key, where in placed.items()}
        result: dict[tuple[str, str], Template] = {}
        if self.needs_key() and self.key_code.reason is None:
            for tab in self.tabs:
                if KEY_COLUMN in self._columns.get(tab, {}):
                    result[(tab, KEY_COLUMN)] = Template(self.key_formula(tab))
        for key, where in placed.items():
            if AUTO + key not in self._columns.get(where.tab, {}):
                continue
            compiler = _Compiler(self, self._columns, where.tab, tab_of)
            try:
                text = (
                    compiler.override(where.override)
                    if where.override is not None
                    else compiler.step(where.step)  # type: ignore[arg-type]
                )
            except Unsupported:
                continue  # a column the workbook lacks; `check` reports it
            result[(where.tab, AUTO + key)] = Template(text)
        return result

    def key_formula(self, tab: str) -> str:
        """`auto:_key` on a tab: blank until the date is typed."""
        assert self._columns is not None
        cols = self._columns[tab]
        date = _ref(tab, cols, self.key_code.date_field, tab)
        code = self.key_code.formula(lambda f: _ref(tab, cols, f, tab))
        return f'IF({date}="","",INT({date})*10{code})'


class _KeyCode:
    """How a record's key becomes `auto:_key`: day x 10 + the code of the second part.

    A child's key is its parent's (the parent key columns it repeats, and `fixed`), so children
    are counted with one criterion (D-031). The second part is an integer (blank = its default)
    or a closed choice (its position, 1-based, as in the mock-ups).
    """

    def __init__(self, module: Module, e_name: str) -> None:
        entity = module.entities[e_name]
        self.module = module
        self.reason: str | None = None
        self.fixed: Value = None
        self.part: str | None = None
        self.choices: list[str] = []
        self.default: Value = None
        key = entity.key or []
        owner, parts = entity, key
        if entity.parent is not None:
            owner = module.entities[entity.parent.entity]
            parts = owner.key or []
        self.date_field = parts[0] if parts else entity.date_field
        if not parts:
            self.reason = "rows have no natural key to be matched on"
            return
        if len(parts) > 2:
            self.reason = "a key of more than a date and one other part"
            return
        if len(parts) == 1:
            return
        second = parts[1]
        field = owner.fields[second]
        if entity.parent is not None and second in entity.parent.fixed:
            self.fixed = entity.parent.fixed[second]
        else:
            self.part = second
            field = entity.fields.get(second, field)
        if isinstance(field, NumberField) and field.type == "integer":
            self.default = field.default if isinstance(field.default, int) else None
            if self.fixed is not None and not isinstance(self.fixed, int):
                self.reason = "the fixed key part is not a whole number"
        elif isinstance(field, ChoiceField) and not field.allow_other:
            choices = field.choices
            if choices is None:
                scale = module.scales[field.scale or ""]
                assert isinstance(scale, OrdinalScale)
                choices = scale.choices
            self.choices = list(choices)
            if len(self.choices) > 9:
                self.reason = "the key's choice has more than 9 values"
        else:
            self.reason = "the second key part is not a whole number or a closed choice"

    def formula(self, ref: Callable[[str], str]) -> str:
        """`+code` to add after `INT(date)*10`, or "" for a date-only key."""
        if self.fixed is not None:
            code = self.choices.index(str(self.fixed)) + 1 if self.choices else self.fixed
            return f"+{code}"
        if self.part is None:
            return ""
        cell = ref(self.part)
        if self.choices:
            codes = ",".join(f'"{c}"' for c in self.choices)
            return f"+IFERROR(MATCH({cell},{{{codes}}},0),0)"
        default = self.default if self.default is not None else 0
        return f'+IF({cell}="",{default},{cell})'


# ---- The compiler ----


@dataclass(frozen=True)
class F:
    """A compiled value: its formula when it isn't blank, and when it is blank.

    Keeping the two apart lets a blank test be written once at the top of an expression
    (`IF(COUNT(D3,E3)=0,"",ROUND(11-AVERAGE(D3,E3),1))`) instead of around every part.
    """

    core: str  # the value, valid when `test` is false
    test: str | None = None  # true when the value is blank; None: never blank
    type: str = "any"
    empty: bool = False  # `core` itself is "" (or an empty cell) when blank
    reference: str | None = None  # the same value as a cell reference, for AVERAGE, COUNT, …

    @property
    def text(self) -> str:
        """The value as a formula of its own: blank ("") when `test` holds."""
        if self.empty or not self.test:
            return self.core
        return f'IF({self.test},"",{self.core})'

    def is_blank(self) -> str:
        return self.test or "FALSE"

    def not_blank(self) -> str:
        if self.test is None:
            return "TRUE"
        if _SIMPLE_BLANK.match(self.test):
            return self.test[:-3] + '<>""'
        return f"NOT({self.test})"


_SIMPLE_BLANK = re.compile(r'^[A-Z]{1,3}\x00=""$')  # D3="" (its negation is D3<>"")


def _cell(ref: str) -> F:
    return F(ref, f'{ref}=""', empty=True, reference=ref)


def _either(*tests: str | None) -> str | None:
    found = list(dict.fromkeys(t for t in tests if t))
    if not found:
        return None
    return found[0] if len(found) == 1 else f"OR({','.join(found)})"


@dataclass(frozen=True)
class _Item:
    """One value of a list (`fields(...)`): counted only where `include` holds (`since`)."""

    value: F
    include: str | None = None


class _Compiler:
    def __init__(
        self, entity: _Entity, columns: ColumnMap, tab: str, tab_of: Mapping[str, str]
    ) -> None:
        self.e = entity
        self.columns = columns
        self.tab = tab
        self.tab_of = tab_of
        self.dialect = entity.dialect
        self.module = entity.module
        self.entity = entity.entity

    # ---- Entry points ----

    def step(self, step: Step) -> str:
        """The column's formula: blank until the row has a date (the formula buffer)."""
        value = self._step(step)
        date = self.cell(self.tab, self.entity.date_field)
        if value.test == f'{date}=""':
            return value.text  # its own blank test already is the date's
        return f'IF({date}="","",{value.text})'

    def override(self, text: str) -> str:
        """A hand-written `sheet_formula`: `{field}` is that field's cell in the row (looked up
        on `auto:_key` when it is on another tab)."""
        body = text.removeprefix("=")
        return re.sub(r"\{([a-z][a-z0-9_]*)\}", lambda m: self.raw(m.group(1)), body)

    def raw(self, name: str) -> str:
        if name in self.entity.fields:
            key = name in (self.entity.key or [])
            tab = self.tab if key else self.e.field_tab.get(name, self.tab)
            return self.lookup(tab, name).core
        if name in self.tab_of:
            return self.lookup(self.tab_of[name], AUTO + name).core
        raise Unsupported(f"{{{name}}} is not a field or a derived value with a column")

    def _step(self, step: Step) -> F:
        match step:
            case ExprStep(expr=expr):
                value = self.expr(expr)
            case BandsStep(value=source, bands=bands, otherwise=otherwise):
                v = self.expr(source)
                result = _text(otherwise)
                for limit, label in reversed(bands):
                    upper = (
                        _number(limit)
                        if isinstance(limit, int | float)
                        else _wrap(self.expr(limit).text)
                    )
                    result = f"IF({_wrap(v.core)}<={upper},{_text(label)},{result})"
                value = F(result, v.test, "text")
            case WhenStep(cases=cases, otherwise=otherwise):
                result = _text(otherwise)
                for condition, label in reversed(cases):
                    result = f"IF({self.condition(condition)},{_text(label)},{result})"
                value = F(result, None, "text")
            case MapStep(source=source, mapping=mapping):
                s = self.expr(source)
                result = '""'
                for given, text in reversed(list(mapping.items())):
                    result = f"IF({_wrap(s.text)}={_text(given)},{_text(text)},{result})"
                value = F(result, None, "text")
            case _:  # pragma: no cover
                raise Unsupported(f"unknown entry {type(step).__name__}")
        if step.require is not None:
            value = F(f'IF({self.condition(step.require)},{value.text},"")', None, value.type)
        return value

    # ---- Expressions ----

    def expr(self, checked: Checked) -> F:
        return self.node(checked.tree)

    def condition(self, checked: Checked) -> str:
        return self.truth(self.node(checked.tree))

    def truth(self, value: F) -> str:
        """A yes/no for IF and AND: only TRUE counts (a blank is no)."""
        assert value.type == "bool"  # `check`: conditions are yes/no
        return value.core

    def node(self, node: Node) -> F:
        match node:
            case Literal(value=bool() as flag):
                return F("TRUE" if flag else "FALSE", None, "bool")
            case Literal(value=int() | float() as number):
                return F(_number(number), None, "number")
            case Literal(value=str() as text):
                return F(_text(text), None, "text")
            case Name(name=name):
                return self.value(name)
            case Setting(name=name):
                if _CELL_LIKE.match(name):
                    raise Unsupported(f"setting '{name}' looks like a cell, not a name")
                if self.module.settings[name].type == "bool":  # written Yes/No (D-052)
                    return F(_truthy(name), None, "bool")
                return F(name, None, "any")
            case Unary(op="-", operand=operand):
                v = self.node(operand)
                return F(f"-{_wrap(v.core)}", v.test, "number")
            case Unary(operand=operand):
                return F(f"NOT({self.truth(self.node(operand))})", None, "bool")
            case Binary(op=op, left=left, right=right):
                return self.binary(op, self.node(left), self.node(right))
            case Compare(op=op, left=left, right=right):
                return self.compare(op, self.node(left), self.node(right))
            case BoolOp(op=op, values=values):
                parts = ",".join(self.truth(self.node(v)) for v in values)
                return F(f"{op.upper()}({parts})", None, "bool")
            case Cond(test=test, then=then, orelse=orelse):
                a, b = self.node(then), self.node(orelse)
                core = f"IF({self.truth(self.node(test))},{a.text},{b.text})"
                blank = f'{_wrap(core)}=""' if a.test or b.test else None
                return F(core, blank, a.type if a.type == b.type else "any")
            case Prev(name=name, n=n):
                return self.prev(name, n)
            case ParentRef(name=name):
                return self.parent(name)
            case Last(entity=entity, field=name, before=before, inclusive=inclusive):
                return self.last(entity, name, self.node(before), inclusive)
            case Call():
                return self.call(node)
            case Fields() | Children() | Window():  # pragma: no cover - lists only in calls
                raise Unsupported("a list of values outside an aggregate")
        raise Unsupported(f"no formula for {type(node).__name__}")  # pragma: no cover

    def binary(self, op: str, left: F, right: F) -> F:
        core = f"{_wrap(left.core)}{op}{_wrap(right.core)}"
        test = _either(left.test, right.test)
        if op == "/" and not (right.test is None and _nonzero(right.core)):
            test = _either(test, f"{_wrap(right.core)}=0")
        return F(core, test, "number")

    def compare(self, op: str, left: F, right: F) -> F:
        if op in ("==", "!="):
            equal = f"{_wrap(left.core)}={_wrap(right.core)}"
            if left.test and right.test:
                same = (
                    f"IF(AND({left.test},{right.test}),TRUE,"
                    f"IF(OR({left.test},{right.test}),FALSE,{equal}))"
                )
            elif left.test or right.test:
                side, other = (left, right) if left.test else (right, left)
                if _is_label(other.core):  # a blank is never equal to a text like "LOW"
                    same = f"{_wrap(_label_side(side))}={other.core}"
                else:
                    same = f"AND({side.not_blank()},{equal})"
            else:
                same = equal
            return F(same if op == "==" else f"NOT({same})", None, "bool")
        plain = f"{_wrap(left.core)}{op}{_wrap(right.core)}"
        tests = [v.not_blank() for v in (left, right) if v.test]
        return F(f"AND({','.join(tests)},{plain})" if tests else plain, None, "bool")

    # ---- Names ----

    def value(self, name: str) -> F:
        """A field or derived value of this record."""
        if name not in self.entity.fields:
            tab = self.tab_of.get(name)
            if tab is None:
                raise Unsupported(f"it reads '{name}', which is app only")
            return self._derived(name, self.lookup(tab, AUTO + name))
        field = self.entity.fields[name]
        assert isinstance(field, BaseField)
        tab = self.tab if name in (self.entity.key or []) else self.e.field_tab.get(name, self.tab)
        return _field_value(field, self.lookup(tab, name))

    def _derived(self, name: str, found: F) -> F:
        info = self.e.plan.names[name]
        if info.type == "bool":
            return F(f"({found.core}=TRUE)", None, "bool")
        return dataclasses.replace(found, type=info.type)

    def lookup(self, tab: str, key: str) -> F:
        """A column's value in this record's row: same tab, or found on `auto:_key`."""
        if tab == self.tab:
            return _cell(self.cell(tab, key))
        own = self.own_key()
        return self.index(tab, key, f"MATCH({own},{self.column(tab, KEY_COLUMN)},0)")

    def index(self, tab: str, key: str, match: str) -> F:
        """The cell of `key` at a matched row: blank when empty or not found.

        As a reference, a row not found points at row 2 (the key, text), which AVERAGE, COUNT,
        MIN and MAX skip like a blank.
        """
        column = self.column(tab, key)
        cell = f"INDEX({column},{match})"
        found = f'IFERROR(IF({cell}="","",{cell}),"")'
        reference = f"INDEX({column},IFERROR({match},2))"
        return F(found, f'{found}=""', empty=True, reference=reference)

    def cell(self, tab: str, key: str) -> str:
        return _ref(tab, self.columns_of(tab), key, self.tab)

    def column(self, tab: str, key: str) -> str:
        col = _letter(self.columns_of(tab), key, tab)
        return f"{_sheet(tab)}!${col}:${col}"

    def columns_of(self, tab: str) -> Mapping[str, int]:
        if tab not in self.columns:
            raise Unsupported(f"tab '{tab}' is not in the workbook")
        return self.columns[tab]

    # ---- Other records ----

    def own_key(self) -> str:
        if self.e.key_code.reason is not None:
            raise Unsupported(self.e.key_code.reason)
        return self.cell(self.tab, KEY_COLUMN)

    def prev(self, name: str, n: int) -> F:
        if n != 1:
            raise Unsupported("prev() further back than one record")
        if self.entity.parent is not None:
            raise Unsupported("prev() on a child entity (its key column is its parent's)")
        own = self.own_key()
        if self.e.key_code.choices:
            raise Unsupported("prev() with a choice in the key (the app orders it by name)")
        if name in self.entity.fields:
            tab, target = self.e.field_tab.get(name, self.tab), name
        elif name in self.tab_of:
            tab, target = self.tab_of[name], AUTO + name
        else:
            raise Unsupported(f"it reads '{name}', which is app only")
        keys = self.column(tab, KEY_COLUMN)
        latest = (
            f'{self.dialect.fn("MAXIFS")}({keys},{keys},"<"&{own},{self.column(tab, DELETED)},"")'
        )
        found = self.index(tab, target, f"MATCH({latest},{keys},0)")
        if name in self.entity.fields:
            field = self.entity.fields[name]
            assert isinstance(field, BaseField)
            return _field_value(field, found)
        return self._derived(name, found)

    def parent(self, name: str) -> F:
        link = self.entity.parent
        assert link is not None
        parent = _Entity(self.module, link.entity, self.e.plans)
        tab = parent.field_tab.get(name, parent.main)
        own = self.own_key()
        found = self.index(tab, name, f"MATCH({own},{self.column(tab, KEY_COLUMN)},0)")
        field = parent.entity.fields[name]
        assert isinstance(field, BaseField)
        return _field_value(field, found)

    def last(self, e_name: str, name: str, before: F, inclusive: bool = False) -> F:
        """The field (or derived value) in the latest record dated before the day, or on it
        (`inclusive`). Records of one day are told apart by `auto:_key` (date, then key), as the
        app orders them, when it has one."""
        other = _Entity(self.module, e_name, self.e.plans)
        tab, key = self._column_of(other, name)
        dates = self.column(tab, other.entity.date_field)
        order = dates
        if other.needs_key() and other.key_code.reason is None and not other.key_code.choices:
            order = self.column(tab, KEY_COLUMN)
        day = f"INT({_wrap(before.core)})"
        deleted = self.column(tab, DELETED)
        op = "<=" if inclusive else "<"
        latest = f'{self.dialect.fn("MAXIFS")}({order},{dates},"{op}"&{day},{deleted},"")'
        value = self._value_of(other, name, self.index(tab, key, f"MATCH({latest},{order},0)"))
        return F(value.core, _either(before.test, value.test), value.type)

    def window(self, node: Window, func: str, equals: F | None) -> F:
        """`count/sum/min/max/mean` of `window(e.x, days=n, until=d)`: the *IFS functions over
        the rows dated from d - n to d; `count(..., equals=v)` counts the values equal to v."""
        other = _Entity(self.module, node.entity, self.e.plans)
        tab, key = self._column_of(other, node.field)
        until = self.node(node.until)
        dates = self.column(tab, other.entity.date_field)
        end = f"INT({_wrap(until.core)})"
        criteria = (
            f'{dates},">="&({end}-{node.days}),{dates},"<="&{end},{self.column(tab, DELETED)},""'
        )
        values = self.column(tab, key)
        numbers = f"COUNTIFS({criteria},{values},{_NUMBER})"
        if func == "count":
            if equals is not None:  # "=" & the value: a text like ">5" is not a condition
                core = f'COUNTIFS({criteria},{values},"="&{_wrap(equals.core)})'
                if equals.test is not None:  # a blank equals nothing (as in the app)
                    core = f"IF({equals.test},0,{core})"
            else:
                kind = self._type_of(other, node.field)
                if kind not in ("number", "date"):
                    text = '"<>"' if key == node.field else '"?*"'
                    numbers = f"COUNTIFS({criteria},{values},{text})"
                core = numbers
            if until.test is not None:  # no day: no rows (as in the app)
                core = f"IF({until.test},0,{core})"
            return F(core, None, "number")
        native = {"sum": "SUMIFS", "min": "MINIFS", "max": "MAXIFS", "mean": "AVERAGEIFS"}[func]
        core = f"{self.dialect.fn(native)}({values},{criteria})"
        return F(core, _either(until.test, f"{numbers}=0"), "number")

    def _column_of(self, other: "_Entity", name: str) -> tuple[str, str]:
        """The tab and column key of another entity's field or derived value."""
        if name in other.entity.fields:
            return other.field_tab.get(name, other.main), name
        return self._child_derived_tab(other, name), AUTO + name

    def _type_of(self, other: "_Entity", name: str) -> str:
        field = other.entity.fields.get(name)
        return _type(field) if isinstance(field, BaseField) else self._kind(other, name)

    def _value_of(self, other: "_Entity", name: str, found: F) -> F:
        """A found cell of another entity's field or derived value, typed as the app reads it."""
        field = other.entity.fields.get(name)
        if isinstance(field, BaseField):
            return _field_value(field, found)
        info = other.plan.names[name]
        if info.type == "bool":
            return F(f"({found.core}=TRUE)", None, "bool")
        return dataclasses.replace(found, type=info.type)

    def children(self, node: Children, func: str) -> F:
        """`count/sum/min/max/mean` of `children(e)` or `children(e).x`."""
        child = _Entity(self.module, node.entity, self.e.plans)
        if child.key_code.reason is not None:
            raise Unsupported(child.key_code.reason)
        own = self.own_key()
        if node.field is None:
            tab = child.main
        elif node.field in child.entity.fields:
            tab = child.field_tab.get(node.field, child.main)
        else:
            tab = self._child_derived_tab(child, node.field)
        # The child's key holds its `fixed` values (they are parent key parts, D-042), so a
        # parent row of another kind never matches.
        criteria = f'{self.column(tab, KEY_COLUMN)},{own},{self.column(tab, DELETED)},""'
        if node.field is None:
            assert func == "count"  # `check`: only count() takes the rows themselves
            return F(f"COUNTIFS({criteria})", None, "number")
        stored = node.field in child.entity.fields
        key = node.field if stored else AUTO + node.field
        values = self.column(tab, key)
        numbers = f"COUNTIFS({criteria},{values},{_NUMBER})"
        if func == "count":
            # Non-blank values: any number; text and Yes/No by type ("" formulas don't count).
            field = child.entity.fields.get(node.field)
            kind = _type(field) if isinstance(field, BaseField) else self._kind(child, node.field)
            if kind not in ("number", "date"):
                text = '"<>"' if stored else '"?*"'
                numbers = f"COUNTIFS({criteria},{values},{text})"
            return F(numbers, None, "number")
        native = {"sum": "SUMIFS", "min": "MINIFS", "max": "MAXIFS", "mean": "AVERAGEIFS"}[func]
        core = f"{self.dialect.fn(native)}({values},{criteria})"
        return F(core, f"{numbers}=0", "number")

    def _kind(self, child: _Entity, name: str) -> str:
        return child.plan.names[name].type

    def _child_derived_tab(self, child: _Entity, name: str) -> str:
        placed, _ = child._placed()
        if name not in placed:
            raise Unsupported(f"it reads {child.name}.{name}, which is app only")
        return placed[name].tab

    # ---- Functions ----

    def call(self, node: Call) -> F:
        func, args = node.func, node.args
        aggregate = func in ("count", "sum", "min", "max", "mean")
        options = dict(node.kwargs)
        if aggregate and len(args) == 1 and isinstance(args[0], Window):
            if "missing_as" in options:
                raise Unsupported("missing_as with window(...)")
            equals = options.get("equals")
            return self.window(args[0], func, self.node(equals) if equals is not None else None)
        if aggregate and len(args) == 1 and isinstance(args[0], Children):
            if node.kwargs:
                raise Unsupported("missing_as or equals with children(...)")
            return self.children(args[0], func)
        if any(isinstance(a, Children | Window) for a in args):
            raise Unsupported("children(...) or window(...) mixed with other values")
        if "equals" in options:
            raise Unsupported("count(..., equals=) of the record's own values")
        match func:
            case "mean" | "min" | "max" | "sum":
                fill = dict(node.kwargs).get("missing_as")
                return self.aggregate(func, self.items(args), self.node(fill) if fill else None)
            case "count":
                return F(self.count(self.items(args), numbers=False), None, "number")
            case "all_present":
                return F(self.all_present(self.items(args)), None, "bool")
            case "notnull":
                return F(self.node(args[0]).not_blank(), None, "bool")
            case "coalesce":
                values = [self.node(a) for a in args]
                core = '""'
                for v in reversed(values):
                    core = f"IF({v.not_blank()},{v.core},{core})"
                return F(core, f'{core}=""', values[0].type)
            case "concat":
                return F("&".join(_wrap(self.node(a).text) for a in args), None, "text")
            case "round":
                v = self.node(args[0])
                digits = self.node(args[1]).core if len(args) > 1 else "0"
                return F(f"ROUND({v.core},{digits})", v.test, "number")
            case "flip":
                return self.flip(self.node(args[0]), node)
            case "worst_of":
                return self.worst_of(node)
            case "days_between":
                a, b = self.node(args[0]), self.node(args[1])
                core = f"INT({_wrap(b.core)})-INT({_wrap(a.core)})"
                return F(core, _either(a.test, b.test), "number")
            case "weekday":
                d = self.node(args[0])
                names = ",".join(f'"{w}"' for w in _WEEKDAYS)
                return F(f"CHOOSE(WEEKDAY({d.core},2),{names})", d.test, "text")
            case "month":  # 2026-10, without TEXT() (its format codes depend on the locale)
                d = self.node(args[0])
                core = f'YEAR({d.core})&"-"&RIGHT("0"&MONTH({d.core}),2)'
                return F(core, d.test, "text")
            case "incomplete":  # ", Physical, Reflect" without its first ", "
                parts = "&".join(
                    f'IF(NOT({self.truth(self.node(value))}),{_text(", " + name)},"")'
                    for name, value in node.kwargs
                )
                return F(f"MID({parts},3,999)", None, "text")
        raise Unsupported(f"{func}() has no formula yet")  # pragma: no cover - all have a case

    def flip(self, value: F, node: Call, turn: bool | None = None) -> F:
        """The value turned to higher-is-better: only if its scale is lower-is-better."""
        if not (node.flips[0] if turn is None else turn):
            return value
        assert node.bounds is not None
        low, high = node.bounds
        return F(f"{_number(low + high)}-{_wrap(value.core)}", value.test, "number")

    def items(self, args: Sequence[Node]) -> list[_Item]:
        """The values an aggregate works on, each with when it counts (`since`)."""
        items: list[_Item] = []
        for arg in args:
            if isinstance(arg, Fields):
                items += self.fields(arg)
            elif isinstance(arg, Call) and arg.func == "flip_all":
                inner = arg.args[0]
                assert isinstance(inner, Fields)
                items += [
                    _Item(self.flip(i.value, arg, turn), i.include)
                    for i, turn in zip(self.fields(inner), arg.flips, strict=True)
                ]
            else:
                items.append(_Item(self.node(arg)))
        return items

    def fields(self, node: Fields) -> list[_Item]:
        date = self.cell(self.tab, node.date_field)
        items = []
        for ref in node.refs:
            include = None
            if ref.since is not None:
                since = ref.since
                include = f"INT({date})>=DATE({since.year},{since.month},{since.day})"
            items.append(_Item(self.value(ref.name), include))
        return items

    def _refs(self, items: Sequence[_Item]) -> str | None:
        """`D3,E3,F3` when every item is a cell (or a looked-up cell) counted always: the
        native functions then skip blanks and text themselves."""
        if items and all(i.include is None and i.value.reference for i in items):
            return ",".join(i.value.reference or "" for i in items)
        return None

    def count(self, items: Sequence[_Item], *, numbers: bool) -> str:
        """How many items are not blank (`numbers`: they are numbers, COUNT may be used)."""
        refs = self._refs(items)
        if numbers and refs is not None:
            return f"COUNT({refs})"
        terms = []
        for item in items:
            present = f"({item.value.not_blank()})"
            terms.append(present if item.include is None else f"({item.include})*{present}")
        return "+".join(terms) if terms else "0"

    def all_present(self, items: Sequence[_Item]) -> str:
        if all(i.value.type == "number" for i in items) and (refs := self._refs(items)):
            return f"COUNT({refs})={len(items)}"
        tests = [
            i.value.not_blank()
            if i.include is None
            else f"OR(NOT({i.include}),{i.value.not_blank()})"
            for i in items
        ]
        return f"AND({','.join(tests)})" if tests else "TRUE"

    def aggregate(self, func: str, items: Sequence[_Item], fill: F | None) -> F:
        """mean/min/max/sum: blanks left out (or filled), blank when all are blank (D-043)."""
        count = self.count(items, numbers=True)
        test = f"{count}=0"
        refs = self._refs(items)
        if fill is None and refs is not None:
            native = {"mean": "AVERAGE", "min": "MIN", "max": "MAX", "sum": "SUM"}[func]
            return F(f"{native}({refs})", test, "number")
        if fill is not None:
            items = [
                _Item(F(f"IF({i.value.is_blank()},{_wrap(fill.core)},{i.value.core})"), i.include)
                for i in items
            ]
            counted = [i.include or "TRUE" for i in items]
        else:
            counted = [self._counts(i) for i in items]
        if func in ("sum", "mean"):
            terms = [_when(c, i.value.core, "0") for c, i in zip(counted, items, strict=True)]
            total = "+".join(terms)
            if func == "sum":
                return F(total, test, "number")
            n = count if fill is None else "+".join(f"({c})" for c in counted)
            return F(f"({total})/({n})", test, "number")
        sentinel = _HUGE if func == "min" else f"-{_HUGE}"
        values = ",".join(
            _when(c, i.value.core, sentinel) for c, i in zip(counted, items, strict=True)
        )
        return F(f"{func.upper()}({values})", test, "number")

    def _counts(self, item: _Item) -> str:
        present = item.value.not_blank()
        return present if item.include is None else f"AND({item.include},{present})"

    def worst_of(self, node: Call) -> F:
        assert node.levels is not None
        values = [self.node(a) for a in node.args]
        core = '""'
        for level in node.levels:  # best first, so the worst ends up outermost
            tests = ",".join(f"{_wrap(_label_side(v))}={_text(level)}" for v in values)
            core = f"IF(OR({tests}),{_text(level)},{core})"
        return F(core, None, "text")


# ---- Helpers ----


def _field_value(field: BaseField, cell: F) -> F:
    """A stored field as the app reads it: blank = its default; Yes/No as TRUE/FALSE."""
    if isinstance(field, BoolField):
        truth = _truthy(cell.core)
        if field.default is not None:
            default = "TRUE" if field.default else "FALSE"
            return F(f"IF({cell.is_blank()},{default},{truth})", None, "bool")
        return F(truth, cell.test, "bool")
    default = field.default
    plain = isinstance(default, int | float | str) and not isinstance(default, bool)
    if plain and default != NEXT:
        literal = _text(default) if isinstance(default, str) else _number(default)
        return F(f"IF({cell.is_blank()},{literal},{cell.core})", None, _type(field))
    return dataclasses.replace(cell, type=_type(field))


def _truthy(value: str) -> str:
    """Yes as the app reads it from a cell (Y, true, Sì, 1, TRUE; any case)."""
    return "OR(" + ",".join(f"{_wrap(value)}={token}" for token in _YES) + ")"


def _type(field: BaseField) -> str:
    kind = getattr(field, "type", "text")
    if kind in ("number", "integer", "scale", "duration", "percent"):
        return "number"
    if kind in ("date", "datetime"):
        return "date"
    return "text"


def _label_side(value: F) -> str:
    """The value to compare with a non-empty text: a reference when there is one (a row not
    found points at the key row, never equal to a label), else the value itself."""
    return value.reference or value.text


def _is_label(core: str) -> bool:
    return len(core) > 2 and core.startswith('"') and core.endswith('"')


def _when(condition: str, value: str, otherwise: str) -> str:
    return value if condition == "TRUE" else f"IF({condition},{value},{otherwise})"


def _nonzero(core: str) -> bool:
    try:
        return float(core) != 0
    except ValueError:
        return False


def _ref(tab: str, columns: Mapping[str, int], key: str, current: str) -> str:
    col = _letter(columns, key, tab)
    prefix = "" if tab == current else f"{_sheet(tab)}!"
    return f"{prefix}{col}{_ROW}"


def _letter(columns: Mapping[str, int], key: str, tab: str) -> str:
    if key not in columns:
        raise Unsupported(f"tab '{tab}' has no column '{key}'")
    return get_column_letter(columns[key])


def _sheet(tab: str) -> str:
    """A tab name as formulas write it: quoted only when it needs it ('Hang sets', Sessions)."""
    if _BARE_SHEET.match(tab) and not _LOOKS_LIKE_CELL.match(tab):
        return tab
    return "'" + tab.replace("'", "''") + "'"


_BARE_SHEET = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_LOOKS_LIKE_CELL = re.compile(r"^([A-Za-z]{1,3}[0-9]+|[RrCc]([0-9]+)?|[Rr][0-9]+[Cc][0-9]+)$")


def normalize(formula: str) -> str:
    """A formula as spreadsheets may rewrite it on save, for comparing (drift check, D-054):
    no `_xlfn.`, `TRUE()` → TRUE, sheet names unquoted where possible, upper case and no
    spaces outside strings."""
    parts = re.split(r'("(?:[^"]|"")*")', formula)  # strings at odd positions
    out = []
    for i, part in enumerate(parts):
        if i % 2:
            out.append(part)
            continue
        part = part.upper().replace(" ", "").replace("_XLFN.", "")
        part = re.sub(r"\b(TRUE|FALSE)\(\)", r"\1", part)
        part = re.sub(r"'([A-Z_][A-Z0-9_]*)'!", r"\1!", part)
        out.append(part)
    return "".join(out)


def _text(value: str | None) -> str:
    if value is None:
        return '""'
    return '"' + value.replace('"', '""') + '"'


def _number(value: float) -> str:
    if isinstance(value, bool):  # pragma: no cover - checked types never give one here
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return repr(value) if isinstance(value, float) else str(value)


def _wrap(text: str) -> str:
    """Parenthesise a compound formula part (an operator outside any parentheses)."""
    depth, quoted = 0, False
    for i, char in enumerate(text):
        if char == '"':
            quoted = not quoted
        elif quoted:
            continue
        elif char == "(":
            depth += 1
        elif char == ")":
            depth -= 1
        elif depth == 0 and char in "+-*/&=<>" and i > 0:
            return f"({text})"
    return text


def _number_format(step: Step) -> str:
    """The column's number format from the result type (a date difference is no date)."""
    kind: Kind = step.kind
    if kind.type == "date":
        return "DD/MM/YYYY"
    if kind.type == "time":
        return "HH:MM"
    if kind.type != "number":
        return "General"
    if isinstance(step, ExprStep):
        tree = step.expr.tree
        if isinstance(tree, Call) and tree.func == "round":
            digits = tree.args[1] if len(tree.args) > 1 else None
            places = (
                digits.value if isinstance(digits, Literal) and isinstance(digits.value, int) else 0
            )
            return "0." + "0" * places if places else "0"
        if isinstance(tree, Call) and tree.func in ("count", "days_between"):
            return "0"
    return "General"


def _label(key: str) -> str:
    return key.replace("_", " ").capitalize()


def _sheet_formula(entry: object) -> bool | str:
    if isinstance(entry, ExprDerived | BandsDerived | WhenDerived | MapDerived | PluginDerived):
        return entry.sheet_formula
    return True


def _deferred_reason(entry: object) -> str:
    if isinstance(entry, PluginDerived):
        return "a plugin"
    return "needs program, reference tables or another app-only value"


def _has(step: Step, kind: type) -> bool:
    return any(isinstance(n, kind) for n in _walk(step))


def _walk(step: Step) -> Iterator[Node]:
    trees = [c.tree for c in _checked(step)]
    for tree in trees:
        yield from _nodes(tree)


def _checked(step: Step) -> Iterator[Checked]:
    if step.require is not None:
        yield step.require
    match step:
        case ExprStep(expr=expr):
            yield expr
        case BandsStep(value=value, bands=bands):
            yield value
            for limit, _ in bands:
                if not isinstance(limit, int | float):
                    yield limit
        case WhenStep(cases=cases):
            for condition, _ in cases:
                yield condition
        case MapStep(source=source):
            yield source


def _nodes(node: object) -> Iterator[Node]:
    if isinstance(node, Node):
        yield node
        for part in dataclasses.fields(node):
            yield from _nodes(getattr(node, part.name))
    elif isinstance(node, tuple):
        for item in node:
            yield from _nodes(item)


def _step_names(step: Step) -> set[str]:
    """The fields and derived values of this record a step reads directly."""
    names: set[str] = set()
    for checked in _checked(step):
        names |= checked.names
    return names


class _DayCompiler(_Compiler):
    """Formulas for a day rather than a record (the recommendation): its names are cells the
    caller gives; other records are read with `last` and `window` only."""

    def __init__(
        self,
        module: Module,
        plans: _Plans,
        columns: ColumnMap,
        refs: Mapping[str, str],
        names: Mapping[str, NameInfo],
        dialect: Dialect,
    ) -> None:
        anchor = _Entity(module, next(iter(module.entities)), plans, columns, dialect)
        super().__init__(anchor, columns, "", {})
        self.refs = refs
        self.names = names

    def value_of(self, step: Step) -> str:
        return self._step(step).text

    def value(self, name: str) -> F:
        if name not in self.refs:
            raise Unsupported(f"no cell for '{name}'")
        return dataclasses.replace(_cell(self.refs[name]), type=self.names[name].type)

    def own_key(self) -> str:
        raise Unsupported("prev(), parent and children() need a record")
